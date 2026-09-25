"""Resolve an IPO issuer cohort from runtime input and a workbook."""
from __future__ import annotations

import datetime as dt
import os
import re
from collections import Counter
from copy import copy
from pathlib import Path
from typing import Any

import yaml

from listing_reports import reports_for_interval
from sample_builder import audit_candidate, load_nlr_candidates

ROOT = Path(__file__).resolve().parent.parent
WS = ROOT.parent


def _slug(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", value).strip("_-")


def _cohort_for_period(start: dt.date, end: dt.date) -> str:
    start_q = (start.month - 1) // 3 + 1
    end_q = (end.month - 1) // 3 + 1
    if start.year == end.year and start_q == end_q:
        return f"{start.year} Q{start_q}"
    return f"{start.isoformat()} to {end.isoformat()}"


def _to_date(value: Any) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str):
        value = value.strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y"):
            try:
                return dt.datetime.strptime(value, fmt).date()
            except ValueError:
                continue
    return None


def _workbook_path(cfg: dict) -> Path:
    path = Path(cfg.get("workbook_path") or cfg.get("_workbook_path") or cfg["workbook"])
    return path if path.is_absolute() else WS / path


def _read_workbook_issuers(cfg: dict) -> list[dict[str, Any]]:
    """Read normalized issuer fields without applying a period filter."""
    import openpyxl
    from openpyxl.utils import column_index_from_string

    wb = openpyxl.load_workbook(_workbook_path(cfg), read_only=True, data_only=True)
    try:
        ws = wb[cfg["sheet"]]
        columns = cfg["id_columns"]
        positions = {
            key: column_index_from_string(column) - 1
            for key, column in columns.items()
        }
        required = {"stock_code", "name", "prospectus_date", "listing_date"}
        missing = required - positions.keys()
        if missing:
            raise ValueError(f"id_columns is missing required fields: {', '.join(sorted(missing))}")

        companies = []
        for row_number, values in enumerate(
            ws.iter_rows(
                min_row=cfg["data_start_row"],
                max_col=max(positions.values()) + 1,
                values_only=True,
            ),
            cfg["data_start_row"],
        ):
            code = values[positions["stock_code"]]
            if code in (None, ""):
                continue
            company = {
                "row": row_number,
                "file_no": values[positions["file_no"]] if "file_no" in positions else None,
                "code": str(code).strip(),
                "name": str(values[positions["name"]] or "").strip(),
                "prospectus_date": _to_date(values[positions["prospectus_date"]]),
                "listing_date": _to_date(values[positions["listing_date"]]),
            }
            offer_price_index = positions.get("offer_price")
            company["offer_price"] = (
                values[offer_price_index] if offer_price_index is not None else None
            )
            companies.append(company)
        return companies
    finally:
        wb.close()


def _membership_date(company: dict[str, Any]) -> dt.date | None:
    return company["listing_date"] or company["prospectus_date"]


def load_cfg(
    workbook_override: str | None = None,
    period_start_override: str | None = None,
    period_end_override: str | None = None,
    *,
    config_path: str | Path | None = None,
) -> dict:
    """Resolve workbook, date interval, paths, and dataset identity for one run.

    With no date overrides, the configured period remains the default. If only
    one endpoint is supplied, the missing endpoint is inferred from the minimum
    or maximum issuer membership date in the selected workbook.
    """
    workbook_override = workbook_override or os.environ.get("HKIPO_WORKBOOK")
    period_start_override = period_start_override or os.environ.get("HKIPO_PERIOD_START")
    period_end_override = period_end_override or os.environ.get("HKIPO_PERIOD_END")
    selected_config = Path(
        config_path or os.environ.get("PIPELINE_CONFIG") or ROOT / "config.yaml"
    ).expanduser()
    if not selected_config.is_absolute():
        selected_config = (Path.cwd() / selected_config).resolve()
    if not selected_config.is_file():
        raise FileNotFoundError(f"Pipeline config not found: {selected_config}")
    with selected_config.open(encoding="utf-8") as stream:
        cfg = yaml.safe_load(stream)

    base_workbook = cfg["workbook"]
    if workbook_override:
        cfg["workbook"] = workbook_override
    cfg["_root"] = ROOT
    cfg["_ws"] = WS
    cfg["_config_path"] = selected_config.resolve()
    cfg["_workbook_path"] = _workbook_path(cfg)

    dataset = cfg.setdefault("dataset", {})
    has_period_override = bool(period_start_override or period_end_override)
    has_workbook_override = bool(
        workbook_override
        and _workbook_path({"workbook": workbook_override}).resolve()
        != _workbook_path({"workbook": base_workbook}).resolve()
    )
    if "filter_period" not in cfg:
        cfg["filter_period"] = not has_workbook_override or has_period_override

    if has_period_override:
        start = dt.date.fromisoformat(period_start_override) if period_start_override else None
        end = dt.date.fromisoformat(period_end_override) if period_end_override else None
        if start is None or end is None:
            membership_dates = [
                date
                for company in _read_workbook_issuers(cfg)
                if (date := _membership_date(company)) is not None
            ]
            if not membership_dates:
                raise ValueError(
                    "Cannot infer a missing period endpoint: the selected workbook has no issuer dates."
                )
            start = start or min(membership_dates)
            end = end or max(membership_dates)
    else:
        start = dt.date.fromisoformat(dataset["period_start"])
        end = dt.date.fromisoformat(dataset["period_end"])

    if start > end:
        raise ValueError("--period-start must be on or before --period-end")
    dataset["period_start"] = start.isoformat()
    dataset["period_end"] = end.isoformat()

    if has_period_override:
        dataset["cohort"] = _cohort_for_period(start, end)
        period_id = f"{start.isoformat()}_{end.isoformat()}"
        workbook_id = f"_{_slug(Path(workbook_override).stem)}" if has_workbook_override else ""
        dataset["id"] = f"HKIPO_{period_id}{workbook_id}"
    elif has_workbook_override:
        dataset["id"] = _slug(Path(workbook_override).stem)
        dataset["cohort"] = Path(workbook_override).stem

    # Default data keeps canonical paths; overrides get isolated state and outputs.
    if has_period_override or has_workbook_override:
        dataset_root = ROOT / "datasets" / _slug(dataset["id"])
        for key, value in cfg["paths"].items():
            relative = Path(value)
            try:
                suffix = relative.relative_to("prospectus_pipeline")
            except ValueError:
                suffix = relative
            cfg["paths"][key] = str(
                Path("prospectus_pipeline/datasets") / _slug(dataset["id"]) / suffix
            )
        cfg["state_dir"] = dataset_root / ".pipeline_state"

    for key, value in cfg["paths"].items():
        path = WS / value
        path.mkdir(parents=True, exist_ok=True)
        cfg["paths"][key] = path
    return cfg


def read_companies(cfg: dict) -> list[dict[str, Any]]:
    """Return normalized issuers in the configured inclusive date interval."""
    companies = _read_workbook_issuers(cfg)
    start = dt.date.fromisoformat(cfg["dataset"]["period_start"])
    end = dt.date.fromisoformat(cfg["dataset"]["period_end"])
    if cfg.get("filter_period", True):
        companies = [
            company
            for company in companies
            if (membership_date := _membership_date(company)) is not None
            and start <= membership_date <= end
        ]
    elif companies:
        observed_dates = [
            date
            for company in companies
            for date in (company["listing_date"], company["prospectus_date"])
            if date is not None
        ]
        if observed_dates:
            cfg["dataset"]["period_start"] = min(observed_dates).isoformat()
            cfg["dataset"]["period_end"] = max(observed_dates).isoformat()

    selected_codes = cfg.get("_selected_codes")
    if selected_codes is not None:
        selected = set(selected_codes)
        companies = [company for company in companies if company["code"] in selected]
    return companies


def cohort_workbook_path(start: dt.date, end: dt.date, *, root: Path = ROOT) -> Path:
    """The stable workbook location used by a date-only collection run."""
    period_id = f"{start.isoformat()}_{end.isoformat()}"
    return root / "datasets" / f"HKIPO_{period_id}_HKIPO-MB" / "HKIPO-MB.xlsx"


def write_cohort_config(
    start: dt.date,
    end: dt.date,
    workbook_path: Path,
    *,
    base_config: Path | None = None,
    source_dir: Path | None = None,
) -> Path:
    """Write a stable config for CLI stages and independent extraction runs."""
    import openpyxl

    base_config = Path(base_config or ROOT / "config.yaml")
    config_path = workbook_path.with_name("cohort.yaml")
    if config_path.exists():
        existing = yaml.safe_load(config_path.read_text(encoding="utf-8"))
        if existing.get("dataset", {}).get("period_start") != start.isoformat() or \
                existing.get("dataset", {}).get("period_end") != end.isoformat():
            raise ValueError(f"Existing cohort config has a different period: {config_path}")
        return config_path
    base = yaml.safe_load(base_config.read_text(encoding="utf-8"))
    dataset_id = workbook_path.parent.name
    base["workbook"] = str(workbook_path.relative_to(WS))
    dataset = base.setdefault("dataset", {})
    dataset.update({
        "id": dataset_id,
        "cohort": _cohort_for_period(start, end),
        "period_start": start.isoformat(),
        "period_end": end.isoformat(),
        "report_source_dir": str(Path(source_dir or WS / "sources").expanduser().resolve()),
    })
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        dataset["expected_companies"] = sum(
            value is not None and str(value).strip() != ""
            for (value,) in workbook["NLR"].iter_rows(min_row=2, min_col=2, max_col=2, values_only=True)
        )
    finally:
        workbook.close()
    base["filter_period"] = True
    for key, value in base["paths"].items():
        relative = Path(value)
        try:
            suffix = relative.relative_to("prospectus_pipeline")
        except ValueError:
            suffix = relative
        base["paths"][key] = str(Path("prospectus_pipeline/datasets") / dataset_id / suffix)
    base["state_dir"] = str(Path("prospectus_pipeline/datasets") / dataset_id / ".pipeline_state")
    config_path.write_text(yaml.safe_dump(base, allow_unicode=True, sort_keys=False), encoding="utf-8")
    return config_path


def build_cohort_workbook(
    start: dt.date,
    end: dt.date,
    *,
    source_dir: Path | None = None,
    template_path: Path | None = None,
    workbook_path: Path | None = None,
    today: dt.date | None = None,
) -> Path:
    """Resolve official issuers and create a styled Main Board workbook.

    Repeated calls preserve an existing workbook and all of its collected data.
    If the official issuer membership changes, stop for reconciliation instead
    of overwriting completed cells or silently omitting new issuers.
    """
    import openpyxl

    if start > end:
        raise ValueError("period start must be on or before period end")
    source_dir = Path(source_dir or WS / "sources")
    template_path = Path(template_path or WS / "templates" / "HKIPO-MB-template-final.xlsx")
    workbook_path = Path(workbook_path or cohort_workbook_path(start, end))
    reports = reports_for_interval(start, end, source_dir, today=today)

    selected = []
    for report in reports:
        match = re.search(r"(?:NLR)?(19\d{2}|20\d{2})", report.stem, re.I)
        if not match:
            raise ValueError(f"Cannot identify annual report year: {report.name}")
        for candidate in load_nlr_candidates(report, int(match.group(1))):
            candidate = audit_candidate(candidate)
            membership = candidate.get("listing_date") or candidate.get("prospectus_date")
            if candidate["inclusion_status"] == "INCLUDED" and membership and start <= membership <= end:
                selected.append(candidate)
    selected.sort(key=lambda item: (
        item.get("listing_date") or item.get("prospectus_date"), item["stock_code"]
    ))
    if not selected:
        raise ValueError(f"No ordinary Main Board IPOs found from {start} through {end}")

    codes = [item["stock_code"] for item in selected]
    duplicates = sorted(code for code, count in Counter(codes).items() if count > 1)
    if duplicates:
        raise ValueError("Reused stock codes in requested interval require separate cohorts: " + ", ".join(duplicates))

    if workbook_path.exists():
        current = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
        try:
            rows = current["NLR"].iter_rows(min_row=2, min_col=2, max_col=2, values_only=True)
            existing = [str(row[0]).strip() for row in rows if row[0] not in (None, "")]
        finally:
            current.close()
        if existing != codes:
            raise ValueError(
                f"Existing cohort workbook differs from current official issuer list: {workbook_path}. "
                "Reconcile it manually; collected cells were not overwritten."
            )
        return workbook_path

    if not template_path.is_file():
        raise FileNotFoundError(f"Main Board workbook template not found: {template_path}")
    workbook = openpyxl.load_workbook(template_path)
    try:
        sheet = workbook["NLR"]
        if any(sheet.cell(row, 2).value not in (None, "") for row in range(2, sheet.max_row + 1)):
            raise ValueError(f"Main Board template is not blank: {template_path}")
        template_last_row = sheet.max_row
        for row_number, item in enumerate(selected, 2):
            if row_number > template_last_row:
                for col in range(1, sheet.max_column + 1):
                    source = sheet.cell(2, col)
                    target = sheet.cell(row_number, col)
                    if source.has_style:
                        target._style = copy(source._style)
            values = (
                item.get("file_no"), item["stock_code"], item["company_name"],
                item.get("prospectus_date"), item.get("listing_date"),
                item.get("sponsors"), item.get("auditor"), item.get("valuer"),
                item.get("funds_raised_public_hkd"), item.get("funds_raised_int_hkd"),
                item.get("offer_price_hkd"),
            )
            for col, value in enumerate(values, 1):
                sheet.cell(row_number, col).value = value
        workbook_path.parent.mkdir(parents=True, exist_ok=True)
        temporary = workbook_path.with_name("." + workbook_path.name)
        try:
            workbook.save(temporary)
            os.replace(temporary, workbook_path)
        finally:
            if temporary.exists():
                temporary.unlink()
    finally:
        workbook.close()
    return workbook_path
