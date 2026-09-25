"""严格验证 AI 抽取结果；任何结构/类型/证据/勾稽错误都阻止写回。"""
from __future__ import annotations

import json
import math
import re
from pathlib import Path
from typing import Any

from contracts import (
    MISSING_NUMERIC,
    MISSING_TEXT,
    evidence_issues,
    is_missing,
    normalize_code,
    strict_load_file,
    validate_record,
)
from storage import atomic_json, official_files

MISSING = {"", None, "nan", "na", "n/a", "none", "-"}


def apply_da(cfg: dict, only=None, log=print) -> int:
    """公告未直接给 free float 时，按已约定公式补 col_DA / col_DC。

    DA = (CS − CK×CS) / CZ；DC = CZ（分母）。仅填缺失项，不覆盖已有数字。
    """
    ext = cfg["paths"]["allot_out"] / "extracted"
    n = 0
    for _, fp in official_files(ext, only=only).items():
        rec = json.loads(fp.read_text(encoding="utf-8"))
        f = rec.setdefault("fields", {})

        def raw(k):
            e = f.get(k) or {}
            return e.get("value") if isinstance(e, dict) else None

        def as_num(k):
            v = raw(k)
            if v in (None, "", "NaN", "NA", "nan", "na"):
                return None
            try:
                return float(v)
            except (TypeError, ValueError):
                return None

        cs, ck, cz = as_num("col_CS"), as_num("col_CK"), as_num("col_CZ")
        if None in (cs, ck, cz) or cz == 0:
            continue
        da_missing = as_num("col_DA") is None or (f.get("col_DA") or {}).get("source") == "da_formula"
        dc_missing = as_num("col_DC") is None or (f.get("col_DC") or {}).get("source") == "da_formula"
        expected = cs * (1 - ck) / cz
        quote = (f"公式：(col_CS {cs:,.0f} − col_CK {ck:.4f}×col_CS) ÷ col_CZ {cz:,.0f}"
                 f" = {expected:.4f}；DC = CZ")
        if da_missing:
            f["col_DA"] = {"value": round(expected, 4), "page": None, "quote": quote,
                           "source": "da_formula", "confidence": "high"}
        if dc_missing:
            f["col_DC"] = {"value": int(round(cz)), "page": (f.get("col_CZ") or {}).get("page"),
                           "quote": quote, "source": "da_formula", "confidence": "high"}
        db = raw("col_DB")
        if db in (None, "", "NA", "NaN"):
            f["col_DB"] = {"value": "Free float denominator = total issued shares upon listing (col_CZ); "
                                    "numerator = offer shares minus cornerstone allocation",
                           "page": None, "quote": quote, "source": "da_formula",
                           "confidence": "high"}
        atomic_json(fp, rec)
        n += 1
    log(f"col_DA/DC 公式回填：{n} 家")
    return n


def num(v):
    if isinstance(v, dict):
        v = v.get("value")
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)) and math.isfinite(float(v)):
        return float(v)
    if isinstance(v, str) and v in {MISSING_NUMERIC, MISSING_TEXT}:
        return None
    return None


def close(a, b, rel_tol=1e-9, abs_tol=1.0) -> bool:
    if a is None or b is None:
        return True
    return abs(a - b) <= max(abs_tol, rel_tol * max(abs(a), abs(b)))


def check_firm(rec: dict, schema: dict, share_rel_tol=1e-9, share_abs_tol=1.0,
               balance_rel_tol=0.005, balance_abs_tol=2000.0) -> list[dict]:
    """只做可证明的勾稽；缺失字段由契约报告为 missing，不静默当作 OK。"""
    f = rec.get("fields", {})
    def g(key):
        return num(f.get(key))
    issues: list[dict] = []
    code = rec.get("code", "?")

    def add(check, detail):
        issues.append({"code": code, "check": check, "detail": detail, "severity": "error"})

    M, R, S = g("col_M"), g("col_R"), g("col_S")
    Q, P, L, N, O = g("col_Q"), g("col_P"), g("col_L"), g("col_N"), g("col_O")
    T, U = g("col_T"), g("col_U")
    if M is not None and R is not None and S is not None and not close(M, R + S, share_rel_tol, share_abs_tol):
        add("全球发售=配售+公开发售", f"M={M:,.0f} vs R+S={R+S:,.0f}")
    if M is not None and Q is not None and P is not None and not close(M, Q + P, share_rel_tol, share_abs_tol):
        add("全球发售=新股+老股", f"M={M:,.0f} vs Q+P={Q+P:,.0f}")
    if L is not None and N is not None and Q is not None and not close(L, N + Q, share_rel_tol, share_abs_tol):
        add("总股本=资本化发行旧股+新股", f"L={L:,.0f} vs N+Q={N+Q:,.0f}")
    if L is not None and O is not None and M is not None and not close(L, O + M, share_rel_tol, share_abs_tol):
        add("总股本=资本化后旧股+全球发售", f"L={L:,.0f} vs O+M={O+M:,.0f}")

    for yr, (a, e, li) in {"year-1": ("col_Y", "col_AB", "col_AE"),
                           "year-2": ("col_X", "col_AA", "col_AD"),
                           "year-3": ("col_W", "col_Z", "col_AC")}.items():
        A, E, LI = g(a), g(e), g(li)
        if A is not None and E is not None and LI is not None and not close(A, E + LI, balance_rel_tol, balance_abs_tol):
            add(f"资产=权益+负债 {yr}", f"资产={A:,.0f} vs 权益+负债={E+LI:,.0f}")
    if T is not None and U is not None and U > T:
        add("价格区间", f"最低价 {U} > 最高价 {T}")

    for key in ("col_AO", "col_AP", "col_AQ", "col_AY"):
        v = g(key)
        if v is not None and not 0 <= v <= 1:
            add("百分比范围", f"{key}={v} 不在 [0,1]")

    # ---------------------------------------------------------
    # 毛利率天花板与未年化期间勾稽 (Gross Margin Ceiling Gate)
    # ---------------------------------------------------------
    CF = g("col_CF")
    AH = g("col_AH")
    AT_val = (f.get("col_AT") or {}).get("value")
    if CF is not None and AH is not None and AH > 0 and AT_val:
        at_str = str(AT_val).strip()
        frac = 1.0
        if any(x in at_str for x in ("30/06", "06/30", "-06-30")):
            frac = 0.5
        elif any(x in at_str for x in ("31/08", "08/31", "-08-31")):
            frac = 8.0 / 12.0
        elif any(x in at_str for x in ("30/09", "09/30", "-09-30")):
            frac = 0.75
        elif any(x in at_str for x in ("31/10", "10/31", "-10-31")):
            frac = 10.0 / 12.0
        unann_sales = AH * frac
        if CF > unann_sales * 1.01:
            add("毛利率超限/期间错位",
                f"year-1 未年化毛利 col_CF={CF:,.0f} 超过未年化销售额 {unann_sales:,.0f} "
                f"(毛利率达 {CF/unann_sales*100:.1f}%)，疑似错采了全年毛利！")
        elif unann_sales > 10_000_000 and CF > 0 and (CF / unann_sales) < 0.005:
            add("毛利数值数量级异常",
                f"col_CF={CF:,.0f} 相比未年化销售额 {unann_sales:,.0f} 仅占 {CF/unann_sales*100:.2f}%，"
                "疑似漏乘表格千元/百万元乘数！")

    # ---------------------------------------------------------
    # 货币基本单位数量级防呆 (Monetary Scale Guard)
    # ---------------------------------------------------------
    curr = str((f.get("col_V") or {}).get("value") or "").upper()
    total_assets = g("col_Y")
    annual_sales = g("col_AH")
    is_large_firm = (total_assets is not None and total_assets > 50_000_000) or (annual_sales is not None and annual_sales > 50_000_000)
    if curr in {"RMB", "HKD"} and is_large_firm:
        for mkey, mname in [("col_CF", "毛利"), ("col_BE", "有息负债")]:
            mval = g(mkey)
            if mval is not None and 0 < mval < 500_000:
                add(f"金额单位未折算 ({mname})",
                    f"{mkey}={mval} 处于 (0, 500,000)，明显未按基本货币单位折算（漏乘千元乘数）")
        cg_val = g("col_CG")
        if cg_val is not None and 0 < cg_val < 100_000:
            add("金额单位未折算 (资本开支)",
                f"col_CG={cg_val} 处于 (0, 100,000)，明显未按基本货币单位折算（漏乘千元乘数）")

    # ---------------------------------------------------------
    # 资本开支合理性勾稽 (CapEx Bounds)
    # ---------------------------------------------------------
    CG = g("col_CG")
    if CG is not None and CG < 0:
        add("资本开支为负数", f"col_CG={CG:,.0f} 不能为负数")
    if CG is not None and total_assets is not None and total_assets > 0 and CG > total_assets * 1.5:
        add("资本开支超出总资产", f"col_CG={CG:,.0f} 超过总资产 col_Y={total_assets:,.0f} 的 1.5 倍")

    return issues


def _sheet_rows(cfg: dict) -> list[tuple]:
    """一次性读入工作表（read-only 模式下逐格 .cell() 是 O(n²)，极慢）。"""
    import openpyxl
    wb = openpyxl.load_workbook(cfg["_ws"] / cfg["workbook"], read_only=True, data_only=True)
    ws = wb[cfg["sheet"]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return rows


def _expected_codes(cfg: dict) -> list[str]:
    from openpyxl.utils import column_index_from_string
    idx = column_index_from_string(cfg["id_columns"]["stock_code"]) - 1
    codes = []
    for row in _sheet_rows(cfg)[cfg["data_start_row"] - 1:]:
        v = row[idx] if idx < len(row) else None
        if v not in (None, ""):
            codes.append(normalize_code(v))
    return codes


def check_allot(rec: dict, ctx: dict) -> list[dict]:
    """配发结果公告字段的交叉校验。ctx 提供招股书初始发售规模与公告日期。"""
    f = rec.get("fields", {})
    code = rec.get("code", "?")
    issues: list[dict] = []

    def g(k):
        return num(f.get(k))

    def add(check, detail):
        issues.append({"code": code, "check": check, "detail": detail})

    CS, CT, CU = g("col_CS"), g("col_CT"), g("col_CU")
    CV, CK, CM = g("col_CV"), g("col_CK"), g("col_CM")
    CY, DA, CN, CO, CX = g("col_CY"), g("col_DA"), g("col_CN"), g("col_CO"), g("col_CX")

    if CS is not None and CT is not None and CU is not None and not close(CT + CU, CS, 1e-9, 1.0):
        add("最终公开+配售=最终全球发售", f"CT+CU={CT+CU:,.0f} vs CS={CS:,.0f}")
    base = ctx.get("initial_offer")
    if base is not None and CS is not None and CS + 1 < base:
        add("最终全球发售≥初始发售", f"CS={CS:,.0f} < 初始 M={base:,.0f}（规模调整权只能扩大）")
    if CV is not None and CS is not None and CV > 0.25 * CS:
        add("超额配售比例异常", f"CV={CV:,.0f} 超过 CS 的 25%")
    if CY is not None and not 0 <= CY <= 1:
        add("百分比范围", f"CY={CY} 不在 [0,1]")
    if CK is not None and not 0 <= CK <= 1:
        add("百分比范围", f"CK={CK} 不在 [0,1]")
    if DA is not None and CY is not None and DA > CY + 1e-9:
        add("不受限公众持股≤公众持股", f"DA={DA} > CY={CY}")
    if CM is not None and CM < 0:
        add("认购倍数", f"CM={CM} 为负")
    if CN is not None and CN < 0:
        add("申请人数", f"CN={CN} 为负")
    if CO is not None and CN is not None and CO > 0 and CN == 0:
        add("申请人数/股数矛盾", f"CO={CO:,.0f} 但 CN=0")
    if CX is not None and CS is not None and ctx.get("price"):
        gross = CS * ctx["price"]
        if not (0 < CX < gross):
            add("净募资额应在(0, 毛募资)之间", f"CX={CX:,.0f} vs 毛募资≈{gross:,.0f}")
    # 公告日期应与港交所元数据一致（确定性核对）
    exp = ctx.get("announce_date")
    got = (f.get("col_CR") or {}).get("value")
    if exp and got not in (None, "NA"):
        if str(got).strip()[:10] != str(exp).strip()[:10]:
            add("公告日期与元数据不符", f"JSON={got!r} vs 索引={exp!r}")
    # 自由流通比例应与「股数 ÷ 分母」自洽（分母 = CZ 或 DC）
    DA, DC, CZ = g("col_DA"), g("col_DC"), g("col_CZ")
    if DA is not None and CS is not None and CK is not None:
        # 自由流通百分比 ≈ (发售股份 − 基石股份) ÷ 上市时已发行股份
        denom = CZ if CZ else None
        if denom:
            expected = CS * (1 - CK) / denom
            if abs(DA - expected) > 0.20 * max(expected, 1e-9):
                add("自由流通比例口径可疑",
                    f"DA={DA:.4f} 与 (CS×(1−CK))/CZ={expected:.4f} 相差过大；"
                    f"注意不要用 CY−CK 的算法（会漏掉其他禁售股东）")

    # --- 超额配售（绿鞋）：只认上市后港交所的行使/失效公告，不接受从配发公告反推 ---
    gre = ctx.get("greenshoe")
    if not gre or gre.get("status") != "ok":
        add("绿鞋未核实", "缺 greenshoe.json 记录：不能从配发公告推断 CV，须先跑 greenshoe 阶段")
    elif gre.get("exercised"):
        s = ctx.get("greenshoe_shares") or {}
        if s.get("status") != "ok" or s.get("shares") is None:
            add("绿鞋已行使但缺股数", "行使公告存在但未取到实际配发股数")
        elif CV is None:
            add("绿鞋已行使但 CV 缺失", f"公告实际配发 {s['shares']:,} 股")
        elif int(CV) != int(s["shares"]):
            add("绿鞋股数与行使公告不符",
                f"CV={CV:,.0f} vs 行使公告 {s['shares']:,}（{s.get('doc_title')}）")
    elif gre.get("window_closed"):
        if CV is None:
            add("绿鞋未行使须填 0", f"上市日起 {gre['window']} 内无行使公告，CV 应为 0（确定的零）")
        elif CV != 0:
            add("绿鞋未行使但 CV 非 0", f"CV={CV:,.0f}，但检索区间内无行使公告")
    else:
        if CV not in (None, 0):
            add("绿鞋窗口未结束", f"CV={CV:,.0f}，但窗口 {gre['window']} 尚未结束，应为 NaN")

    # --- 基石：确认无基石时必须填 0（手册：确定的零写 0，不写 NaN）---
    if CK is None:
        ca = ctx.get("cornerstone")
        if not ca:
            add("基石未核实", "缺 cornerstone_absence.json：须先跑 cornerstone 阶段")
        elif ca.get("verdict") == "none":
            add("基石获配缺失",
                f"招股书 {ca['prospectus_pages']} 页 0 处「基石投资者」-> 确认无基石，须填 0")
        else:
            add("基石获配缺失", "招股书有基石投资者披露，col_CK 不得为缺失")


    # 全空报警
    if all(g(k) is None for k in ("col_CS", "col_CT", "col_CU", "col_CM", "col_CK")):
        add("核心字段全缺", "CS/CT/CU/CM/CK 全为缺失，疑似未读到公告正文")
    return issues


def _allot_context(cfg: dict) -> dict:
    """给 allot 校验准备上下文：初始发售规模、发售价、公告日期。"""
    from openpyxl.utils import column_index_from_string
    ctx: dict = {}
    rows = _sheet_rows(cfg)
    if not rows:
        return ctx
    header = rows[0]
    hdr = {}
    for i, v in enumerate(header):
        if v not in (None, ""):
            hdr[" ".join(str(v).replace("\n", " ").split()).strip().lower()] = i
    m_i = hdr.get("global offering (without option)")
    k_i = hdr.get("ipo subscription price (hk$)")
    code_i = column_index_from_string(cfg["id_columns"]["stock_code"]) - 1
    for row in rows[cfg["data_start_row"] - 1:]:
        code = row[code_i] if code_i < len(row) else None
        if code in (None, ""):
            continue
        key = normalize_code(code)
        ctx.setdefault(key, {})
        if m_i is not None and m_i < len(row):
            ctx[key]["initial_offer"] = row[m_i]
        if k_i is not None and k_i < len(row):
            ctx[key]["price"] = row[k_i]
    idx = cfg["paths"]["out"] / "allotment_index.json"
    if idx.exists():
        for rec in json.loads(idx.read_text()):
            ctx.setdefault(normalize_code(rec["code"]), {})["announce_date"] = rec.get("datetime")
    # 绿鞋检索结果 + 行使股数（col_CV 的确定性来源）+ 基石核实（col_CK 的零值依据）
    for name, key in (("greenshoe.json", "greenshoe"),
                      ("greenshoe_shares.json", "greenshoe_shares"),
                      ("cornerstone_absence.json", "cornerstone")):
        p = cfg["paths"]["allot_out"] / name
        if not p.exists():
            continue
        for code, rec in json.loads(p.read_text(encoding="utf-8")).items():
            ctx.setdefault(normalize_code(code), {})[key] = rec
    return ctx


def _greenshoe_haystack(cfg: dict, code: str) -> str | None:
    """把行使公告的页级文本拼成带 <<<PAGE n>>> 标记的核对用全文。"""
    f = (cfg["_root"] / "data" / "allot" / "greenshoe" / "text"
         / f"HKIPO-MB{''.join(ch for ch in code if ch.isdigit())}.jsonl")
    if not f.exists():
        return None
    return "\n".join(f"<<<PAGE {json.loads(x)['page']}>>>\n{json.loads(x)['text']}"
                     for x in f.open(encoding="utf-8"))


def validate_all(cfg: dict, only: list[str] | None = None, limit: int = 0,
                 target: str = "prospectus", log=print) -> dict:
    if target == "allot":
        schema = strict_load_file(cfg["_root"] / "schema" / "allot_fields.json")
        ext_dir: Path = cfg["paths"]["allot_out"] / "extracted"
        packets_dir: Path = cfg["paths"]["allot_packets"]
        ctx_map = _allot_context(cfg)
        out_file = cfg["paths"]["allot_out"] / "validation.json"
    else:
        schema = strict_load_file(cfg["_root"] / "schema" / "fields.json")
        ext_dir = cfg["paths"]["out"] / "extracted"
        packets_dir = cfg["paths"]["packets"]
        ctx_map = {}
        out_file = cfg["paths"]["out"] / "validation.json"
    wanted = [normalize_code(x) for x in only] if only else _expected_codes(cfg)
    if limit:
        wanted = wanted[:limit]
    by_code = {normalize_code(x): x for x in wanted}
    files = official_files(ext_dir) if ext_dir.exists() else {}
    records, errors, missing_files, pending = [], [], [], []
    for code in wanted:
        fp = files.get(code)
        if fp is None:
            missing_files.append(code)
            continue
        try:
            rec = strict_load_file(fp)
            context = ctx_map.get(code, {}) if target == "allot" else {}
            if target == "allot":
                from cornerstone import assess
                context = dict(context)
                if "cornerstone" not in context:
                    context["cornerstone"] = assess(cfg, code)
            structural = validate_record(rec, schema, code, target, context)
            packet = packets_dir / f"HKIPO-MB{code.split('.')[0]}.md"
            alt: dict[str, str] = {}
            if target == "allot":
                gh = _greenshoe_haystack(cfg, code)
                if gh:
                    alt["greenshoe"] = gh
            # 招股书全文（页级 jsonl）作为合法证据源（招股书与配发阶段均可引用）
            text_file = cfg["paths"]["text"] / f"HKIPO-MB{''.join(ch for ch in code if ch.isdigit())}.jsonl"
            if text_file.exists():
                parts = []
                for line in text_file.open(encoding="utf-8"):
                    p = json.loads(line)
                    parts.append(f"<<<PAGE {p['page']}>>>\n{p['text']}")
                alt["prospectus"] = "\n".join(parts)
            ev = evidence_issues(rec, packet, schema, alt, target, context) if packet.exists() else ["packet missing"]
            if target == "allot":
                checks = check_allot(rec, context)
                checks = [{"check": c["check"], "detail": c["detail"]} for c in checks]
            else:
                checks = check_firm(rec, schema, cfg["validate"].get("share_rel_tol", 1e-9),
                                    cfg["validate"].get("share_abs_tol", 1),
                                    cfg["validate"].get("balance_rel_tol", 0.005),
                                    cfg["validate"].get("balance_abs_tol", 2000))
            all_err = structural + ev + [x["check"] + ": " + x["detail"] for x in checks]
            fields = rec.get("fields", {}) if isinstance(rec, dict) else {}
            missing = []
            for field in schema["fields"]:
                value = fields.get(field["key"], {}).get("value") if isinstance(fields.get(field["key"]), dict) else None
                if (not isinstance(fields.get(field["key"]), dict)
                        or is_missing(value, field.get("kind", "text"))):
                    missing.append(field["key"])
            status = "error" if all_err else ("warning_missing" if missing else "pass")
            from state import digest, file_hash, save_record
            evidence = {"packet": packet.read_text(encoding="utf-8"), "alt": alt, "context": context}
            version = digest(rec, schema, evidence)
            records.append({"code": code, "file": str(fp), "hash": version,
                            "json_sha256": file_hash(fp), "status": status,
                            "fields_total": len(schema["fields"]),
                            "fields_missing": missing, "errors": all_err})
            save_record(cfg, target, code, "validated", {**records[-1], "target": target,
                                                         "gate_pass": not all_err})
            if all_err:
                errors.extend({"code": code, "detail": x} for x in all_err)
            if missing:
                pending.append({"code": code, "fields": missing})
            log(f"{status.upper():15s} {code:9s} missing={len(missing):2d} errors={len(all_err):2d}")
        except Exception as exc:
            errors.append({"code": code, "detail": f"cannot parse/validate {fp}: {exc}"})
            records.append({"code": code, "file": str(fp), "status": "error", "errors": [str(exc)]})
    if missing_files:
        errors.extend({"code": c, "detail": "missing extraction JSON"} for c in missing_files)
        log("MISSING FILES: " + ", ".join(missing_files))
    run_report = {
        "schema_version": 1,
        "target": target,
        "expected_codes": wanted,
        "records": records,
        "missing_files": missing_files,
        "pending_missing_fields": pending,
        "errors": errors,
        "gate_pass": not errors,
    }
    # Save immutable run-level report
    import datetime as dt
    import uuid
    runs_dir = (cfg["paths"]["allot_out"] if target == "allot" else cfg["paths"]["out"]) / "validation_runs"
    runs_dir.mkdir(parents=True, exist_ok=True)
    ts_str = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%S_%fZ")
    run_id = uuid.uuid4().hex[:10]
    run_report["run_id"] = run_id
    run_report["generated_at"] = dt.datetime.now(dt.timezone.utc).isoformat()
    atomic_json(runs_dir / f"validation_{target}_{ts_str}_{run_id}.json", run_report)

    # Rebuild or update aggregate validation.json
    if only and out_file.exists():
        try:
            agg_report = strict_load_file(out_file)
            agg_records = {r["code"]: r for r in agg_report.get("records", [])}
            for r in records:
                agg_records[r["code"]] = r
            all_expected = _expected_codes(cfg)
            all_records = [agg_records[c] for c in all_expected if c in agg_records]
            all_missing_files = [c for c in all_expected if c not in agg_records]
            all_errors = [e for r in all_records for e in [{"code": r["code"], "detail": d} for d in r.get("errors", [])]]
            all_pending = [{"code": r["code"], "fields": r.get("fields_missing", [])} for r in all_records if r.get("fields_missing")]
            report = {
                "schema_version": 1,
                "target": target,
                "expected_codes": all_expected,
                "records": all_records,
                "missing_files": all_missing_files,
                "pending_missing_fields": all_pending,
                "errors": all_errors,
                "gate_pass": not all_errors and not all_missing_files,
            }
        except Exception:
            report = run_report
    else:
        report = run_report

    out = out_file
    atomic_json(out, report)
    log(f"\n校验完成[{target}]：目标 {len(wanted)} 家，解析 {len(records)} 家，错误 {len(errors)} 项；"
        f"写回闸门={'PASS' if not errors else 'BLOCK'} -> {out}")
    return report
