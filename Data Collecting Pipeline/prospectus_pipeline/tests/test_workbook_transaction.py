import tempfile
import unittest
from pathlib import Path

import openpyxl

from workbook_transaction import (
    WorkbookLockError,
    commit_prepared_workbook,
    workbook_transaction,
)
from storage import file_sha256


class WorkbookTransactionTests(unittest.TestCase):
    def make_book(self, folder: str) -> Path:
        path = Path(folder) / "sample.xlsx"
        wb = openpyxl.Workbook()
        wb.active["A1"] = "original"
        wb.save(path)
        wb.close()
        return path

    def test_transaction_commits_and_snapshots(self):
        with tempfile.TemporaryDirectory() as tmp:
            book = self.make_book(tmp)
            with workbook_transaction(book, operation="test") as wb:
                wb.active["A1"] = "updated"
            check = openpyxl.load_workbook(book, read_only=True)
            self.assertEqual(check.active["A1"].value, "updated")
            check.close()
            snapshots = list((Path(tmp) / "backups" / "excel_snapshots").glob("*.xlsx"))
            self.assertEqual(len(snapshots), 1)

    def test_prepared_commit_rejects_concurrent_change(self):
        with tempfile.TemporaryDirectory() as tmp:
            book = self.make_book(tmp)
            original_hash = file_sha256(book)
            prepared = openpyxl.load_workbook(book)
            prepared.active["A1"] = "prepared"
            with workbook_transaction(book, operation="other") as other:
                other.active["A1"] = "concurrent"
            with self.assertRaises(WorkbookLockError):
                commit_prepared_workbook(book, prepared, original_hash, operation="prepared")


if __name__ == "__main__":
    unittest.main()
