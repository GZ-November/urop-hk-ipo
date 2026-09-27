import datetime as dt
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / "src")]

from cohort import check_cross_cohort_conflicts, check_listing_dates_within_period


def write_cohort_workbook(path: Path, rows: list[tuple[str, dt.date | None]]) -> None:
    """rows: (stock code, listing date)；列布局与 NLR 一致（B=代码，E=上市日）。"""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "NLR"
    sheet["B1"] = "Stock Code"
    sheet["E1"] = "Date of Listing (dd/mm/yy)"
    for i, (code, listing) in enumerate(rows, start=2):
        sheet.cell(row=i, column=2, value=code)
        if listing is not None:
            sheet.cell(row=i, column=5, value=listing)
    workbook.save(path)


class CrossCohortGuardTests(unittest.TestCase):
    def test_conflict_detected_against_sibling_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            write_cohort_workbook(tmp / "HKIPO-MB2025Q1.xlsx", [("0001.HK", None)])
            with self.assertRaises(ValueError) as ctx:
                check_cross_cohort_conflicts(
                    tmp / "HKIPO-MB2025Q2.xlsx", ["0001.HK", "0002.HK"])
            self.assertIn("0001.HK", str(ctx.exception))
            self.assertIn("HKIPO-MB2025Q1.xlsx", str(ctx.exception))

    def test_self_workbook_is_not_a_conflict(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            target = tmp / "HKIPO-MB2025Q2.xlsx"
            write_cohort_workbook(target, [("0001.HK", None)])
            write_cohort_workbook(tmp / "HKIPO-MB2025Q1.xlsx", [("0009.HK", None)])
            check_cross_cohort_conflicts(target, ["0001.HK"])

    def test_only_cohort_named_workbooks_are_scanned(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            write_cohort_workbook(tmp / "HKIPO-MB-template-final.xlsx", [("0001.HK", None)])
            check_cross_cohort_conflicts(tmp / "HKIPO-MB2025Q2.xlsx", ["0001.HK"])


class ListingDateGuardTests(unittest.TestCase):
    def test_out_of_period_listing_date_raises(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            path = tmp / "HKIPO-MB2025Q2.xlsx"
            write_cohort_workbook(path, [
                ("0001.HK", dt.datetime(2025, 5, 6)),
                ("0002.HK", dt.datetime(2025, 7, 2)),
                ("0003.HK", None),
            ])
            with self.assertRaises(ValueError) as ctx:
                check_listing_dates_within_period(
                    path, dt.date(2025, 4, 1), dt.date(2025, 6, 30))
            self.assertIn("0002.HK", str(ctx.exception))
            self.assertNotIn("0003.HK", str(ctx.exception))

    def test_unknown_listing_dates_are_tolerated(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            path = tmp / "HKIPO-MB2025Q2.xlsx"
            write_cohort_workbook(path, [("0001.HK", None), ("0002.HK", None)])
            check_listing_dates_within_period(
                path, dt.date(2025, 4, 1), dt.date(2025, 6, 30))


if __name__ == "__main__":
    unittest.main()
