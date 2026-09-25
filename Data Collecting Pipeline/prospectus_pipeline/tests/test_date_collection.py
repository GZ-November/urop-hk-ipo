"""Date-only cohort creation from official annual listing reports."""
from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "src"))

from cohort import build_cohort_workbook, load_cfg, read_companies, write_cohort_config  # noqa: E402
from listing_reports import report_links, reports_for_interval  # noqa: E402
from sample_builder import clean_code  # noqa: E402
import run as pipeline_run  # noqa: E402

TEMPLATE = ROOT.parent / "templates" / "HKIPO-MB-template-final.xlsx"
TODAY = dt.date(2026, 9, 25)


def make_report(path: Path, year: int, as_of: str, rows: list[list]) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([f"Main Board New Listings up to {as_of}"])
    sheet.append(["File no", "Stock Code", "Company Name", "Date of Prospectus",
                  "Date of Listing", "Sponsor(s)", "Accountant", "Valuer",
                  "Funds Raised HK (a)", "Offer Price", "Fund Type"])
    for row in rows:
        sheet.append(row)
    workbook.save(path / f"NLR{year}_Eng.xlsx")
    workbook.close()


class SourceFixture(unittest.TestCase):
    def setUp(self):
        self.temporary_source = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary_source.cleanup)
        self.source = Path(self.temporary_source.name)
        make_report(self.source, 2025, "30 December 2025", [
            [1, 9998, "Prior Year IPO Ltd", dt.datetime(2025, 12, 1),
             dt.datetime(2025, 12, 8), "Sponsor", "Accountant", None, 100, 10],
        ])
        make_report(self.source, 2026, "30 June 2026", [
            [1, 6082, "First IPO Ltd", dt.datetime(2025, 12, 22),
             dt.datetime(2026, 1, 2), "Sponsor", "Accountant", None, 100, 10],
            [2, 2513, "Second IPO Ltd", dt.datetime(2026, 2, 1),
             dt.datetime(2026, 2, 9), "Sponsor", "Accountant", None, 200, 20],
            [3, 8888, "Introduction Ltd", dt.datetime(2026, 2, 5),
             dt.datetime(2026, 2, 10), "Sponsor", "Accountant", None, 0, None],
            [4, 7777, "Later IPO Ltd", dt.datetime(2026, 4, 20),
             dt.datetime(2026, 5, 1), "Sponsor", "Accountant", None, 300, 30],
        ])


class OfficialReportTests(SourceFixture):
    def test_numeric_stock_code_from_xls_does_not_gain_a_zero(self):
        self.assertEqual(clean_code(9998.0), "9998.HK")

    def test_report_link_discovery_accepts_xlsx_and_historical_xls(self):
        html = """
        <a href="/-/media/HKEXnews/Homepage/New-Listings/New-Listing-Information/New-Listing-Report/Main/NLR2026_Eng.xlsx">Download</a>
        <a href="/-/media/HKEXnews/Homepage/New-Listings/New-Listing-Information/New-Listing-Report/Main/2010.XLS">Download</a>
        <a href="https://untrusted.example/NLR2025_Eng.xlsx">Download</a>
        """
        links = report_links(html)
        self.assertEqual(set(links), {2010, 2026})
        self.assertTrue(links[2010].endswith("2010.XLS"))

    def test_cached_report_covers_q1_without_network(self):
        session = Mock()
        paths = reports_for_interval(
            dt.date(2026, 1, 1), dt.date(2026, 3, 31), self.source,
            session=session, today=TODAY,
        )
        self.assertEqual([p.name for p in paths], ["NLR2026_Eng.xlsx"])
        session.get.assert_not_called()

    def test_missing_report_downloads_from_hkex_and_reuses_cache(self):
        url = ("https://www2.hkexnews.hk/-/media/HKEXnews/Homepage/New-Listings/"
               "New-Listing-Information/New-Listing-Report/Main/NLR2026_Eng.xlsx")
        page = Mock(text=f'<a href="{url}">2026</a>')
        report = Mock(content=(self.source / "NLR2026_Eng.xlsx").read_bytes())
        session = Mock()
        session.get.side_effect = [page, report]
        with tempfile.TemporaryDirectory() as temporary:
            cache = Path(temporary)
            first = reports_for_interval(
                dt.date(2026, 1, 1), dt.date(2026, 3, 31), cache,
                session=session, today=TODAY,
            )
            second = reports_for_interval(
                dt.date(2026, 1, 1), dt.date(2026, 3, 31), cache,
                session=session, today=TODAY,
            )
            self.assertEqual(first, second)
            self.assertEqual(first[0].read_bytes(), report.content)
            self.assertEqual(session.get.call_count, 2)

    def test_archived_report_can_end_on_final_listing_day(self):
        session = Mock()
        paths = reports_for_interval(
            dt.date(2025, 12, 31), dt.date(2025, 12, 31), self.source,
            session=session, today=TODAY,
        )
        self.assertEqual([path.name for path in paths], ["NLR2025_Eng.xlsx"])
        session.get.assert_not_called()

    def test_stale_current_report_fails_closed(self):
        session = Mock()
        session.get.side_effect = OSError("network unavailable")
        with self.assertRaisesRegex(ValueError, "2026"):
            reports_for_interval(
                dt.date(2026, 9, 1), dt.date(2026, 9, 20), self.source,
                session=session, today=TODAY,
            )

    def test_ipo_count_window_includes_prior_year(self):
        cfg = {"dataset": {"report_source_dir": str(self.source)}}
        issuers = [{"prospectus_date": dt.date(2026, 1, 10)}]
        with patch("listing_reports.reports_for_interval", return_value=[]) as reports:
            pipeline_run._ipo_count_reports(cfg, Mock(source_dir=None), issuers)
        reports.assert_called_once_with(
            dt.date(2025, 10, 12), dt.date(2026, 1, 9), self.source.resolve()
        )

    def test_external_stops_before_writes_when_prior_year_is_missing(self):
        with patch.object(pipeline_run, "_ipo_count_reports", side_effect=ValueError("2025 missing")), \
                patch("subprocess.run") as subprocess_run:
            result = pipeline_run.cmd_external(
                {"dataset": {}}, argparse.Namespace(),
                [{"prospectus_date": dt.date(2026, 1, 10)}],
            )
        self.assertEqual(result, 1)
        subprocess_run.assert_not_called()


class DateCohortWorkbookTests(SourceFixture):
    def test_date_only_run_selects_ordinary_q1_issuers_and_preserves_resume(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "cohort.xlsx"
            result = build_cohort_workbook(
                dt.date(2026, 1, 1), dt.date(2026, 3, 31),
                source_dir=self.source, template_path=TEMPLATE,
                workbook_path=path, today=TODAY,
            )
            self.assertEqual(result, path)
            workbook = openpyxl.load_workbook(path)
            sheet = workbook["NLR"]
            codes = [sheet.cell(row, 2).value for row in range(2, sheet.max_row + 1)
                     if sheet.cell(row, 2).value]
            self.assertEqual(codes, ["6082.HK", "2513.HK"])
            self.assertEqual(codes[0], "6082.HK")
            self.assertEqual(sheet["E2"].value.date(), dt.date(2026, 1, 2))
            sheet["L2"] = 123456
            workbook.save(path)
            workbook.close()

            build_cohort_workbook(
                dt.date(2026, 1, 1), dt.date(2026, 3, 31),
                source_dir=self.source, template_path=TEMPLATE,
                workbook_path=path, today=TODAY,
            )
            workbook = openpyxl.load_workbook(path, read_only=True)
            self.assertEqual(workbook["NLR"]["L2"].value, 123456)
            workbook.close()

    def test_refuses_to_replace_changed_issuer_membership(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "cohort.xlsx"
            build_cohort_workbook(
                dt.date(2026, 1, 1), dt.date(2026, 3, 31),
                source_dir=self.source, template_path=TEMPLATE,
                workbook_path=path, today=TODAY,
            )
            workbook = openpyxl.load_workbook(path)
            workbook["NLR"]["B2"] = "9999.HK"
            workbook.save(path)
            workbook.close()
            with self.assertRaisesRegex(ValueError, "differs"):
                build_cohort_workbook(
                    dt.date(2026, 1, 1), dt.date(2026, 3, 31),
                    source_dir=self.source, template_path=TEMPLATE,
                    workbook_path=path, today=TODAY,
                )

    def test_generated_config_keeps_workbook_and_artifacts_together(self):
        (ROOT / "datasets").mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=ROOT / "datasets") as temporary:
            destination = Path(temporary) / "cohort.xlsx"
            build_cohort_workbook(
                dt.date(2026, 1, 1), dt.date(2026, 3, 31),
                source_dir=self.source, template_path=TEMPLATE,
                workbook_path=destination, today=TODAY,
            )
            config = write_cohort_config(dt.date(2026, 1, 1), dt.date(2026, 3, 31), destination)
            cfg = load_cfg(config_path=config)
            self.assertEqual(cfg["_workbook_path"], destination)
            self.assertEqual(cfg["paths"]["out"], destination.parent / "out")
            self.assertEqual(len(read_companies(cfg)), 2)

    def test_collect_keeps_writeback_blocked_without_independent_review(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            code = "6082.HK"
            paths = {
                "out": root / "out",
                "packets": root / "packets",
                "allot_out": root / "allot_out",
                "allot_packets": root / "allot_packets",
            }
            for path in paths.values():
                path.mkdir()
            for out_key, packets_key in (("out", "packets"), ("allot_out", "allot_packets")):
                (paths[out_key] / "packets.json").write_text(json.dumps([{"code": code}]))
                (paths[packets_key] / "HKIPO-MB6082.md").write_text("packet")
                extracted = paths[out_key] / "extracted"
                extracted.mkdir()
                (extracted / "HKIPO-MB6082.json").write_text("{}")
            cfg = {"paths": paths, "_config_path": root / "cohort.yaml"}
            with patch.object(pipeline_run, "cmd_validate", return_value={"gate_pass": True}), \
                    patch("state.read_record", return_value={}), \
                    patch.object(pipeline_run, "cmd_write") as write:
                result = pipeline_run.cmd_collect(
                    cfg, argparse.Namespace(), [{"code": code}]
                )
            self.assertEqual(result, 3)
            write.assert_not_called()


if __name__ == "__main__":
    unittest.main()
