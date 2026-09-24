import csv
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "src"), str(ROOT)]
from write_back_expansion import WorkbookExpansionWriter


class ExpansionWriterTests(unittest.TestCase):
    def test_writer_uses_mapper_and_only_writes_selected_cohort_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook_path = root / "cohort.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "NLR"
            sheet.append(["File", "Code", "Name", "Prospectus", "Listing", None, None, None, None, None, "Offer"])
            sheet.append([1, "1234.HK", "Selected", datetime(2026, 1, 2), datetime(2026, 1, 10), None, None, None, None, None, 10.0])
            sheet.append([2, "5678.HK", "Outside", datetime(2026, 2, 2), datetime(2026, 2, 10), None, None, None, None, None, 20.0])
            workbook.save(workbook_path)
            workbook.close()

            out_master = root / "out" / "master"
            out_master.mkdir(parents=True)
            with (out_master / "stabilization_events.csv").open("w", newline="", encoding="utf-8-sig") as stream:
                writer = csv.DictWriter(
                    stream,
                    fieldnames=["stock_code", "stabilizing_manager", "stabilization_period_end"],
                )
                writer.writeheader()
                writer.writerow({"stock_code": "1234.HK", "stabilizing_manager": "Example Sponsor", "stabilization_period_end": "2026-02-09"})

            cfg = {
                "workbook": str(workbook_path),
                "workbook_path": workbook_path,
                "_workbook_path": workbook_path,
                "sheet": "NLR",
                "data_start_row": 2,
                "id_columns": {
                    "file_no": "A",
                    "stock_code": "B",
                    "name": "C",
                    "prospectus_date": "D",
                    "listing_date": "E",
                    "offer_price": "K",
                },
                "dataset": {"period_start": "2026-01-01", "period_end": "2026-01-31"},
                "filter_period": True,
                "paths": {"out": root / "out"},
            }
            result = WorkbookExpansionWriter(cfg=cfg).write_expansion()

            check = openpyxl.load_workbook(workbook_path, data_only=True)
            output = check["NLR"]
            self.assertEqual(result, 0)
            self.assertEqual(output.cell(1, 162).value, "Stabilizing manager")
            self.assertEqual(output.cell(2, 162).value, "Example Sponsor")
            self.assertEqual(output.cell(2, 163).value, "2026-02-09")
            self.assertIsNone(output.cell(3, 162).value)
            check.close()


if __name__ == "__main__":
    unittest.main()
