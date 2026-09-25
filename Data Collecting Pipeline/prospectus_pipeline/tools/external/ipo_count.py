#!/usr/bin/env python3
"""DE = 招股书日前 90 个自然日内港股主板普通 IPO 家数。

来源：HKEX Main Board 年度 New Listing Reports，按窗口自动确定年份。
窗口：[D−90, D)，按**上市日**计；同一股票代码只计一次。
剔除：介绍上市 / 无募资（募资额两行均为空或 0，或公司名含 Introduction）。
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2] if Path(__file__).resolve().parent.name == "external" else Path(__file__).resolve().parent
WS = ROOT.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "src"))
from run import load_cfg
from listing_reports import reports_for_interval
from sample_builder import audit_candidate, load_nlr_candidates

_cfg = load_cfg()
_configured_book = Path(_cfg["workbook"])
BOOK = _configured_book if _configured_book.is_absolute() else WS / _configured_book


SHEET = _cfg.get("sheet", "NLR")
HEADER = "HK ordinary IPO count in 90 calendar days before prospectus"


def to_date(v) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def load_listings(path: Path) -> list[dict]:
    match = re.search(r"(?:NLR)?(19\d{2}|20\d{2})", path.stem, re.I)
    if not match:
        raise ValueError(f"Cannot identify annual report year: {path.name}")
    candidates = load_nlr_candidates(path, int(match.group(1)))
    kept = []
    for candidate in candidates:
        candidate = audit_candidate(candidate)
        if candidate["inclusion_status"] == "INCLUDED" and candidate.get("listing_date"):
            kept.append({
                "code": candidate["stock_code"],
                "name": candidate["company_name"],
                "listing": candidate["listing_date"],
            })
    return kept


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--book", default=str(BOOK))
    ap.add_argument("--nlr", nargs="+", default=None,
                    help="指定已校验的年度报告；默认按工作簿招股书日期自动获取")
    ap.add_argument("--only", nargs="*", default=None, help="只处理指定股票代码")
    args = ap.parse_args()
    book = Path(args.book)

    wb_read = openpyxl.load_workbook(book, data_only=True)
    ws_read = wb_read[SHEET]
    col = None
    col_code = 2
    col_pd = 4
    for c in range(1, ws_read.max_column + 1):
        v = " ".join(str(ws_read.cell(1, c).value or "").replace("\n", " ").split()).strip().lower()
        if v == HEADER.lower():
            col = c
        elif "stock code" in v:
            col_code = c
        elif "date of prospectus" in v:
            col_pd = c
    if col is None:
        wb_read.close()
        raise SystemExit(f"找不到列 {HEADER!r}")

    companies = []
    for r in range(2, ws_read.max_row + 1):
        code = ws_read.cell(r, col_code).value
        pd = to_date(ws_read.cell(r, col_pd).value)
        if code in (None, "") or pd is None:
            continue
        c_str = str(code).strip()
        if args.only and c_str not in args.only:
            continue
        companies.append((r, c_str, pd))
    wb_read.close()
    if not companies:
        raise SystemExit("No issuers with prospectus dates to calculate")
    if args.nlr is None:
        source_dir = Path(_cfg["dataset"].get("report_source_dir") or WS / "sources")
        paths = reports_for_interval(
            min(item[2] for item in companies) - dt.timedelta(days=90),
            max(item[2] for item in companies) - dt.timedelta(days=1),
            source_dir,
        )
    else:
        paths = [Path(path) for path in args.nlr]

    listings, seen = [], set()
    for path in paths:
        recs = load_listings(path)
        print(f"{path.name}: {len(recs)} 家")
        for rec in recs:
            if rec["code"] in seen:
                continue
            seen.add(rec["code"])
            listings.append(rec)
    if not listings:
        raise SystemExit("No ordinary IPO listings in the supplied annual reports")
    print(f"合计普通 IPO {len(listings)} 家，上市日 "
          f"{min(x['listing'] for x in listings)} ~ {max(x['listing'] for x in listings)}")

    print(f"\n{'code':9s} {'招股书':12s} {'窗口':26s} {'家数':>4s}")
    values = []
    for r, code, pd in companies:
        lo, hi = pd - dt.timedelta(days=90), pd
        n = sum(1 for x in listings if lo <= x["listing"] < hi)
        values.append((r, n))
        print(f"{code:9s} {str(pd):12s} [{lo} , {hi}) {n:4d}")

    if args.dry_run:
        print("\n--dry-run：未写回")
        return 0

    sys.path.insert(0, str(ROOT / "src"))
    from workbook_transaction import workbook_transaction

    with workbook_transaction(book, operation="ipocount") as wb:
        ws = wb[SHEET]
        for r, n in values:
            ws.cell(r, col).value = int(n)

    print(f"\n已写回 DE（{len(values)} 家） -> {book.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
