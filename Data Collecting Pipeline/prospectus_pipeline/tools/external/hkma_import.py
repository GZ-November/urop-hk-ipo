#!/usr/bin/env python3
"""把手工导出的市场数据（HIBOR / 银行体系总结余）按日期并入工作簿。

输入：prospectus_pipeline/data/manual/hibor_balance.csv
      表头必须为 date,hibor_1m_pct,aggregate_balance_hkd_mn

口径（手册 §4.6）：对每家公司取"观察日严格早于招股书日期"的最后一条记录。
单位换算：HIBOR 百分数 -> 小数；总结余 百万港元 -> 基本单位。

写入：DF `1-month HIBOR before prospectus (%)`、
      DG `Banking system aggregate balance before prospectus (HK$)`。
      只改这两列的值，不动样式，不动其它列。
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import shutil
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
WS = ROOT.parent
BOOK = WS / "HKIPO-MB2026Q1.xlsx"
CSV_PATH = ROOT / "data" / "manual" / "hibor_balance.csv"

HIBOR_HEADER = "1-month HIBOR before prospectus (%)"
BALANCE_HEADER = "Banking system aggregate balance before prospectus (HK$)"
SHEET = "NLR"
CODE_COL = "B"
DATE_COL = "D"


def norm(s) -> str:
    return " ".join(str(s or "").replace("\n", " ").split()).strip().lower()


def parse_date(s: str) -> dt.date:
    s = str(s).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%Y/%m/%d"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"无法识别的日期：{s!r}")


def to_float(s: str, field: str, row_no: int) -> float:
    t = str(s).strip().replace(",", "")
    if t in ("", "-", "N/A", "na", "NA"):
        raise ValueError(f"第 {row_no} 行 {field} 为空")
    return float(t)


def load_csv(path: Path) -> list[dict]:
    if not path.exists():
        raise SystemExit(f"找不到 {path}\n请按 data/manual/README.md 的格式导出后重试。")
    rows = []
    with path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        need = {"date", "hibor_1m_pct", "aggregate_balance_hkd_mn"}
        missing = need - {norm(k) for k in (reader.fieldnames or [])}
        if missing:
            raise SystemExit(f"CSV 缺少列 {sorted(missing)}；实际表头 = {reader.fieldnames}")
        for i, raw in enumerate(reader, start=2):
            r = {norm(k): v for k, v in raw.items()}
            if not any(str(v).strip() for v in r.values()):
                continue
            rows.append({
                "date": parse_date(r["date"]),
                "hibor": to_float(r["hibor_1m_pct"], "hibor_1m_pct", i) / 100.0,
                "balance": to_float(r["aggregate_balance_hkd_mn"],
                                    "aggregate_balance_hkd_mn", i) * 1_000_000,
            })
    if not rows:
        raise SystemExit("CSV 没有数据行。")
    rows.sort(key=lambda x: x["date"])
    dup = [d for i, d in enumerate(rows) if i and rows[i - 1]["date"] == d["date"]]
    if dup:
        print(f"!! 警告：CSV 中有重复日期 {sorted({str(d) for d in dup})[:5]}（取最后一条）")
        dedup = {}
        for r in rows:
            dedup[r["date"]] = r
        rows = [dedup[d] for d in sorted(dedup)]
    print(f"CSV：{len(rows)} 条，区间 {rows[0]['date']} ~ {rows[-1]['date']}")
    return rows


def pick(rows: list[dict], d: dt.date) -> dict | None:
    """观察日严格早于招股书日期的最后一条。"""
    best = None
    for r in rows:
        if r["date"] < d:
            best = r
        else:
            break
    return best


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true", help="只预览，不写回")
    ap.add_argument("--book", default=str(BOOK))
    ap.add_argument("--csv", default=str(CSV_PATH))
    ap.add_argument("--only", nargs="*", default=None, help="只处理指定股票代码")
    args = ap.parse_args()
    book = Path(args.book)

    rows = load_csv(Path(args.csv))

    wb_read = openpyxl.load_workbook(book, data_only=True)
    ws_read = wb_read[SHEET]
    col_of = {}
    for c in range(1, ws_read.max_column + 1):
        v = ws_read.cell(1, c).value
        if v in (None, ""):
            continue
        n = norm(v)
        if n == norm(HIBOR_HEADER):
            col_of["hibor"] = c
        elif n == norm(BALANCE_HEADER):
            col_of["balance"] = c
        elif "stock code" in n:
            col_of["code"] = c
        elif "date of prospectus" in n:
            col_of["date"] = c
    if "hibor" not in col_of or "balance" not in col_of:
        wb_read.close()
        raise SystemExit(f"工作簿里找不到这两列，实际得到 {col_of}")

    code_col = col_of.get("code", openpyxl.utils.column_index_from_string(CODE_COL))
    date_col = col_of.get("date", openpyxl.utils.column_index_from_string(DATE_COL))

    # 取公司行
    companies = []
    for r in range(2, ws_read.max_row + 1):
        code = ws_read.cell(r, code_col).value
        date = ws_read.cell(r, date_col).value
        if code in (None, ""):
            continue
        c_str = str(code).strip()
        if args.only and c_str not in args.only:
            continue
        if not isinstance(date, (dt.date, dt.datetime)):
            print(f"!! 第 {r} 行 {code} 的招股书日期不是日期型，跳过")
            continue
        if isinstance(date, dt.datetime):
            date = date.date()
        companies.append((r, c_str, date))
    wb_read.close()


    print(f"工作簿：{len(companies)} 家公司\n")
    print(f"{'code':10s} {'招股书日':12s} {'观察日':12s} {'HIBOR':>9s} {'总结余(HK$)':>16s}")
    filled = 0
    no_row = []
    for r, code, d in companies:
        h = pick(rows, d)
        if h is None:
            no_row.append((code, d))
            print(f"{code:10s} {str(d):12s} {'—':12s} {'无可用观察值（日期早于 CSV 起点）':>28s}")
            continue
        filled += 1
        print(f"{code:10s} {str(d):12s} {str(h['date']):12s} "
              f"{h['hibor']*100:8.5f}% {h['balance']:16,.0f}")

    print(f"\n可填 {filled}/{len(companies)} 家；无观察值 {len(no_row)} 家")
    if no_row:
        print("  建议把 CSV 起点提前到早于最早招股书日期：", min(d for _, d in no_row))

    if args.dry_run:
        print("\n--dry-run：未写回。确认无误后去掉 --dry-run 重跑。")
        return 0

    sys.path.insert(0, str(ROOT / "src"))
    from workbook_transaction import workbook_transaction

    with workbook_transaction(book, operation="market-hkma") as wb:
        ws = wb[SHEET]
        for r, code, d in companies:
            h = pick(rows, d)
            if h is None:
                continue
            ws.cell(r, col_of["hibor"]).value = round(h["hibor"], 8)
            ws.cell(r, col_of["balance"]).value = h["balance"]

    print(f"\n已写回 DF/DG 两列（{filled} 家） -> {book.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
