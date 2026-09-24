#!/usr/bin/env python3
"""Generate a VC/PE summary for the cohort selected by PIPELINE_CONFIG."""
from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
WS = ROOT.parent
sys.path[:0] = [str(ROOT), str(ROOT / "src")]

from run import load_cfg  # noqa: E402
from write_back import norm_header  # noqa: E402

FIELD_KEYS = {
    "Pre-IPO VC/PE backing (1=yes; 0=no)": "col_BA",
    "Pre-IPO VC backing (1=yes; 0=no)": "col_vc_backed",
    "Pre-IPO PE backing (1=yes; 0=no)": "col_pe_backed",
    "Pre-IPO CVC backing (1=yes; 0=no)": "col_cvc_backed",
    "Pre-IPO State/Gov backing (1=yes; 0=no)": "col_gov_backed",
    "Top-tier VC/PE backing (1=yes; 0=no)": "col_top_tier_vc",
    "Key Pre-IPO investors": "col_pre_ipo_investors",
    "Pre-IPO institutional shareholding (%)": "col_vc_pe_stake",
    "Pre-IPO investor board seat (1=yes; 0=no)": "col_vc_board_seat",
    "Earliest Pre-IPO investment round": "col_earliest_round",
    "Pre-IPO holding duration (years)": "col_holding_duration",
}


def main() -> int:
    parser = argparse.ArgumentParser(description="为当前配置的 IPO cohort 生成 VC/PE 统计报告")
    parser.add_argument("--config", help="cohort config；也可设置 PIPELINE_CONFIG")
    args = parser.parse_args()
    if args.config:
        os.environ["PIPELINE_CONFIG"] = str(Path(args.config).expanduser().resolve())
    cfg = load_cfg()
    book = Path(cfg.get("workbook_path") or (WS / cfg["workbook"]))
    if not book.is_file():
        raise FileNotFoundError(f"Workbook not found: {book}")
    wb = openpyxl.load_workbook(book, read_only=True, data_only=True)
    try:
        ws = wb[cfg["sheet"]]
        headers: dict[str, list[int]] = {}
        for column in range(1, ws.max_column + 1):
            value = ws.cell(1, column).value
            if value not in (None, ""):
                headers.setdefault(norm_header(value), []).append(column)
        cols = {}
        for header, key in FIELD_KEYS.items():
            matched = headers.get(norm_header(header), [])
            if len(matched) != 1:
                raise ValueError(f"Expected exactly one workbook column for {header!r}; got {matched}")
            cols[key] = matched[0]
        id_col = openpyxl.utils.column_index_from_string(cfg["id_columns"]["stock_code"])
        name_col = openpyxl.utils.column_index_from_string(cfg["id_columns"]["name"])
        rows = []
        for row_no in range(cfg["data_start_row"], ws.max_row + 1):
            code = ws.cell(row_no, id_col).value
            if code in (None, ""):
                continue
            rows.append({"code": str(code).strip(), "name": ws.cell(row_no, name_col).value or "",
                         **{key: ws.cell(row_no, col).value for key, col in cols.items()}})
    finally:
        wb.close()

    n = len(rows)
    cohort = cfg.get("dataset", {}).get("cohort", cfg.get("dataset", {}).get("id", "active cohort"))
    cohort_id = cfg.get("dataset", {}).get("id", "cohort")
    out_path = cfg["paths"]["out"] / f"{cohort_id}_VC_PE_Research_Report.md"

    def flag_summary(key):
        values = [row[key] for row in rows]
        known = [int(value) for value in values if value in (0, 1, "0", "1")]
        yes = sum(value == 1 for value in known)
        return yes, len(known), n - len(known)

    def as_number(value):
        if value in (None, "", "NA", "NaN", "N/A", "-"):
            return None
        try:
            number = float(value)
            return number if __import__("math").isfinite(number) else None
        except (TypeError, ValueError):
            return None

    backed = [row for row in rows if row["col_BA"] in (1, "1")]
    stakes = [as_number(r["col_vc_pe_stake"]) for r in backed]
    stakes = [value for value in stakes if value is not None]
    durations = [as_number(r["col_holding_duration"]) for r in backed]
    durations = [value for value in durations if value is not None]
    stats = [
        ("Pre-IPO 机构总覆盖", "col_BA"),
        ("早期风险投资 (VC)", "col_vc_backed"),
        ("私募股权基金 (PE)", "col_pe_backed"),
        ("产业资本/企业创投 (CVC)", "col_cvc_backed"),
        ("国资/产业引导基金 (Gov)", "col_gov_backed"),
        ("顶级机构认证 (Top-tier)", "col_top_tier_vc"),
        ("董事会席位派驻", "col_vc_board_seat"),
    ]
    md = [
        f"# 香港主板 {cohort} IPO Pre-IPO VC/PE 汇总报告", "",
        f"> 数据源：`{book.name}`，工作表 `{cfg['sheet']}`；字段按规范化表头定位。",
        f"> 样本范围：当前配置 cohort（N = {n}）。", "",
        "## 核心统计", "", "| 指标 | 字段 | Yes / 已知值 | 缺失/无效 | Yes 占已知值 |", "|---|---|---:|---:|---:|",
    ]
    for label, key in stats:
        yes, known, missing = flag_summary(key)
        share = f"{yes / known * 100:.1f}%" if known else "n/a"
        md.append(f"| {label} | `{key}` | {yes} / {known} | {missing} | {share} |")
    md += [
        f"| 上市前机构平均持股比例（有 VC/PE 样本） | `col_vc_pe_stake` | {sum(stakes)/len(stakes)*100:.1f}% | n={len(stakes)} |" if stakes else "| 上市前机构平均持股比例 | `col_vc_pe_stake` | n/a | n=0 |",
        f"| 平均持有年限（有 VC/PE 样本） | `col_holding_duration` | {sum(durations)/len(durations):.2f} 年 | n={len(durations)} |" if durations else "| 平均持有年限 | `col_holding_duration` | n/a | n=0 |",
        "", "## 公司明细", "",
        "| 股票代码 | 公司 | VC/PE | VC | PE | CVC | 国资 | 顶级 | 董事席位 | 最早轮次 | 持有年限 | 机构持股 | 主要机构投资者 |",
        "|---|---|---:|---:|---:|---:|---:|---:|---:|---|---:|---:|---|",
    ]
    for row in rows:
        stake = row["col_vc_pe_stake"]
        duration = row["col_holding_duration"]
        stake_s = f"{stake*100:.2f}%" if isinstance(stake, (int, float)) else "NA"
        duration_s = f"{duration:.2f}" if isinstance(duration, (int, float)) else "NA"
        md.append("| " + " | ".join(str(v if v is not None else "NA") for v in (
            row["code"], row["name"], row["col_BA"], row["col_vc_backed"], row["col_pe_backed"],
            row["col_cvc_backed"], row["col_gov_backed"], row["col_top_tier_vc"],
            row["col_vc_board_seat"], row["col_earliest_round"], duration_s, stake_s,
            row["col_pre_ipo_investors"])) + " |")
    out_path.write_text("\n".join(md) + "\n", encoding="utf-8")
    print(f"Generated {out_path} ({n} companies, cohort={cohort_id})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
