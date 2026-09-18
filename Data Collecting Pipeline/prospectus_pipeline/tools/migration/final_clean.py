#!/usr/bin/env python3
"""终检清理：把空数据行对齐成与已填数据行完全一致的格式，并修正冻结窗格。

问题来源：老师模板在 40 行以后存在"拖拽式"条带状格式（同一列在不同行段
格式不同，例如 A 列 `0.00_);[Red]`、B 列 `0000`、H 列 `dd/mm/yy;`）。这些是
空行，但将来录入会继承错误格式，导致同一列前后显示不一致。

处理：
  1) 把 A 列到最后一列、第 40 行起到 1427 行的字体/数字格式/对齐，
     统一复制自第 2 行（已填数据行）——即该列应有的格式。
  2) 冻结窗格显式写成 D2（XML 里 xSplit=3, ySplit=1, topLeftCell=D2），
     并把活动单元格复位到 A2。
  3) 第 1–39 行（表头 + 38 家数据）完全不碰。
"""
from __future__ import annotations

import datetime as dt
import shutil
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.views import Pane, Selection

ROOT = Path(__file__).resolve().parent
BOOK = ROOT.parent / "HKIPO-MB2026Q1.xlsx"
FIRST_EMPTY_ROW = 40          # 数据到第 39 行
COPY_FROM_ROW = 2             # 以该行为各列的格式样板


def main() -> None:
    backup = BOOK.with_name(f"{BOOK.stem}.backup-before-finalclean-{dt.datetime.now():%Y%m%d-%H%M%S}.xlsx")
    shutil.copy2(BOOK, backup)

    wb = openpyxl.load_workbook(BOOK)
    ws = wb["NLR"]
    last_col = ws.max_column
    max_row = ws.max_row

    changed = 0
    for c in range(1, last_col + 1):
        src = ws.cell(COPY_FROM_ROW, c)
        font = copy(src.font)
        fmt = src.number_format
        align = copy(src.alignment)
        for r in range(FIRST_EMPTY_ROW, max_row + 1):
            cell = ws.cell(r, c)
            cell.font = copy(font)
            cell.alignment = copy(align)
            cell.number_format = fmt
            changed += 1

    # 冻结窗格：显式 D2
    ws.freeze_panes = None
    ws.freeze_panes = "D2"
    ws.sheet_view.pane = Pane(xSplit=3, ySplit=1, topLeftCell="D2", activePane="bottomRight", state="frozen")
    ws.sheet_view.selection = [Selection(pane="bottomRight", activeCell="A2", sqref="A2")]

    wb.save(BOOK)
    wb.close()
    print(f"backup: {backup.name}")
    print(f"已对齐第 {FIRST_EMPTY_ROW}–{max_row} 行、{last_col} 列的格式（{changed} 个单元格）")


if __name__ == "__main__":
    main()
