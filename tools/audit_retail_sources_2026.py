"""Read-only source follow-up: never edits frozen tiers, reviews or workbooks."""
from __future__ import annotations

import hashlib
import json
from decimal import Decimal
from pathlib import Path
import re
import sys

import numpy as np
import pandas as pd
import pymupdf
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "analysis"))
from research_inputs import load_panel, select_2026  # noqa: E402
from shared.allocation_review import reviewed_tiers  # noqa: E402

OUT = ROOT / "pipeline/reports/data_gap_collection/source_followup_2026-10-03"
PDF = ROOT / "pipeline/prospectus_pipeline/data/allot/pdf"
CODES = "2290 1392 2335 6106 3952 2667 6880 7656 7687 9971 1377 1770 2475 2797 3752 6951 2249 6745 9976".split()
URLS = {
    "2649_clarification": "https://www.hkexnews.hk/listedco/listconews/sehk/2026/0311/2026031100837_c.pdf",
    "3355_clarification": "https://www1.hkexnews.hk/listedco/listconews/sehk/2026/0327/2026032702666.pdf",
    "2649_ccass": "https://www.hkex.com.hk/-/media/HKEX-Market/Services/Circulars-and-Notices/Participant-and-Members-Circulars/HKSCC/2026/ce_HKSCC_SKA_074_2026.pdf",
}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    tiers = reviewed_tiers(ROOT)
    panel = select_2026(load_panel()).set_index("Stock Code")
    index = pd.read_csv(OUT.parent / "allocation_source_index.csv").set_index("code")
    coverage = pd.read_csv(OUT.parent / "allocation_coverage.csv").set_index("code")
    inputs = [OUT.parent / "allocation_tiers_clean.csv", OUT.parent / "allocation_tiers_candidates.csv",
              OUT.parent / "allocation_review.json", OUT.parent / "allocation_coverage.csv",
              OUT.parent / "board_lots.csv", ROOT / "pipeline/exports/HKIPO-MB-MASTER_clean.csv", Path(__file__)]
    sources = []
    for name, url in URLS.items():
        path = PDF / "source_followup_2026-10-03" / (name + ".pdf")
        inputs.append(path)
        with pymupdf.open(path) as document:
            pages = [p.get_text() for p in document]
        sources.append({"id": name, "url": url, "pdf": str(path.relative_to(ROOT)),
                        "sha256": digest(path), "pages": len(pages)})
        if name == "2649_clarification":
            assert "500股H股" in pages[0] and "200股H股" in pages[0]
            assert "2026年3月11日" in pages[1]
        elif name == "2649_ccass":
            assert re.search(r"ALSCO Pooling Service Co., Ltd.\s+2649\s+500", pages[0])
        else:
            pattern = (r"(?m)^([\d,]+)\s*\n([\d,]+)\s*\n([\d,]+) H Shares plus "
                       r"([\d,]+) out of ([\d,]+)\s+applicants to receive an additional\s+"
                       r"([\d,]+)\s+H Shares\s+([\d.]+)%")
            records = []
            for match in re.finditer(pattern, pages[1]):
                values = [int(s.replace(",", "")) for s in match.groups()[:6]]
                records.append(dict(zip(["applied_shares", "applicants", "guaranteed_shares",
                                         "ballot_winners", "ballot_denominator", "ballot_extra_shares"], values),
                                    printed_allocation_pct=float(match[7]), pdf_page=2))
            assert len(records) == 10
            official = pd.DataFrame(records).sort_values("applied_shares")
            prior = tiers[tiers.code.eq("3355.HK") & tiers.pool.eq("B")].sort_values("applied_shares")
            for col in official.columns.difference(["pdf_page"]):
                assert np.allclose(official[col], prior[col]), col
            official["allocated_shares"] = official.guaranteed_shares * official.applicants + official.ballot_winners * official.ballot_extra_shares
            assert official.applicants.sum() == 11492
            assert official.allocated_shares.sum() == 2000000
            official["source_url"] = url
            official["source_sha256"] = digest(path)
            official["status"] = "directly_disclosed_in_clarification_matches_prior_inference"
            official.to_csv(OUT / "3355_official_pool_b.csv", index=False)

    rows, pairs_all = [], []
    workbook_values = {}
    for quarter in ["Q2", "Q3"]:
        path = ROOT / f"pipeline/cohorts/HKIPO-MB2026{quarter}.xlsx"
        inputs.append(path)
        workbook = load_workbook(path, read_only=True, data_only=True)
        headers = [cell.value for cell in workbook["NLR"][1]]
        applied_column = headers.index("Public valid applied shares")
        for row in workbook["NLR"].iter_rows(min_row=2, values_only=True):
            if row[1] in {c + ".HK" for c in CODES}:
                workbook_values[row[1]] = row[applied_column]
        workbook.close()
    for bare in CODES:
        code = bare + ".HK"
        current = tiers[tiers.code.eq(code)]
        path = PDF / f"HKIPO-MB{bare}.pdf"
        formal = ROOT / f"pipeline/prospectus_pipeline/out/allot/extracted/HKIPO-MB{bare}.json"
        inputs.extend([path, formal])
        before = OUT / "formal_json_repair/before" / formal.name
        historical = before if before.exists() else formal
        if historical != formal:
            inputs.append(historical)
        fields = json.loads(historical.read_text())["fields"]
        current_formal_value = json.loads(formal.read_text())["fields"]["col_CO"]["value"]
        old = fields["col_CO"]
        pairs = []
        with pymupdf.open(path) as document:
            for page in sorted(current.pdf_page.unique()):
                text = document[int(page) - 1].get_text(sort=True)
                # Independently read the first two numeric columns from PDF layout.
                # Wrapped allocation wording can start above the numeric row.
                pattern = r"(?m)^\s*([\d,]+)\s+([\d,]+)\s+(?:\D|(?=[\d,]+\s+(?:out\s+of|(?:H\s+)?Shares)))"
                for match in re.finditer(pattern, text):
                    applied, applicants = [int(x.replace(",", "")) for x in match.groups()]
                    pairs.append((applied, applicants))
                    pairs_all.append({"code": code, "applied_shares": applied, "applicants": applicants,
                                      "pdf_page": int(page), "pdf_layout_excerpt": match[0].strip(),
                                      "source_url": index.loc[code, "official_pdf_url"], "pdf_sha256": digest(path)})
        expected = sorted(zip(current.applied_shares.astype(int), current.applicants.astype(int)))
        assert sorted(pairs) == expected, f"{code}: PDF operands differ from stored tiers"
        exact = sum(a * n for a, n in pairs)
        assert sum(n for _, n in pairs) == panel.loc[code, "Public applicants"]
        assert exact == panel.loc[code, "Public valid applied shares"]
        assert exact == workbook_values[code]
        reserved = int(coverage.loc[code, "overseas_employee_reserved_shares"])
        assert current.allocated_shares.sum() == panel.loc[code, "Final public offer shares"] - reserved
        note = old.get("note", "")
        cause = "documented_rounded_subscription_reconstruction" if note else "original_error_mechanism_not_documented"
        reconstructed = None
        arithmetic_residual = None
        if note:
            numbers = re.findall(r"\d[\d,]*(?:\.\d+)?", note)
            initial, ratio = [Decimal(s.replace(",", "")) for s in numbers[:2]]
            reconstructed = float(initial * ratio)
            arithmetic_residual = float(Decimal(int(old["value"])) - initial * ratio)
            if abs(arithmetic_residual) > 0.5:
                cause = "documented_reconstruction_with_arithmetic_mismatch"
        if bare == "9976":
            cause = "compatible_with_rounded_subscription_reconstruction_not_documented"
        rows.append({"code": code, "status": "pdf_tier_operands_verified_exact_sum_derived",
                     "tiers": len(pairs), "pdf_pages": ";".join(map(str, sorted(current.pdf_page.unique()))),
                     "old_formal_value": old["value"], "current_master_value": int(panel.loc[code, "Public valid applied shares"]),
                     "current_formal_value": current_formal_value,
                     "exact_pdf_tier_sum": exact, "delta": exact - int(old["value"]),
                     "current_workbook_value": int(workbook_values[code]),
                     "cause_status": cause, "old_formal_note": note,
                     "old_note_product_recomputed": reconstructed,
                     "old_value_minus_note_product": arithmetic_residual,
                     "old_formal_page": old.get("page"), "source_url": index.loc[code, "official_pdf_url"],
                     "pdf_sha256": digest(path), "formal_sha256": digest(historical),
                     "current_formal_sha256": digest(formal),
                     "formal_master_consistent": int(current_formal_value) == exact})
    pd.DataFrame(rows).to_csv(OUT / "nineteen_applied_totals.csv", index=False)
    pd.DataFrame(pairs_all).to_csv(OUT / "nineteen_pdf_operands.csv", index=False)
    # A study-specific evidence overlay; not a replacement pipeline approval.
    pd.DataFrame([{"code": "2649.HK", "old_board_lot_units": 200, "board_lot_units": 500,
                   "minimum_application_units": 500, "status": "official_clarification_and_ccass_confirmed",
                   "source_url": URLS["2649_clarification"], "pdf_page": 1,
                   "source_sha256": sources[0]["sha256"]}]).to_csv(OUT / "board_lot_followup.csv", index=False)
    for name in ["repair_manifest.json", "independent_review.json", "applied.json"]:
        repair_record = OUT / "formal_json_repair" / name
        if repair_record.exists():
            inputs.append(repair_record)
    manifest = {"review_date": "2026-10-03", "observation_cutoff": "2026-09-30",
                "scope": "2649 lot correction; 3355 ten Pool B rules; 19 issuers' PDF application quantities/counts and exact sums",
                "not_a_pipeline_gate_approval": True, "canonical_data_changed": False,
                "formal_json_repair_applied": (OUT / "formal_json_repair/applied.json").exists(),
                "current_formal_totals_match": sum(r["formal_master_consistent"] for r in rows),
                "allocation_rules_for_19_independently_rereviewed": False,
                "sources": sources, "inputs": {str(p.relative_to(ROOT)): digest(p) for p in inputs},
                "outputs": {p.name: digest(p) for p in sorted(OUT.glob("*.csv"))},
                "nineteen_issuer_count": len(rows), "pdf_operand_rows": len(pairs_all),
                "documented_reconstruction_note_count": sum(r["cause_status"].startswith("documented") for r in rows),
                "old_formula_reproduces_count": sum(r["cause_status"] == "documented_rounded_subscription_reconstruction" for r in rows),
                "old_formula_arithmetic_mismatch_count": sum(r["cause_status"] == "documented_reconstruction_with_arithmetic_mismatch" for r in rows)}
    (OUT / "source_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    print(f"Verified {len(pairs_all)} tier operand pairs for 19 issuers; 3355 Pool B 10/10 match; 2649 lot=500.")


if __name__ == "__main__":
    main()
