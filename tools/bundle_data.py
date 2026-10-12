#!/usr/bin/env python3
"""Package selected cohort snapshots with CSV, definitions, evidence and quality reports."""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
from pathlib import Path
import re
import shutil
import sys
import tempfile
import zipfile

ROOT = Path(__file__).resolve().parents[1]
PIPE = ROOT / "pipeline/prospectus_pipeline"
sys.path.insert(0, str(PIPE / "src"))

from audit import run_audit  # noqa: E402
from auto_fill import inventory  # noqa: E402
from boards import normalize_board  # noqa: E402
from codebook import export_all  # noqa: E402
from cohort import load_cfg, read_companies  # noqa: E402
from gem_template import detail_schema, update_gem_view  # noqa: E402
from state import record_path  # noqa: E402


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def sheet_csv(workbook, sheet_name: str, destination: Path) -> None:
    sheet = workbook[sheet_name]
    with destination.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow([cell.value for cell in sheet[1]])
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if any(value not in (None, "") for value in row):
                writer.writerow(row)


def build_bundle(configs: list[Path], output: Path, require_reviewed: bool = False) -> dict:
    import openpyxl
    if output.exists():
        raise ValueError(f"Delivery already exists; choose another --output: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="hkipo-delivery-") as temporary:
        staging = Path(temporary)
        instruction_dir = staging / "instructions"
        instruction_dir.mkdir()
        for name in ("AGENT_INSTRUCTIONS.md", "GEM_PIPELINE.md"):
            source = ROOT / "docs" / name
            if source.is_file():
                shutil.copy2(source, instruction_dir / name)
        cohort_summaries = []
        used_ids = set()
        for config in configs:
            cfg = load_cfg(config_path=config)
            cohort_id = re.sub(r"[^A-Za-z0-9_.-]", "_", cfg["dataset"]["id"])
            if cohort_id in used_ids:
                raise ValueError(f"Duplicate cohort delivery id: {cohort_id}")
            used_ids.add(cohort_id)
            destination = staging / "cohorts" / cohort_id
            destination.mkdir(parents=True)
            configured_period = [cfg["dataset"]["period_start"], cfg["dataset"]["period_end"]]
            companies = read_companies(cfg)
            # Work from a workbook copy: packaging never changes the research workbook.
            workbook_path = destination / Path(cfg["_workbook_path"]).name
            shutil.copy2(cfg["_workbook_path"], workbook_path)
            book = openpyxl.load_workbook(workbook_path)
            try:
                selected_rows = {item["row"] for item in companies}
                for sheet in book:
                    for row in range(sheet.max_row, cfg["data_start_row"] - 1, -1):
                        if row not in selected_rows:
                            sheet.delete_rows(row)
                if normalize_board(cfg["dataset"].get("board")) == "gem":
                    update_gem_view(book)
                    sheet_csv(book, "GEM", destination / "GEM_research_view.csv")
                    sheet_csv(book, "GEM_details", destination / "GEM_extensions.csv")
                    (destination / "GEM_extension_codebook.json").write_text(json.dumps(detail_schema(), indent=2) + "\n")
                book.save(workbook_path)
            finally:
                book.close()
            snapshot_cfg = dict(cfg, workbook=str(workbook_path), _workbook_path=workbook_path, workbook_path=workbook_path)
            exports = export_all(snapshot_cfg, out_dir=destination)
            # Reports describe the original selected cohort, preserving credential cell addresses.
            report = run_audit(cfg, only=[item["code"] for item in companies], target="all",
                               out_json=destination / "audit.json", out_md=destination / "audit.md", log=lambda *_: None)
            if normalize_board(cfg["dataset"].get("board")) == "gem":
                gem_report = run_audit(cfg, only=[item["code"] for item in companies], target="gem",
                                       out_json=destination / "audit_gem.json", out_md=destination / "audit_gem.md", log=lambda *_: None)
                report["stats"].update(gem_report["stats"])
            status = inventory(cfg)
            gates = [{key: item[key] for key in ("code", "prospectus_validated", "prospectus_reviewed",
                      "prospectus_aligned", "allot_validated", "allot_reviewed", "allot_aligned", "gem_extension_pending") if key in item}
                     for item in status]
            fully_reviewed = bool(gates) and all(item["prospectus_aligned"] and item["allot_aligned"]
                                 and not item.get("gem_extension_pending", False) for item in gates)
            reconciled = not any(stats.get(key, 0) for stats in report["stats"].values()
                                 for key in ("excel_missing", "json_missing", "value_mismatch", "json_files_absent"))
            if require_reviewed and (not fully_reviewed or not reconciled):
                raise ValueError(f"{cohort_id}: evidence/writeback or reconciliation incomplete; see normal bundle quality reports")
            (destination / "status.json").write_text(json.dumps(gates, indent=2) + "\n")
            targets = ["prospectus", "allot"] + (["gem"] if normalize_board(cfg["dataset"].get("board")) == "gem" else [])
            from extraction_targets import target_paths
            for target in targets:
                schema_path, target_out, _ = target_paths(cfg, target)
                evidence = destination / "evidence" / target
                evidence.mkdir(parents=True)
                shutil.copy2(schema_path, evidence / "schema.json")
                for company in companies:
                    code = company["code"]
                    payload = target_out / "extracted" / f"HKIPO-MB{code.split('.')[0]}.json"
                    if payload.exists():
                        shutil.copy2(payload, evidence / payload.name)
                    for stage in ("extracted", "validated", "reviewed", "written"):
                        record = record_path(cfg, target, code, stage)
                        if record.exists():
                            shutil.copy2(record, evidence / f"{code}.{stage}.json")
            # Keep disclosure URL indices, not raw PDFs or executable credentials.
            for name in ("found.json", "allotment_index.json"):
                index = cfg["paths"]["out"] / name
                if index.exists():
                    shutil.copy2(index, destination / name)
            # Absolute runtime paths are provenance only; never ship them as runnable configs.
            cohort_summaries.append({"id": cohort_id, "board": cfg["dataset"].get("board", "Main Board"),
                                     "requested_period": configured_period,
                                     "listing_dates": [item["listing_date"].isoformat() if item.get("listing_date") else None for item in companies],
                                     "issuers": len(companies), "fully_reviewed_and_aligned": fully_reviewed,
                                     "reconciled": reconciled, "audit_stats": report["stats"],
                                     "csv": str(Path(exports["csv_path"]).relative_to(staging))})
        summary = {"created_at": dt.datetime.now(dt.timezone.utc).isoformat(), "cohorts": cohort_summaries,
                   "note": "A snapshot can contain legacy values and evidence gaps. Passing code checks is not semantic certification. Review audit/states before analysis."}
        (staging / "quality_summary.json").write_text(json.dumps(summary, indent=2) + "\n")
        (staging / "README.md").write_text("# IPO research data delivery\n\nOpen each cohort workbook or clean CSV and its codebook. GEM also includes a research-layout CSV and a separate extension dictionary.\n\nRead quality_summary.json, audit reports and status.json before analysis. Blank/NaN/NA are missing, not zero. Monetary values retain documented currencies and units. Review credentials refer to original source packets that can be re-downloaded; hashes do not establish semantic truth. Raw PDFs, full text, virtual environments and secrets are excluded.\n\nCurrent built-in research scripts select Main Board 2026 listings; analyze GEM separately under an explicit sample and specification.\n", encoding="utf-8")
        files = sorted(path for path in staging.rglob("*") if path.is_file())
        manifest = [{"path": str(path.relative_to(staging)), "bytes": path.stat().st_size, "sha256": digest(path)} for path in files]
        (staging / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
        temporary_archive = output.with_name("." + output.name + ".tmp")
        try:
            with zipfile.ZipFile(temporary_archive, "w", zipfile.ZIP_DEFLATED) as archive:
                for path in sorted(staging.rglob("*")):
                    if path.is_file():
                        archive.write(path, str(path.relative_to(staging)))
            temporary_archive.replace(output)
        finally:
            temporary_archive.unlink(missing_ok=True)
    return summary


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", action="append", type=Path, help="Repeat for multiple cohorts; paths relative to repository root")
    parser.add_argument("--output", type=Path, default=ROOT / "deliverables/ipo-data.zip")
    parser.add_argument("--require-reviewed", action="store_true", help="Refuse legacy/unreviewed/mismatched evidence")
    args = parser.parse_args()
    configs = args.config or sorted(PIPE.glob("config_*.yaml"))
    configs = [(ROOT / path).resolve() if not path.is_absolute() else path for path in configs]
    if not configs:
        parser.error("No cohort configs found")
    try:
        summary = build_bundle(configs, args.output.resolve(), args.require_reviewed)
    except (ValueError, OSError) as exc:
        print(f"Delivery failed: {exc}", file=sys.stderr)
        return 1
    print(f"Delivery: {args.output.resolve()} ({sum(c['issuers'] for c in summary['cohorts'])} issuers)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
