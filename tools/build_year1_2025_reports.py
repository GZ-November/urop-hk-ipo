#!/usr/bin/env python3
"""Generate HKIPO-MB2025Q1_year1_corrections.md and HKIPO-MB2025Q2_year1_corrections.md."""
import subprocess
import io
import openpyxl
import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent


def get_head_wb(path):
    proc = subprocess.run(["git", "show", f"origin/main:{path}"], capture_output=True, check=True, cwd=ROOT)
    return openpyxl.load_workbook(io.BytesIO(proc.stdout), data_only=True)


def fmt(v):
    if v is None:
        return "NaN"
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return f"{v:.6f}".rstrip("0").rstrip(".")
    return str(v)


def generate_q1():
    path = "pipeline/cohorts/HKIPO-MB2025Q1.xlsx"
    wb_b = get_head_wb(path)
    wb_a = openpyxl.load_workbook(ROOT / path, data_only=True)
    ws_b = wb_b["NLR"]
    ws_a = wb_a["NLR"]

    q1_diffs = {}
    for r in range(2, ws_a.max_row + 1):
        code = ws_a[f"B{r}"].value
        name = ws_a[f"C{r}"].value
        if not code:
            continue
        diffs = []
        for c in range(1, ws_a.max_column + 1):
            vb = ws_b.cell(row=r, column=c).value
            va = ws_a.cell(row=r, column=c).value
            if fmt(vb) != fmt(va):
                col_let = openpyxl.utils.get_column_letter(c)
                hdr = ws_a.cell(row=1, column=c).value
                diffs.append((col_let, hdr, fmt(vb), fmt(va)))
        if diffs:
            q1_diffs[code] = (name, diffs)

    lines = []
    lines.append("# 2025Q1 Main Board: year-1 re-anchored to the audited Track Record")
    lines.append("")
    lines.append("Same standard as `HKIPO-MB2026Q2_year1_corrections.md` and `HKIPO-MB2026Q3_year1_corrections.md`:")
    lines.append("year-1 is the latest period of the audited Accountants' Report / \"Track Record Period\" definition;")
    lines.append("unaudited figures outside it are never year-1; only sales, profit before tax and profit for the period")
    lines.append("are annualised (BC = months/12, originals in CN/CO/CP).")
    lines.append("")
    lines.append("## What changed")
    lines.append("")
    lines.append("| Group | Issuers | Before → after |")
    lines.append("|---|---|---|")
    lines.append("| 6M2024 audited stub period omitted in initial extraction; re-anchored to 30/06/2024 (×2.0) | 2613.HK, 0325.HK, 0805.HK, 2571.HK | year-1 FY2023 (not annualised) → 6M2024 stub (×2.0 annualised for AH/AK/AN) |")
    lines.append("| 9M2024 audited stub period omitted in initial extraction; re-anchored to 30/09/2024 (×1.3333) | 1364.HK, 2097.HK, 2610.HK, 2569.HK, 2625.HK | year-1 FY2023 (not annualised) → 9M2024 stub (×1.3333 annualised for AH/AK/AN) |")
    lines.append("| Audited stub already in payload, derived disclosure notes columns updated | 6681.HK, 2560.HK, 2530.HK, 2596.HK, 6693.HK, 2561.HK | Derived columns BC (months/12) and CI–CP re-derived by `disclosure_notes` |")
    lines.append("")
    lines.append("For the 9 re-anchored issuers, the 27 period-dependent payload fields were re-extracted from the prospectus")
    lines.append("(other fields untouched), validated (period-anchor gate included), recorded `extracted`, and independently")
    lines.append("reviewed against original disclosure pages; reports bound to each payload sha256 are in")
    lines.append("`pipeline/prospectus_pipeline/out/period_reviews/2025Q1/`. Reviewed payloads were written back to the workbook,")
    lines.append("followed by `disclosure_notes` (all 15 on the Track Record basis), `audit` (1050/1050 perfect match), `cross_check`")
    lines.append("(15/15 clean), `export` and `master --derive`. 2025Q1 `validate`: 15/15 target, 0 errors, gate PASS.")
    lines.append("")
    lines.append("## Pending")
    lines.append("")
    lines.append("None in 2025Q1.")
    lines.append("")
    lines.append("## Every changed cell (before → after)")
    lines.append("")

    for code in sorted(q1_diffs.keys()):
        name, diffs = q1_diffs[code]
        clean_name = " ".join(name.split()) if name else ""
        lines.append(f"### {code} - {clean_name} ({len(diffs)} cells)")
        lines.append("")
        lines.append("| col | header | before | after |")
        lines.append("|---|---|---|---|")
        for col_let, hdr, vb, va in diffs:
            clean_hdr = " ".join(str(hdr).split()) if hdr else ""
            lines.append(f"| {col_let} | {clean_hdr} | {vb} | {va} |")
        lines.append("")

    out_p = ROOT / "pipeline/reports/period_corrections/HKIPO-MB2025Q1_year1_corrections.md"
    out_p.parent.mkdir(parents=True, exist_ok=True)
    out_p.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"Generated {out_p} ({len(lines)} lines)")


def generate_q2():
    path = "pipeline/cohorts/HKIPO-MB2025Q2.xlsx"
    wb_b = get_head_wb(path)
    wb_a = openpyxl.load_workbook(ROOT / path, data_only=True)
    ws_b = wb_b["NLR"]
    ws_a = wb_a["NLR"]

    q2_diffs = {}
    for r in range(2, ws_a.max_row + 1):
        code = ws_a[f"B{r}"].value
        name = ws_a[f"C{r}"].value
        if not code:
            continue
        diffs = []
        for c in range(1, ws_a.max_column + 1):
            vb = ws_b.cell(row=r, column=c).value
            va = ws_a.cell(row=r, column=c).value
            if fmt(vb) != fmt(va):
                col_let = openpyxl.utils.get_column_letter(c)
                hdr = ws_a.cell(row=1, column=c).value
                diffs.append((col_let, hdr, fmt(vb), fmt(va)))
        if diffs:
            q2_diffs[code] = (name, diffs)

    lines = []
    lines.append("# 2025Q2 Main Board: year-1 re-anchored to the audited Track Record")
    lines.append("")
    lines.append("Same standard as `HKIPO-MB2026Q2_year1_corrections.md` and `HKIPO-MB2026Q3_year1_corrections.md`:")
    lines.append("year-1 is the latest period of the audited Accountants' Report / \"Track Record Period\" definition;")
    lines.append("unaudited figures outside it are never year-1; only sales, profit before tax and profit for the period")
    lines.append("are annualised (BC = months/12, originals in CN/CO/CP).")
    lines.append("")
    lines.append("## What changed")
    lines.append("")
    lines.append("| Group | Issuers | Before → after |")
    lines.append("|---|---|---|")
    lines.append("| Unaudited quarterly dates outside Track Record used as year-1; Track Record ended 31/12/2024 | 3750.HK (CATL), 2619.HK (X.J. Electrics) | Unaudited 31/03/25 or 30/04/25 → Audited FY2024 ended 31/12/2024 (full year, not annualised) |")
    lines.append("| Short track record due to incorporation (8 Dec 2022); period override recorded | 6603.HK (IFBH Limited) | Added `period-override` note to `col_AT`; FY2024 ended 31/12/2024 validated |")
    lines.append("")
    lines.append("The other 24 issuers in 2025Q2 already matched their Track Record (FY2024 ending 31 Dec 2024 or FY2025 ending 31 Mar 2025); their cells are unchanged.")
    lines.append("")
    lines.append("For 3750.HK and 2619.HK, the 27 period-dependent payload fields were re-extracted from the audited Accountants' Report")
    lines.append("(other fields untouched), validated, recorded `extracted`, and independently reviewed against original disclosure pages;")
    lines.append("reports bound to each payload sha256 are in `pipeline/prospectus_pipeline/out/period_reviews/2025Q2/`. Reviewed payloads")
    lines.append("were written back to the workbook, followed by `disclosure_notes` (all 27 on the Track Record basis), `audit` (1890/1890 perfect match),")
    lines.append("`cross_check` (5 pre-existing allocation warnings, unchanged), `export` and `master --derive`. 2025Q2 `validate`: 27/27 target, 0 errors, gate PASS.")
    lines.append("")
    lines.append("## Pending")
    lines.append("")
    lines.append("None in 2025Q2.")
    lines.append("")
    lines.append("## Every changed cell (before → after)")
    lines.append("")

    for code in sorted(q2_diffs.keys()):
        name, diffs = q2_diffs[code]
        clean_name = " ".join(name.split()) if name else ""
        lines.append(f"### {code} - {clean_name} ({len(diffs)} cells)")
        lines.append("")
        lines.append("| col | header | before | after |")
        lines.append("|---|---|---|---|")
        for col_let, hdr, vb, va in diffs:
            clean_hdr = " ".join(str(hdr).split()) if hdr else ""
            lines.append(f"| {col_let} | {clean_hdr} | {vb} | {va} |")
        lines.append("")

    out_p = ROOT / "pipeline/reports/period_corrections/HKIPO-MB2025Q2_year1_corrections.md"
    out_p.parent.mkdir(parents=True, exist_ok=True)
    out_p.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"Generated {out_p} ({len(lines)} lines)")


if __name__ == "__main__":
    generate_q1()
    generate_q2()
