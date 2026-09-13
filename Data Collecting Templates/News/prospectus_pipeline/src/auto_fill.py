#!/usr/bin/env python3
"""HKIPO 全流程外壳：准备 →（DSH 抽取）→ 校验 → 写回 Excel。

完成定义 = 工作簿对应列已有值，不是「JSON 落盘」。
抽取（Flash）由 DSH skill `hkipo-one-shot` 调 workflow；本脚本不调用 LLM。

用法：
  python3 prospectus_pipeline/auto_fill.py status
  python3 prospectus_pipeline/auto_fill.py prepare --only 6809.HK --allot
  python3 prospectus_pipeline/auto_fill.py next-batch --n 2
  python3 prospectus_pipeline/auto_fill.py write-ready          # JSON 齐但表空的，全部进表
  python3 prospectus_pipeline/auto_fill.py finish               # 只写回待写的；没有待写则 0
"""
from __future__ import annotations

import argparse
import json
import subprocess
import sys
from datetime import datetime, timezone, timedelta
from pathlib import Path

PIPE = Path(__file__).resolve().parent.parent
ROOT = PIPE.parent
EXT = PIPE / "out" / "extracted"
ALLOT_EXT = PIPE / "out" / "allot" / "extracted"
PACKETS = PIPE / "out" / "packets.json"
ALLOT_PACKETS = PIPE / "out" / "allot" / "packets.json"
PDF_DIR = PIPE / "data" / "pdf"
PKT_DIR = PIPE / "data" / "packets"
ALLOT_PKT_DIR = PIPE / "data" / "allot" / "packets"
BJ = timezone(timedelta(hours=8))


def digits(code: str) -> str:
    d = "".join(c for c in str(code) if c.isdigit())
    return f"{int(d):04d}"


def stem(code: str) -> str:
    return f"HKIPO-MB{digits(code)}"


def load_json(path: Path, default):
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def load_cfg():
    sys.path.insert(0, str(PIPE))
    from run import load_cfg as _load
    return _load()


def workbook_rows() -> list[dict]:
    from run import read_companies
    return read_companies(load_cfg())


def excel_written_map() -> dict[str, dict]:
    """code -> {prospectus: bool, allot: bool} based on live workbook cells."""
    import openpyxl
    cfg = load_cfg()
    book = cfg["_ws"] / cfg["workbook"]
    wb = openpyxl.load_workbook(book, read_only=True, data_only=True)
    ws = wb[cfg["sheet"]]
    headers = {}
    for col in range(1, ws.max_column + 1):
        v = ws.cell(1, col).value
        if v not in (None, ""):
            headers[" ".join(str(v).replace("\n", " ").split()).strip().lower()] = col
    # sentinel: Total (without option) must be filled for prospectus write-back
    l_col = headers.get("total (without option)")
    ck_col = headers.get("final cornerstone allocation (% of base offer)")
    code_col = headers.get("stock code")
    out: dict[str, dict] = {}
    if not code_col:
        wb.close()
        return out
    for r, row in enumerate(ws.iter_rows(min_row=cfg["data_start_row"], values_only=True),
                            cfg["data_start_row"]):
        code = row[code_col - 1] if code_col else None
        if code in (None, ""):
            continue
        key = str(code).strip()
        l_val = row[l_col - 1] if l_col else None
        ck_val = row[ck_col - 1] if ck_col else None
        out[key] = {
            "prospectus": l_val not in (None, ""),
            "allot": ck_val not in (None, ""),
            "row": r,
        }
    wb.close()
    return out


def json_ok(path: Path) -> bool:
    if not path.exists() or path.stat().st_size <= 200:
        return False
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
        return isinstance(value, dict) and set(value) == {"code", "fields"} and isinstance(value["fields"], dict)
    except (OSError, UnicodeError, json.JSONDecodeError):
        return False


def peak_now(now: datetime | None = None) -> bool:
    t = now or datetime.now(BJ)
    if t.weekday() >= 5:
        return False
    minutes = t.hour * 60 + t.minute
    return (9 * 60 <= minutes < 12 * 60) or (14 * 60 <= minutes < 18 * 60)


def inventory() -> list[dict]:
    cfg = load_cfg()
    companies = workbook_rows()
    written = excel_written_map()
    prospectus_packets = {r["code"]: r for r in load_json(PACKETS, [])}
    allot_packets = {r["code"]: r for r in load_json(ALLOT_PACKETS, [])}
    rows = []
    for co in companies:
        code = co["code"]
        s = stem(code)
        xl = written.get(code, {})
        rec = {
            "row": co["row"],
            "code": code,
            "name": co.get("name") or "",
            "pdf": (PDF_DIR / f"{s}.pdf").exists(),
            "packet": (PKT_DIR / f"{s}.md").exists(),
            "extracted": json_ok(EXT / f"{s}.json"),
            "excel_prospectus": bool(xl.get("prospectus")),
            "excel_allot": bool(xl.get("allot")),
            "allot_packet": (ALLOT_PKT_DIR / f"{s}.md").exists(),
            "allot_extracted": json_ok(ALLOT_EXT / f"{s}.json"),
            "prospectus_packet_rec": prospectus_packets.get(code),
            "allot_packet_rec": allot_packets.get(code),
            "prospectus_write_tracked": False,
            "allot_write_tracked": False,
        }
        # Once the repaired pipeline has written a row, the per-company manifest
        # becomes authoritative. A changed JSON must be reviewed and written again.
        sys.path.insert(0, str(PIPE / "src"))
        from state import file_hash, read_record
        for target, json_path, excel_key in (
            ("prospectus", EXT / f"{s}.json", "excel_prospectus"),
            ("allot", ALLOT_EXT / f"{s}.json", "excel_allot"),
        ):
            written_state = read_record(cfg, target, code, "written")
            if written_state and json_path.exists():
                rec[excel_key] = written_state.get("json_sha256") == file_hash(json_path)
                rec[target + "_write_tracked"] = True
        rec["need_prepare"] = not (rec["pdf"] and rec["packet"])
        rec["need_extract"] = rec["packet"] and not rec["extracted"]
        rec["need_write"] = rec["extracted"] and not rec["excel_prospectus"]
        rec["need_allot_prepare"] = not rec["allot_packet"]
        rec["need_allot_extract"] = rec["allot_packet"] and not rec["allot_extracted"]
        rec["need_allot_write"] = rec["allot_extracted"] and not rec["excel_allot"]
        rec["done"] = rec["excel_prospectus"]  # 招股书进表才算这家完成
        rows.append(rec)
    return rows


def print_status(rows: list[dict]) -> None:
    need_prep = [r for r in rows if r["need_prepare"]]
    need_ex = [r for r in rows if r["need_extract"]]
    need_write = [r for r in rows if r["need_write"]]
    need_allot_ex = [r for r in rows if r["need_allot_extract"]]
    need_allot_write = [r for r in rows if r["need_allot_write"]]
    in_excel = [r for r in rows if r["excel_prospectus"]]
    have_json = [r for r in rows if r["extracted"]]
    legacy = [r for r in rows if ((r["excel_prospectus"] and not r["prospectus_write_tracked"])
                                  or (r["excel_allot"] and not r["allot_write_tracked"]))]
    print(f"工作簿 {len(rows)} 家 | 北京 {datetime.now(BJ):%Y-%m-%d %H:%M} "
          f"| {'高峰（建议等空闲）' if peak_now() else '空闲（可抽）'}")
    print(f"  招股书进表 {len(in_excel)}/{len(rows)}  JSON {len(have_json)}/{len(rows)}")
    if legacy:
        print(f"  注意：{len(legacy)} 家为旧流程写入，暂无哈希绑定的 written 记录；"
              "不能据此证明当前 JSON 与 Excel 完全一致")
    print(f"  待准备 {len(need_prep)}  待抽取 {len(need_ex)}  "
          f"待写回Excel {len(need_write)}  配发待抽 {len(need_allot_ex)}  "
          f"配发待写回 {len(need_allot_write)}")
    if need_prep:
        print("待 prepare:", " ".join(r["code"] for r in need_prep))
    if need_ex:
        print("待招股书抽取:")
        for r in need_ex:
            print(f"  {r['code']:9s}  row {r['row']:<3d}  {r['name'][:60]}")
    if need_write:
        print("JSON已齐、尚未写进 Excel:")
        for r in need_write:
            print(f"  {r['code']:9s}  row {r['row']:<3d}  → write-ready")
    if need_allot_ex:
        print("待配发抽取:", " ".join(r["code"] for r in need_allot_ex))
    if need_allot_write:
        print("配发待写回:", " ".join(r["code"] for r in need_allot_write))
    if (not need_prep and not need_ex and not need_write and not need_allot_ex
            and not need_allot_write and not legacy):
        print("全流程完成：抽取 JSON 与 Excel 已对齐。")
    elif not need_prep and not need_ex and not need_write and not need_allot_ex and not need_allot_write:
        print("表格均已有值，但旧写入尚无可验证的版本绑定；状态为 legacy-untracked。")
    elif need_write and not need_ex:
        print("下一步：python3 prospectus_pipeline/auto_fill.py write-ready")


def run_py(args: list[str]) -> int:
    cmd = [sys.executable, str(PIPE / "run.py"), *args]
    print("+", " ".join(cmd), flush=True)
    return subprocess.call(cmd, cwd=ROOT)


def cmd_prepare(codes: list[str] | None, also_allot: bool) -> int:
    extra = ["--only", *codes] if codes else []
    rc = run_py(["all", *extra])
    if rc:
        return rc
    if also_allot:
        rc = run_py(["allot", *extra])
    return rc


def cmd_write_ready(codes: list[str] | None, target: str, rewrite: bool = False) -> int:
    """Validate then write JSON into Excel.

    Default: only firms whose Excel sentinel is still empty.
    --rewrite: write every firm that has JSON (overwrite cells).
    """
    rows = inventory()
    if target == "allot":
        ready = [r["code"] for r in rows if r["need_allot_write"] or (rewrite and r["allot_extracted"])]
        have = {r["code"] for r in rows if r["allot_extracted"]}
        pending_key = "need_allot_write"
    else:
        ready = [r["code"] for r in rows if r["need_write"] or (rewrite and r["extracted"])]
        have = {r["code"] for r in rows if r["extracted"]}
        pending_key = "need_write"
    if codes:
        wanted = set(codes)
        ready = [c for c in ready if c in wanted]
        missing = [c for c in codes if c not in have]
        if missing:
            print("没有 JSON，跳过写回：", " ".join(missing))
        already = [c for c in codes if c in have and c not in ready]
        if already and not rewrite:
            print("Excel 已有值，跳过（需要覆盖加 --rewrite）：", " ".join(already))
    if not ready:
        print("没有待写回 Excel 的公司（JSON 空或表里已有值）")
        return 0
    print("即将校验并写回 Excel：", " ".join(ready), flush=True)
    extra = ["--only", *ready]
    rc = run_py(["validate", "--target", target, *extra])
    if rc:
        print("validate 非零退出，中止写回", flush=True)
        return rc
    gate = PIPE / ("out/allot/validation.json" if target == "allot" else "out/validation.json")
    report = load_json(gate, {})
    blocked = {str(e.get("code")) for e in report.get("errors", [])}
    passed = [c for c in ready if c not in blocked]
    failed = [c for c in ready if c in blocked]
    if failed:
        print("闸门 ERROR，本批不写：", " ".join(failed), flush=True)
    if not passed:
        print("没有通过闸门的公司，Excel 未改", flush=True)
        return 1
    extra = ["--only", *passed]
    rc = run_py(["write", "--target", target, "--fill-missing", *extra])
    if rc:
        return rc
    ready = passed  # verify Excel only for firms we attempted to write
    after = {r["code"]: r for r in inventory()}
    still = [c for c in ready if after.get(c, {}).get(pending_key)]
    if still:
        print("写回后 Excel 仍无值：", " ".join(still), flush=True)
        return 1
    print("写回完成，Excel 已有值：", " ".join(ready), flush=True)
    return 0


def cmd_finish() -> int:
    """Write every JSON-ready empty Excel row. Does not call the LLM."""
    rows = inventory()
    rc = 0
    if any(r["need_write"] for r in rows):
        rc = cmd_write_ready(None, "prospectus") or rc
    if any(r["need_allot_write"] for r in rows):
        rc = cmd_write_ready(None, "allot") or rc
    print_status(inventory())
    return rc


def next_batch(n: int, target: str) -> list[dict]:
    rows = inventory()
    if target == "allot":
        cand = [r for r in rows if r["need_allot_extract"] and r["allot_packet_rec"]]
        key = "allot_packet_rec"
    else:
        cand = [r for r in rows if r["need_extract"] and r["prospectus_packet_rec"]]
        key = "prospectus_packet_rec"
    out = []
    for r in cand[:n]:
        rec = r[key]
        out.append({
            "code": rec["code"],
            "name": rec.get("name") or r["name"],
            "packet_path": rec["packet_path"],
            "out_path": rec["out_path"],
        })
    return out


def cmd_next_batch(n: int, target: str, force: bool, out: Path | None) -> int:
    if peak_now() and not force:
        print("现在是高峰时段（北京时间工作日 9:00–12:00、14:00–18:00）。"
              "空闲再抽才能压到约 ¥1/家。确认要现在跑就加 --force。", file=sys.stderr)
        payload = {"packets": [], "verify": True, "blocked": "peak"}
        text = json.dumps(payload, ensure_ascii=False, indent=2)
        print(text)
        if out:
            out.parent.mkdir(parents=True, exist_ok=True)
            out.write_text(text + "\n", encoding="utf-8")
        return 2
    batch = next_batch(n, target)
    payload = {
        "packets": batch,
        "verify": True,
        "out_dir": (
            "prospectus_pipeline/out/allot/extracted" if target == "allot"
            else "prospectus_pipeline/out/extracted"
        ),
    }
    text = json.dumps(payload, ensure_ascii=False, indent=2)
    print(text)
    if out:
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(text + "\n", encoding="utf-8")
    print(("# 没有待抽公司" if not batch else f"# {len(batch)} 家待抽 → workflow args"),
          file=sys.stderr)
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description="HKIPO 自动填表外壳（不含 LLM）")
    sub = ap.add_subparsers(dest="cmd", required=True)

    sub.add_parser("status", help="看还剩谁")
    p = sub.add_parser("prepare", help="find+download+prepare（¥0）")
    p.add_argument("--only", nargs="*")
    p.add_argument("--allot", action="store_true", help="同时准备配发公告")
    w = sub.add_parser("write-ready", help="把 JSON 校验后写入 Excel（默认只写表空的行）")
    w.add_argument("--only", nargs="*")
    w.add_argument("--target", choices=["prospectus", "allot"], default="prospectus")
    w.add_argument("--rewrite", action="store_true", help="覆盖 Excel 里已有值")
    sub.add_parser("finish", help="把所有「JSON 齐但表空」的公司写进 Excel")
    n = sub.add_parser("next-batch", help="打印下一批 workflow args")
    n.add_argument("--n", type=int, default=2)
    n.add_argument("--target", choices=["prospectus", "allot"], default="prospectus")
    n.add_argument("--force", action="store_true", help="高峰也打印批次")
    n.add_argument("--out", type=Path, default=PIPE / "out" / "next_batch.json")

    args = ap.parse_args()
    if args.cmd == "status":
        print_status(inventory())
        return 0
    if args.cmd == "prepare":
        return cmd_prepare(args.only, args.allot)
    if args.cmd == "write-ready":
        return cmd_write_ready(args.only, args.target, rewrite=args.rewrite)
    if args.cmd == "finish":
        return cmd_finish()
    if args.cmd == "next-batch":
        return cmd_next_batch(args.n, args.target, args.force, args.out)
    return 2


if __name__ == "__main__":
    raise SystemExit(main())
