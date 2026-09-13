#!/usr/bin/env python3
"""腾讯港股日线 → DD（恒指 20 日回报）+ DH–DM（上市首日 OHLC/量/额）。

接口：https://web.ifzq.gtimg.cn/appstock/app/hkfqkline/get
  param=hk06082,day,<from>,<to>,<n>,qfq
  恒指代码 hkHSI。

K 线字段顺序（实测）：date, open, close, high, low, volume, {}, ?, turnover_万元, ...
口径：
  DD  t = 招股书日前最后一个恒指交易日，P(t)/P(t−20)−1，用收盘。
  DH–DK 上市日（或其后第一个有 bar 的交易日）的 收/开/高/低。
  DL 成交量（股）；DM 成交额（HK$）= 接口第 9 列 × 10,000。

只赋值，不改样式。缓存落到 data/market/。
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import shutil
import time
import urllib.request
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2] if Path(__file__).resolve().parent.name == "external" else Path(__file__).resolve().parent
WS = ROOT.parent
BOOK = WS / "HKIPO-MB2026Q1.xlsx"
CACHE = ROOT / "data" / "market"
SHEET = "NLR"
UA = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36"

HEADERS = {
    "DD": "HSI return over 20 trading days before prospectus (%)",
    "DH": "First trading day closing price (HK$)",
    "DI": "First trading day opening price (HK$)",
    "DJ": "First trading day high (HK$)",
    "DK": "First trading day low (HK$)",
    "DL": "First trading day volume (shares)",
    "DM": "First trading day turnover (HK$)",
}


def norm(s) -> str:
    return " ".join(str(s or "").replace("\n", " ").split()).strip().lower()


def to_date(v) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def hk_symbol(code: str) -> str:
    d = "".join(ch for ch in str(code) if ch.isdigit())
    return f"hk{int(d):05d}"


def fetch(symbol: str, frm: str, to: str, n: int = 120, retries: int = 4) -> list[list]:
    url = ("https://web.ifzq.gtimg.cn/appstock/app/hkfqkline/get"
           f"?param={symbol},day,{frm},{to},{n},qfq")
    last = None
    for i in range(retries):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": UA})
            with urllib.request.urlopen(req, timeout=25) as r:
                payload = json.loads(r.read().decode())
            if payload.get("code") != 0:
                raise RuntimeError(f"{symbol} code={payload.get('code')} msg={payload.get('msg')}")
            data = (payload.get("data") or {}).get(symbol) or {}
            rows = data.get("day") or data.get("qfqday") or []
            if not rows:
                raise RuntimeError(f"{symbol} empty bars")
            return rows
        except Exception as exc:  # noqa: BLE001
            last = exc
            time.sleep(0.6 * (i + 1))
    raise RuntimeError(f"{symbol} 拉取失败：{last}")


def parse_bar(row: list) -> dict:
    # date, open, close, high, low, volume, {}, ?, turnover_万元
    turnover = None
    if len(row) > 8 and row[8] not in (None, "", "{}", {}):
        try:
            turnover = float(row[8]) * 10_000
        except (TypeError, ValueError):
            turnover = None
    return {
        "date": dt.datetime.strptime(str(row[0])[:10], "%Y-%m-%d").date(),
        "open": float(row[1]),
        "close": float(row[2]),
        "high": float(row[3]),
        "low": float(row[4]),
        "volume": float(row[5]),
        "turnover": turnover,
    }


def last_before(bars: list[dict], d: dt.date) -> int | None:
    idx = None
    for i, b in enumerate(bars):
        if b["date"] < d:
            idx = i
        else:
            break
    return idx


def first_on_or_after(bars: list[dict], d: dt.date) -> dict | None:
    for b in bars:
        if b["date"] >= d:
            return b
    return None


def companies(ws) -> list[tuple[int, str, dt.date, dt.date]]:
    out = []
    for r in range(2, ws.max_row + 1):
        code = ws[f"B{r}"].value
        if code in (None, ""):
            continue
        pd, ld = to_date(ws[f"D{r}"].value), to_date(ws[f"E{r}"].value)
        if not pd or not ld:
            continue
        out.append((r, str(code).strip(), pd, ld))
    return out


def resolve_cols(ws) -> dict[str, int]:
    col_of = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        if v in (None, ""):
            continue
        n = norm(v)
        for key, header in HEADERS.items():
            if n == norm(header):
                col_of[key] = c
    missing = [k for k in HEADERS if k not in col_of]
    if missing:
        raise SystemExit(f"工作簿缺列 {missing}")
    return col_of


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--book", default=str(BOOK))
    ap.add_argument("--only", nargs="*", default=None, help="只处理指定股票代码")
    args = ap.parse_args()
    book = Path(args.book)
    CACHE.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(book)
    ws = wb[SHEET]
    col_of = resolve_cols(ws)
    rows = companies(ws)
    if args.only:
        rows = [r for r in rows if r[1] in args.only]
    print(f"工作簿 {len(rows)} 家")

    # 恒指：覆盖最早招股书日前 40 个交易日
    hsi_from = (min(p for _, _, p, _ in rows) - dt.timedelta(days=80)).isoformat()
    hsi_to = (max(p for _, _, p, _ in rows) + dt.timedelta(days=5)).isoformat()
    hsi_raw = fetch("hkHSI", hsi_from, hsi_to, n=200)
    (CACHE / "hsi.json").write_text(json.dumps(hsi_raw, ensure_ascii=False), encoding="utf-8")
    hsi = [parse_bar(x) for x in hsi_raw]
    print(f"HSI {len(hsi)} 条 {hsi[0]['date']} ~ {hsi[-1]['date']}")

    results = []
    for r, code, pd, ld in rows:
        rec = {"code": code, "row": r, "prospectus": str(pd), "listing": str(ld)}
        i = last_before(hsi, pd)
        if i is None or i < 20:
            rec["DD"] = None
            rec["DD_note"] = f"恒指序列不够：i={i}"
        else:
            rec["DD"] = hsi[i]["close"] / hsi[i - 20]["close"] - 1
            rec["DD_t"] = str(hsi[i]["date"])
            rec["DD_t20"] = str(hsi[i - 20]["date"])
        sym = hk_symbol(code)
        frm = (ld - dt.timedelta(days=3)).isoformat()
        to = (ld + dt.timedelta(days=10)).isoformat()
        try:
            raw = fetch(sym, frm, to, n=20)
            (CACHE / f"{sym}.json").write_text(json.dumps(raw, ensure_ascii=False), encoding="utf-8")
            bars = [parse_bar(x) for x in raw]
            bar = first_on_or_after(bars, ld)
            if bar is None:
                rec["first"] = None
                rec["first_note"] = "无上市日及之后的 K 线"
            else:
                rec["first"] = bar
                if bar["date"] != ld:
                    rec["first_note"] = f"上市日 {ld} 无 K 线，改用 {bar['date']}"
        except Exception as exc:  # noqa: BLE001
            rec["first"] = None
            rec["first_note"] = str(exc)
        results.append(rec)
        time.sleep(0.25)

    (CACHE / "computed.json").write_text(json.dumps(results, ensure_ascii=False, indent=2, default=str) + "\n",
                                         encoding="utf-8")

    print(f"\n{'code':9s} {'招股书':12s} {'DD%':8s} {'上市日':12s} {'收':>8s} {'开':>8s} {'量':>14s}")
    n_ok = 0
    for rec in results:
        first = rec.get("first") or {}
        dd = rec.get("DD")
        print(f"{rec['code']:9s} {rec['prospectus']:12s} "
              f"{(f'{dd*100:7.3f}' if dd is not None else '   NaN ')} "
              f"{rec['listing']:12s} "
              f"{first.get('close', float('nan')):8.3f} {first.get('open', float('nan')):8.3f} "
              f"{first.get('volume', float('nan')):14,.0f}"
              f"{'  ' + rec.get('first_note','') if rec.get('first_note') else ''}"
              f"{'  ' + rec.get('DD_note','') if rec.get('DD_note') else ''}")
        if dd is not None and first:
            n_ok += 1
    print(f"\n完整 {n_ok}/{len(results)}")

    if args.dry_run:
        wb.close()
        print("--dry-run：未写回")
        return 0

    wb.close()
    backup_dir = book.parent / "backups" / "excel_snapshots"

    backup_dir.mkdir(parents=True, exist_ok=True)

    backup = backup_dir / (f"{book.stem}.backup-before-tencent-"
                            f"{dt.datetime.now():%Y%m%d-%H%M%S}.xlsx")
    shutil.copy2(book, backup)
    wb = openpyxl.load_workbook(book)
    ws = wb[SHEET]
    col_of = resolve_cols(ws)
    for rec in results:
        r = rec["row"]
        dd = rec.get("DD")
        ws.cell(r, col_of["DD"]).value = round(dd, 8) if dd is not None else "NaN"
        first = rec.get("first")
        if not first:
            for k in ("DH", "DI", "DJ", "DK", "DL", "DM"):
                ws.cell(r, col_of[k]).value = "NaN"
            continue
        ws.cell(r, col_of["DH"]).value = first["close"]
        ws.cell(r, col_of["DI"]).value = first["open"]
        ws.cell(r, col_of["DJ"]).value = first["high"]
        ws.cell(r, col_of["DK"]).value = first["low"]
        ws.cell(r, col_of["DL"]).value = int(round(first["volume"]))
        ws.cell(r, col_of["DM"]).value = (round(first["turnover"], 2)
                                          if first.get("turnover") is not None else "NaN")
    tmp = book.with_suffix(".saving.xlsx")
    wb.save(tmp)
    wb.close()
    tmp.replace(book)
    print(f"\n已写回 DD, DH–DM（{n_ok} 家完整）\n备份：{backup.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
