"""Create an English STE-style report and standalone LaTeX from verified inputs."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
import re

import numpy as np
import pandas as pd
import yaml
from openpyxl import load_workbook

from research_inputs import C, MASTER, ROOT, load_panel, select_2026

OUT = ROOT / 'analysis/out/retail_2026_report'
BRIEF = ROOT / 'analysis/out/retail_evidence_brief'
STEM = ROOT / 'docs/reports/HK_IPO_2026_COMPREHENSIVE_REPORT'


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def tex(value):
    replacements = {'\\':r'\textbackslash{}', '&':r'\&', '%':r'\%', '$':r'\$', '#':r'\#',
                    '_':r'\_', '{':r'\{', '}':r'\}', '~':r'\textasciitilde{}', '^':r'\textasciicircum{}'}
    return ''.join(replacements.get(c,c) for c in str(value))


def latex_table(block):
    n=len(block['headers'])
    weights=block['widths'] or [1]*n
    # Available width excludes the spaces between columns. Outer spaces are absent.
    available=17.0-2*(4/28.45274)*(n-1)
    widths=[available*w/sum(weights) for w in weights]
    cols='@{}'+''.join((r'>{\raggedright\arraybackslash}' if i==0 else r'>{\raggedleft\arraybackslash}')+f'p{{{w:.4f}cm}}' for i,w in enumerate(widths))+'@{}'
    if block.get('text_columns'):
        cols='@{}'+''.join(r'>{\raggedright\arraybackslash}'+f'p{{{w:.4f}cm}}' for w in widths)+'@{}'
    header=' & '.join(r'\textbf{'+tex(x)+'}' for x in block['headers'])+r' \\'
    result=[r'\par\addvspace{5pt}',r'\noindent\textbf{'+tex(block['title'])+'}',r'\par\nobreak\vspace{2pt}',r'{\small',r'\setlength{\tabcolsep}{4pt}',r'\renewcommand{\arraystretch}{1.08}',r'\setlength{\LTpre}{0pt}\setlength{\LTpost}{3pt}',r'\begin{longtable}{'+cols+'}',r'\toprule',
            header,r'\midrule\endfirsthead',r'\toprule',header,r'\midrule\endhead',r'\bottomrule\endfoot']
    result+=[' & '.join(tex(x) for x in row)+r' \\' for row in block['rows']]
    result += [r'\end{longtable}',r'}\par\addvspace{3pt}']
    return '\n'.join(result)


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    provenance=json.loads((BRIEF/'run_manifest.json').read_text())
    for rel,expected in provenance['inputs'].items(): assert sha(ROOT/rel)==expected,rel
    for name,expected in provenance['outputs'].items(): assert sha(BRIEF/name)==expected,name
    panel=select_2026(load_panel())
    one=pd.read_csv(BRIEF/'one_lot_issuer_results.csv')
    assert len(panel)==len(one)==113 and set(panel['Stock Code'])==set(one.code)
    panel=panel.set_index('Stock Code').loc[one.code].reset_index()
    registry_path=ROOT/'pipeline/registry/HKIPO_Variable_Registry.yaml'
    registry=yaml.safe_load(registry_path.read_text())
    workbooks=[ROOT/f'pipeline/cohorts/HKIPO-MB2026Q{q}.xlsx' for q in [1,2,3]]
    excel_records={}
    for path in workbooks:
        wb=load_workbook(path,read_only=True,data_only=True)
        ws=wb['NLR']; headers=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
        normalize=lambda value:' '.join(str(value or '').split()).lower()
        index_by_header={normalize(value):i for i,value in enumerate(headers) if value}
        # Two legacy workbook labels differ from their registered export names.
        aliases={
            'Company Name at time of listing':'Company Name at time of listing (exclude Chapter 20 cases)',
            'Comments / Annualization factor':'Comments (nearest sales& profit adjustment factor - original data duration in year, eg. 6 month pls input 0.5)'}
        field_indexes={}
        for variable in registry['variables']:
            field=variable['header']; key=normalize(field)
            if key not in index_by_header:key=normalize(aliases.get(field,field))
            assert key in index_by_header,(path.name,field)
            field_indexes[field]=index_by_header[key]
        code_index=field_indexes['Stock Code']
        for row in ws.iter_rows(min_row=2,values_only=True):
            code=row[code_index]
            if code in set(one.code):
                assert code not in excel_records,code
                excel_records[code]={field:row[index] for field,index in field_indexes.items()}
        wb.close()
    assert set(excel_records)==set(one.code)
    def present(value):
        return value is not None and str(value).strip().lower() not in {'','na','nan','n/a','none'}
    coverage=[]
    for variable in registry['variables']:
        field=variable['header']
        assert field in panel,field
        coverage.append({'field':field,'workbook_letter':variable['letter'],'source_layer':variable['layer'],
                         'registered_unit':variable.get('unit'),'definition_from_registry':variable.get('description_zh'),
                         'workbook_nonmissing_n':sum(present(row.get(field)) for row in excel_records.values()),
                         'analysis_nonmissing_n':int(panel[field].notna().sum()),'sample_n':113})
    assert len(coverage)==202
    pd.DataFrame(coverage).to_csv(OUT/'workbook_field_coverage.csv',index=False)
    # The exported headline return has up to five decimal places in decimal units.
    np.testing.assert_allclose(panel.ir,one.initial_return,atol=5.1e-6,rtol=0)
    assert panel.listing_date.max()==pd.Timestamp('2026-09-30')
    np.testing.assert_allclose(one.application_return,one.allocation_rate*one.initial_return)
    summaries=pd.read_csv(BRIEF/'sample_sensitivity.csv')
    base=summaries.set_index('sample').loc['corrected_all_113']
    quarters=pd.read_csv(BRIEF/'quarter_results.csv')
    fees=pd.read_csv(BRIEF/'fee_sensitivity.csv')
    winners=pd.read_csv(BRIEF/'winner_sensitivity.csv')
    pages=[]
    table_count=0
    def page(title):
        p={'title':title,'blocks':[]};pages.append(p);return p['blocks']
    def paragraph(b,s):b.append({'type':'paragraph','text':s})
    def heading(b,s):b.append({'type':'heading','text':s})
    def table(b,key,title,headers,rows,widths=None,text_columns=False):
        nonlocal table_count
        table_count+=1
        title=re.sub(r'^Table \d+',f'Table {table_count}',title)
        pd.DataFrame(rows,columns=headers).to_csv(OUT/f'{key}.csv',index=False)
        b.append({'type':'table','title':title,'headers':headers,'rows':rows,'widths':widths,'text_columns':text_columns})
    pct=lambda x:f'{100*x:.2f}%'
    num=lambda x:f'{x:,.2f}'

    b=page('Hong Kong IPOs in 2026')
    paragraph(b,'Data description and retail application returns. Report for mentor discussion. Prepared on 3 October 2026. Revised on 4 October 2026. Data cutoff: 30 September 2026.')
    heading(b,'Main result')
    paragraph(b,'We study 113 ordinary Main Board IPOs in the project database. These IPOs listed from 2 January to 30 September 2026. The mean first-day return is 53.15%. The mean gross application return for one lot is 2.08%. The difference is 51.07 percentage points.')
    paragraph(b,'The allocation rate and the first-day return have a negative covariance. This covariance reduces the application return by 3.65 percentage points. It offsets 63.6% of the product of the two means. This result describes the sample. It does not show a causal effect.')
    table(b,'overview','Table 1  Sample scope and main results',['Measure','Result'],[
        ['Listing dates','2 January to 30 September 2026'],['Sample size','113 IPOs: Q1 38, Q2 45, Q3 30'],
        ['Total base offer proceeds',f'HK${panel.base_proceeds.sum()/1e9:,.2f} billion'],
        ['First-day return: mean / median',f'{pct(base.mean_initial_return)} / {pct(base.median_initial_return)}'],
        ['Gross application return: mean / median',f'{pct(base.mean_application_return)} / {pct(base.median_application_return)}'],
        ['Expected gross profit: mean / median',f'HK${num(base.mean_expected_gross_profit_hkd)} / HK${num(base.median_expected_gross_profit_hkd)}'],
        ['Mean expected allocation rate for one lot',pct(one.allocation_rate.mean())]], [10,7])
    heading(b,'How to read this report')
    paragraph(b,'The sample covers Q1 to Q3. It does not cover the full year. Each IPO has equal weight. We use the first-day closing price. The base result excludes all costs.')
    paragraph(b,'An application return is a calculated expectation. It is not an observed account return. HK$88 is an assumed fee for a sensitivity check. We do not use a portfolio simulation.')
    paragraph(b,'The database contains the research sample. This report does not certify coverage of every type of new listing on the exchange.')

    b=page('The 2026 sample')
    monthly=[]
    for month,ids in panel.groupby(panel.listing_date.dt.strftime('%Y-%m')):
        current=one[one.code.isin(ids['Stock Code'])]
        monthly.append([month,len(ids),f'{ids.base_proceeds.sum()/1e9:.2f}',pct(current.initial_return.mean()),pct(current.initial_return.median())])
    table(b,'monthly','Table 2  Listings and base offer proceeds by month',['Month','N','Proceeds HK$bn','Mean return','Median return'],monthly,[2.2,.8,3.2,3.4,3.4])
    paragraph(b,'Base offer proceeds equal the offer price times the base offer quantity. They exclude the over-allotment option. They are not proceeds after issue costs. We do not add share quantities to HDR quantities.')
    names={'Conventional':'Conventional','A+H (19A)':'A+H','18C specialist tech':'Chapter 18C','18A biotech':'Chapter 18A'}
    rows=[]
    for route,name in names.items():
        sub=panel[panel.route.eq(route)]; current=one[one.code.isin(sub['Stock Code'])]
        rows.append([name,len(sub),f'{100*len(sub)/113:.1f}%',pct(current.initial_return.mean()),pct(current.application_return.mean())])
    table(b,'routes','Table 3  Listing route groups',['Route','N','Sample share','First-day mean','Application mean'],rows,[3.3,.8,2.5,3.3,3.3])
    paragraph(b,'The route groups do not overlap. The project assigns Chapter 18A and Chapter 18C before it assigns A+H. There are 38 issuers with an A+H flag. Two also have a Chapter 18C flag. Thus, the separate A+H route group has 36 issuers.')

    b=page('Industry and issuer characteristics')
    taxonomy_path=ROOT/'pipeline/prospectus_pipeline/data/manual/hsics.json'
    taxonomy={x['code']:x['industry'] for x in json.loads(taxonomy_path.read_text())}
    english={'資訊科技業':'Information technology','工業':'Industrials','醫療保健業':'Healthcare',
             '非必需性消費':'Consumer discretionary','原材料業':'Materials','必需性消費':'Consumer staples',
             '地產建築業':'Property and construction','金融業':'Financials'}
    sector=panel['Industry classification code'].map(lambda x:str(int(x)).zfill(6)).map(taxonomy).map(english)
    assert sector.notna().all()
    industry_rows=[]
    for name,count in sector.value_counts().items():
        codes=panel.loc[sector.eq(name),'Stock Code']; current=one[one.code.isin(codes)]
        industry_rows.append([name,int(count),f'{100*count/113:.1f}%',pct(current.initial_return.mean()),pct(current.application_return.mean())])
    table(b,'industry','Table 4  Industry groups',['Industry','N','Sample share','First-day mean','Application mean'],industry_rows,[5.5,.7,2.5,3.0,3.3])
    paragraph(b,'Industry groups use the stored HSICS 2026 code and the project classification table. All 113 codes have a group. Information technology accounts for 49 IPOs. Industrials account for 24 IPOs. Some groups have only one IPO. Their means do not show a general industry effect.')
    table(b,'issuer_flags','Table 5  Issuer characteristics',['Characteristic','Yes','No','Unknown','Yes share of known'],[
        ['A+H flag',38,75,0,'33.63%'],['Weighted voting rights',4,109,0,'3.54%'],
        ['VC/PE backing',87,19,7,'82.08%'],['Prior-year loss',51,61,1,'45.54%']],[6.5,1.5,1.5,2.0,5.5])
    paragraph(b,'The flags can overlap. Unknown flags remain unknown. We do not replace them with zero. The route groups on the previous page use a separate classification rule.')
    paragraph(b,'Financial periods and currencies differ across issuers. We do not add prior-year profit values. The loss flag uses each issuer\'s recorded prior-year profit. It does not describe a common calendar-year accounting period.')

    b=page('Descriptive statistics')
    measures={
        'Base proceeds (HK$m)':panel.base_proceeds/1e6,
        'Public subscription (times)':panel[C['sub']],
        'Public applicants (count)':panel['Public applicants'],
        'Firm age (years)':panel[C['age']],
        'Cornerstone allocation (%)':panel[C['corner']]*100,
        'Unrestricted public holding (%)':panel[C['float']]*100,
        'One-lot principal (HK$)':one.application_principal_hkd,
        'One-lot allocation rate (%)':one.allocation_rate*100,
        'First-day return (%)':one.initial_return*100,
        'Gross application return (%)':one.application_return*100,
        'Expected gross profit (HK$)':one.expected_gross_profit_hkd,
        'Liabilities / assets (%)':panel['leverage_y1']*100,
        'Profit / assets (%)':panel['roa_y1']*100,
        'Sales growth (%)':panel['sales_growth_y1']*100,
        'Pre-IPO institutional holding (%)':panel['Pre-IPO institutional shareholding (%)']*100,
        'Controller economic interest (%)':panel['Controller economic interest at listing (%)']*100,
        'Proceeds for debt repayment (%)':panel['Debt repayment (% of planned net IPO proceeds)']*100}
    raw=[];rows=[]
    for name,s in measures.items():
        s=pd.to_numeric(s,errors='coerce').dropna()
        rows.append([name,len(s),num(s.mean()),num(s.std(ddof=1)),num(s.quantile(.25)),num(s.median()),num(s.quantile(.75))])
        raw.append({'variable':name,'n':len(s),'mean':s.mean(),'sd':s.std(ddof=1),'p25':s.quantile(.25),'median':s.median(),'p75':s.quantile(.75),'min':s.min(),'max':s.max()})
    pd.DataFrame(raw).to_csv(OUT/'descriptive_statistics_raw.csv',index=False)
    table(b,'descriptive_statistics','Table 4  Raw data without winsorization',['Variable','N','Mean','SD','P25','Median','P75'],rows,[4.5,.75,2.35,2.35,2.35,2.35,2.35])
    paragraph(b,'N is the number of known values. SD uses the N minus one denominator. P25 and P75 are the 25th and 75th percentiles. The underlying CSV also gives minimum and maximum values.')
    positive=int(one.initial_return.gt(0).sum());zero=int(one.initial_return.eq(0).sum());negative=int(one.initial_return.lt(0).sum())
    paragraph(b,f'First-day prices increased for {positive} IPOs, did not change for {zero}, and decreased for {negative}. The share with a negative first-day return is {100*negative/113:.2f}%. The return range is {pct(one.initial_return.min())} to {pct(one.initial_return.max())}.')
    paragraph(b,'The mean first-day return is much higher than the median. Large positive returns affect the mean. Public subscription also has a long upper tail. Its mean is 1,985.19 times, compared with a median of 1,073.37 times.')
    paragraph(b,'The one-lot allocation rate uses the application tier for one lot. It differs from the average rate for all public applications. The sample includes one HDR issuer, 6228. We calculate its quantities in HDR units.')
    paragraph(b,'A median expected profit is a median across IPO-specific expectations. It is not the median profit of actual investor accounts. Missing values remain missing.')

    b=page('What the collected Excel data can do')
    paragraph(b,'The three workbooks contain 202 registered fields for 113 issuers. Eleven fields come from official listing records. Seventy-seven fields come from prospectuses. The other 114 fields contain allotment, market, and external data. A separate CSV lists coverage for every workbook field.')
    paragraph(b,'The tables below explain the meaning and possible use of selected fields. N counts known analysis values out of 113. Each field has its own count. A filled cell is not proof of a correct source. Some missing cells are not applicable.')
    def counts(pairs):
        return '. '.join(f'{label}: {int(panel[field].notna().sum())}' for label,field in pairs)+'.'
    company_groups=[
        ['Identity and dates','Issuer names, listing dates, and industry identify the observation.',counts([('Dates','Date of Listing (dd/mm/yy)'),('Industry','Industry classification code')]),'Build cohorts. Join prices and event data. Control for time and industry.'],
        ['Offer price and structure','Price limits and offer quantities describe pricing and capital raised.',counts([('Upper price','Maximum Offer Price'),('Lower price','Minimum Offer Price'),('Final price',C['offer'])]),'Study price revision and offer size. Treat fixed-price offers separately.'],
        ['Financial accounts','Assets, liabilities, revenue, and profit describe size, debt, growth, and profitability.',counts([('Assets','total assets in year-1'),('Revenue','Net sales in year-1'),('Profit','Profit for the year in year-1')]),'Construct liability ratios, profit ratios, and sales growth. Compare issuer risk with pricing.'],
        ['Operations and cash','Research costs, customer concentration, cash flow, and capital spending describe operating risk.',counts([('R&D','R&D expensed in year-1 (before annualization)'),('Customers','Top 5 customers (% of year-1 revenue)'),('Cash flow','Operating cash flow in year-1 (before annualization)')]),'Study innovation, cash needs, and customer risk. Align the financial periods first.'],
        ['Pre-IPO investors','Backing, ownership, holding time, and board seats describe institutional involvement.',counts([('VC/PE',C['vc']),('Holding','Pre-IPO institutional shareholding (%)'),('Board seat','Pre-IPO investor board seat (1=yes; 0=no)')]),'Compare pricing and later returns across backing groups. Separate investor selection from a causal effect.'],
        ['Control and governance','Controller ownership, voting rights, and WVR describe control rights.',counts([('Economic rights','Controller economic interest at listing (%)'),('Voting rights','Controller voting rights at listing (%)'),('WVR','WVR flag')]),'Measure ownership and voting differences. Study governance and investor protection.'],
        ['Sponsors and issue costs','Sponsors, accountants, and commissions describe intermediaries and issue costs.',counts([('Sponsor','Sponsor(s)'),('Lead sponsor','Lead sponsor name'),('Base fee','Underwriting base commission rate (%)')]),'Compare sponsor groups and underwriting costs. Treat reputation as an observed association.'],
        ['Cornerstone investors','Allocation and investor names describe demand commitments. Unlock dates describe later supply events.',counts([('Allocation',C['corner']),('Names','Cornerstone investor names'),('Unlock date','Earliest cornerstone unlock date (dd/mm/yy)')]),'Study commitments, available supply, and pricing. Separate final allotment from pre-pricing information.'],
    ]
    table(b,'company_data_uses','Table 7  Company and offer data',['Data group','What it means','Known N','Possible research use'],company_groups,[2.7,4.5,3.2,6.6],text_columns=True)
    paragraph(b,'Financial ratios use the recorded periods and definitions. Liabilities can exceed assets. The liability ratio can exceed 100%. Raw sales growth has large upper-tail values. Check period lengths, currencies, and units before comparisons. Do not mix annualized sales with unadjusted research costs.')
    paragraph(b,'Underwriting fees are issuer costs. They are not the assumed applicant fee. A sponsor group can reflect issuer selection. These data can show an association without showing a sponsor effect.')
    market_groups=[
        ['Public demand and allocation','Applicants, applied quantity, and final tranches describe demand and supply.',counts([('Applicants','Public applicants'),('Applied shares','Public valid applied shares'),('Public tranche','Final public offer shares')]),'Calculate macro allocation rates. Link exact tier rules to application returns. Test the allocation-return relation.'],
        ['Public holding and free float','Public holding includes eligible holdings. Unrestricted holding measures available supply under the stored definition.',counts([('Public holding','Public shareholding at listing (%)'),('Unrestricted holding',C['float'])]),'Study liquidity and price pressure. Keep the denominator and locked holdings explicit.'],
        ['Market conditions','HSI returns, HIBOR, bank balances, and recent IPO counts describe conditions before the offer.',counts([('HSI','HSI return over 20 trading days before prospectus (%)'),('HIBOR','1-month HIBOR before prospectus (%)'),('IPO count','HK ordinary IPO count in 90 calendar days before prospectus')]),'Control for market returns, funding conditions, and offer waves. Use only information known at the test date.'],
        ['Regimes and offer rules','Mechanism and regime fields classify the offer and settlement setting.',counts([('Mechanism','Offer mechanism'),('Rules','Applicable IPO rules / transition basis'),('FINI','FINI digital settlement regime')]),'Compare mechanisms when they differ. A constant regime flag cannot identify an effect within 2026.'],
        ['First-day trading','Opening and closing prices, volume, and turnover describe listing-day trading.',counts([('Close','First trading day closing price (HK$)'),('Volume','First trading day volume (shares)'),('Turnover','First trading day turnover (HK$)')]),'Study price discovery and trading intensity. The flipping ratio is a volume proxy, not investor-level selling.'],
        ['Later returns and liquidity','Later prices and index benchmarks separate the IPO jump from later trading returns.',counts([('One month','1-month BHR from Day-1 close (%)'),('Three months','3-month BHR from Day-1 close (%)'),('Six months','6-month BHR from Day-1 close (%)')]),'Calculate later returns from the first-day close. Use mature horizons and matched benchmark coverage.'],
        ['Stabilization and greenshoe','Purchases, option exercise, and end dates describe price support after listing.',counts([('Purchases','Stabilization purchases occurred'),('End date','Stabilization period end date'),('Exercise rate','Greenshoe exercise rate (%)')]),'Study support and prices after support ends. These are later outcomes, not ex-ante pricing controls.'],
        ['Lockup events','Contract dates and event returns describe changes around investor unlock dates.',counts([('Controller date','Controlling shareholder 6-month disposal lockup expiry date'),('CAR 5','Cornerstone unlock CAR [-5, +5] (%)'),('CAR 20','Cornerstone unlock CAR [-20, +20] (%)')]),'Run an event study of price and volume changes. Require contract evidence and complete event windows.'],
    ]
    table(b,'market_data_uses','Table 8  Demand and aftermarket data',['Data group','What it means','Known N','Possible research use'],market_groups,[2.7,4.5,3.2,6.6],text_columns=True)
    paragraph(b,'The three-month raw return has 82 observations. Its HSI wealth relative has only 65 observations. Joint tests must use the matched sample. One-year and three-year reserved returns have no observations. Future or unknown returns must not become zero.')
    paragraph(b,'HIBOR is a market funding measure. It is not an observed borrowing cost for an applicant. Later stabilization outcomes cannot explain pricing with information known before the offer.')
    heading(b,'Research that the current data support')
    paragraph(b,'First, compare application returns across demand, offer size, and public-tranche groups. Second, study price revision with issuer finances, institutional backing, and cornerstone commitments. Third, test later returns and unlock events on separate mature samples. Use a common sample within each model. Report associations until a credible design identifies causality.')

    b=page('From first-day returns to application returns')
    paragraph(b,'Let the offer price be P0 and the first-day closing price be P1. Let q be the requested quantity. Let E[A] be the expected allotted quantity. Let a be the allocation rate, and let r be the first-day return.')
    b.append({'type':'equation','tex':r'r_i=\frac{P_{1i}}{P_{0i}}-1,\quad a_i=\frac{E[A_i]}{q_i},\quad B_i=q_iP_{0i},\quad G_i=E[A_i](P_{1i}-P_{0i}),\quad y_i=\frac{G_i}{B_i}=a_ir_i.'})
    paragraph(b,'B is the application principal. G is expected gross profit. y is the gross application return. The expected allotted quantity includes guaranteed quantities and additional ballots.')
    table(b,'decomposition','Table 5  The allocation and return decomposition',['Measure','Result'],[
        ['Mean first-day return',pct(base.mean_initial_return)],['Mean gross application return',pct(base.mean_application_return)],
        ['Difference',f'{base.return_gap_pp:.2f} pp'],['Mean allocation rate times mean first-day return',pct(base.product_of_means)],
        ['Allocation-return covariance',f'{base.covariance_contribution_pp:.2f} pp'],['Share of the product offset by negative covariance',f'{base.covariance_reduction_pct:.1f}%']], [12,5])
    b.append({'type':'equation','tex':r'\overline{ar}=\bar a\,\bar r+\operatorname{Cov}_{N}(a,r),\qquad \operatorname{Cov}_{N}(a,r)=\frac{1}{N}\sum_{i=1}^{N}(a_i-\bar a)(r_i-\bar r).'})
    paragraph(b,'The covariance uses N as the denominator. The identity is exact before rounding. Thus, 5.73% minus 3.65 percentage points gives 2.08%. The product of means is an algebraic comparison. It is not the return from a new allocation policy.')
    table(b,'quarters','Table 6  Returns by quarter',['Quarter','N','First-day mean','Application mean','Covariance pp','Gross profit HK$'],[
        [r['sample'],int(r['n']),pct(r['mean_initial_return']),pct(r['mean_application_return']),f"{r['covariance_contribution_pp']:.2f}",num(r['mean_expected_gross_profit_hkd'])]
        for r in quarters.to_dict('records')],[2,.7,3.2,3.4,3.3,3.4])
    paragraph(b,'The covariance is negative in all three quarters. Mean gross profit is only HK$7.24 in Q3. Profit in dollars also depends on principal and allotted quantity. Demand and issuer characteristics can affect both allocation and return. This decomposition does not identify a causal mechanism.')

    b=page('Sensitivity to fees and large winners')
    table(b,'fees','Table 7  Assumed fixed fees per application',['Fee HK$','Mean profit HK$','Median profit HK$','Mean return','Negative share'],[
        [int(r['fee_hkd']),num(r['mean_expected_profit_hkd']),num(r['median_ipo_expected_profit_hkd']),pct(r['mean_application_return_after_fee']),pct(r['negative_expected_profit_ipo_share'])]
        for r in fees.to_dict('records')],[1.8,3.8,3.8,3.3,3.3])
    paragraph(b,'The fee applies to every application, including applications with no allocation. The return after this fee is (G minus f) divided by B. The negative share counts IPOs with negative expected profit. It does not count actual accounts.')
    paragraph(b,'HK$88 is an assumed parameter. It is not an observed account fee. These scenarios exclude subscription levies, selling costs, financing costs, and opportunity cost. The gross result remains the main result.')
    table(b,'winners','Table 8  Remove the largest expected gross profits',['Removed','N','Gross profit HK$','Profit with HK$88 fee','Gross return'],[
        [int(r['sample'].split('_')[2]),int(r['n']),num(r['mean_expected_gross_profit_hkd']),num(r['mean_expected_profit_fee88_hkd']),pct(r['mean_application_return'])]
        for r in winners.to_dict('records')],[1.9,.7,3.9,5.6,3.9])
    paragraph(b,'The largest-profit issuer is 0501. The largest three are 0501, 2672, and 3231. We rank by expected one-lot profit in dollars. Gross mean profit stays positive after their removal. Profit under the HK$88 scenario becomes negative.')
    paragraph(b,'These checks use information known after listing. They measure influence on the result. They are not application selection rules.')
    sample_names={'corrected_all_113':'All IPOs','prior_112_exclude_2649':'Remove 2649','exclude_3355':'Remove 3355',
                  'exclude_19_total_corrections':'Remove 19 corrected issuers','exclude_2649_3355_and_19':'Remove all 21 issuers'}
    table(b,'source_exclusions','Table 9  Remove issuers with source questions',['Sample','N','Gross return','Covariance pp','Gross profit HK$'],[
        [sample_names[r['sample']],int(r['n']),pct(r['mean_application_return']),f"{r['covariance_contribution_pp']:.2f}",num(r['mean_expected_gross_profit_hkd'])]
        for r in summaries.to_dict('records')],[5.5,.7,3.3,3.3,3.2])
    paragraph(b,'The source questions now have verified answers. The exclusions show dependence on these issuers. Total public applied shares are not the one-lot return denominator. Removing the 19 issuers changes the sample. It does not identify the effect of their data corrections.')

    b=page('Data evidence and remaining limits')
    table(b,'evidence','Table 10  Verified evidence and derived values',['Item','Verified evidence','Derived value or limit'],[
        ['2649 lot size','Official correction: lot and minimum application both equal 500','The frozen lot file still shows 200. This report uses the documented correction.'],
        ['3355 Pool B','An official clarification confirms all ten guaranteed quantities','Expected allocation still combines guarantees with additional ballot probabilities.'],
        ['19 application totals','Independent PDF calculations cover 763 tiers. All 19 formal JSON totals are repaired.','Each total is an exact tier sum. It is not a directly printed aggregate.'],
        ['Other allotment fields','342 cells match the workbooks for these 19 issuers','Eleven cells are missing on both sides. Matching values do not prove every source.']], [3,6.7,7.3])
    paragraph(b,'Fourteen recorded rounded-ratio formulas reproduce the old totals. Two recorded formulas also have arithmetic differences. One old total is only consistent with rounding. The original error causes for two issuers are unrecorded. The repaired totals use source tier sums, not assumed error causes.')
    heading(b,'Data coverage')
    paragraph(b,'Core prices, one-lot rules, and returns cover all 113 IPOs. The VC/PE flag is unknown for seven issuers. Prior-year profit is missing for one issuer. Unrestricted public holding is missing for one issuer. This report does not certify all database fields.')
    paragraph(b,'The application principal uses the final offer price. It excludes levies. It does not measure peak cash blocked at the upper offer price. No actual account records enter the calculations. We do not calculate leverage strategies, cash scheduling, or annualized returns.')
    heading(b,'Sources and reproduction')
    paragraph(b,'The master export supplies the sample description. Verified allocation rules supply expected returns. The report tables retain separate counts for missing values. The source manifest records official PDF links and file hashes.')
    b.append({'type':'source','label':'Master export','path':'../../pipeline/exports/HKIPO-MB-MASTER_clean.csv'})
    b.append({'type':'source','label':'Verified return tables and run manifest','path':'../../analysis/out/retail_evidence_brief/README.md'})
    b.append({'type':'source','label':'Source evidence assessment','path':'RETAIL_EVIDENCE_ASSESSMENT_2026-10-03.md'})
    b.append({'type':'source','label':'The 19 JSON repairs and independent review','path':'../../pipeline/reports/data_gap_collection/source_followup_2026-10-03/formal_json_repair/README.md'})
    b.append({'type':'source','label':'Report tables and run manifest','path':'../../analysis/out/retail_2026_report/run_manifest.json'})
    b.append({'type':'source','label':'Coverage for all 202 Excel fields','path':'../../analysis/out/retail_2026_report/workbook_field_coverage.csv'})
    source_manifest=json.loads((ROOT/'pipeline/reports/data_gap_collection/source_followup_2026-10-03/source_manifest.json').read_text())
    b.append({'type':'url','label':'2649 official correction','url':source_manifest['sources'][0]['url']})
    b.append({'type':'url','label':'3355 official clarification','url':source_manifest['sources'][1]['url']})

    b=page('Short draft for the mentor')
    paragraph(b,'We ask how much of a high first-day IPO return reaches retail application capital. We study 113 ordinary Hong Kong Main Board IPOs. They listed from January to September 2026. We calculate the expected result of one application for one lot. We assume a sale at the first-day closing price.')
    paragraph(b,'The mean first-day return is 53.15%, and the median is 14.82%. The mean gross application return is 2.08%, and the median is 0.59%. The mean expected gross profit is HK$91.77 per application. Its median is HK$27.00. These are IPO-specific expectations, not actual account results.')
    paragraph(b,'The mean return difference is 51.07 percentage points. The denominators differ. First-day return applies to allotted securities. Application return divides expected profit by all requested principal. Low allocation rates limit the return that reaches application capital.')
    paragraph(b,'The product of mean allocation and mean first-day return is 5.73%. Their covariance contributes minus 3.65 percentage points. This negative contribution offsets 63.6% of the product and leaves 2.08%. Higher-return IPOs tend to give smaller allocations. The decomposition describes a sample relationship, not a causal effect.')
    paragraph(b,'The covariance is negative in Q1, Q2, and Q3. However, expected profit changes across quarters. The Q3 mean is only HK$7.24. Large winners also affect the mean. Removing the largest-profit issuer reduces mean gross profit to HK$77.41. Removing the largest three reduces it to HK$58.48.')
    paragraph(b,'Costs require a separate sensitivity check. With an assumed HK$88 fee, mean profit is HK$3.77. Removing the largest winner changes this scenario profit to minus HK$10.59. HK$88 is not an observed account fee. The main finding therefore uses gross returns. Other transaction and financing costs remain outside these calculations.')
    paragraph(b,'We have checked the key source evidence. The official 2649 correction confirms a lot of 500. The 3355 clarification confirms ten Pool B guarantees. Independent calculations verify 763 tiers for 19 application totals. The formal JSON files now contain these totals. Source exclusions retain the large return gap and negative covariance.')
    paragraph(b,'We propose application capital as the main return denominator for this research question. We seek feedback on the next mechanism test. Demand measures or institutional variation could help separate allocation mechanics from information effects. We will choose that test before we add a cash constraint model or simulation.')

    b=page('Terms and methods')
    paragraph(b,'IPO means initial public offering. HDR means Hong Kong depositary receipt. VC/PE means venture capital or private equity. HSI means Hang Seng Index. HIBOR means Hong Kong Interbank Offered Rate. CAR means cumulative abnormal return relative to a benchmark.')
    heading(b,'Method and language checks')
    paragraph(b,'The main results use raw returns and equal IPO weights. There are no regression estimates or significance tests in this report. All displayed numbers come from saved tables or the master export. We retain the full precision in the CSV files.')
    paragraph(b,'The text follows ASD-STE100 writing rules for short descriptive sentences and consistent terms. The report defines the required finance terms. The automated check tests sentence length and paragraph size. It does not certify full dictionary compliance.')
    b.append({'type':'url','label':'ASD-STE100 official writing specification, Issue 9','url':'https://www.asd-ste100.org/assets/files/ASD-STE100_ISSUE9.pdf'})
    paragraph(b,'To reproduce the tables, run the report script from the project root. Compile the standalone LaTeX source to produce the PDF.')
    b.append({'type':'code','text':'python analysis/retail_2026_report.py'})
    paragraph(b,'Source and calculation code remain separate. The collection pipeline supplies evidence. The analysis code calculates the report results. No pipeline analysis step or simulation is added.')

    audit=[]
    for p in pages:
        for block in p['blocks']:
            if block['type']!='paragraph':continue
            sentences=re.split(r'(?<=[.!?])\s+(?=[A-Z])',block['text'])
            # Digits in percentages and decimals stay within the same sentence.
            counts=[len(re.findall(r"[A-Za-z0-9]+(?:[-'][A-Za-z0-9]+)*",s)) for s in sentences]
            audit.append({'page':p['title'],'text':block['text'],'sentences':len(sentences),'word_counts':counts})
            assert len(sentences)<=6,(p['title'],block['text'])
            assert max(counts,default=0)<=25,(p['title'],counts,block['text'])
    (OUT/'language_check.json').write_text(json.dumps({'writing_specification':'ASD-STE100 Issue 9','scope':'descriptive sentence length and paragraph sentence count; consistent technical terms',
        'full_dictionary_compliance_certified':False,'max_sentence_words':max(max(x['word_counts']) for x in audit),'paragraphs':audit},indent=2)+'\n')

    preamble=r'''\documentclass[10pt,a4paper]{article}
\usepackage[margin=20mm]{geometry}
\usepackage[T1]{fontenc}
\usepackage{lmodern}
\usepackage{amsmath,booktabs,array,longtable}
\usepackage[hidelinks]{hyperref}
\usepackage{fancyhdr}
\setlength{\parindent}{0pt}
\setlength{\parskip}{3pt}
\renewcommand{\baselinestretch}{1.02}
\setlength{\emergencystretch}{2em}
\pagestyle{fancy}
\fancyhf{}
\fancyhead[L]{\small Hong Kong IPOs in 2026}
\fancyhead[R]{\small Cutoff 30 September 2026}
\fancyfoot[C]{\small\thepage}
\renewcommand{\headrulewidth}{0pt}
\setlength{\headheight}{14pt}
\hypersetup{pdftitle={Hong Kong IPOs in 2026},pdfsubject={Data description and retail application returns},pdfauthor={}}
\begin{document}
'''
    latex=[preamble];md=[]
    for i,p in enumerate(pages):
        if p['title']=='Short draft for the mentor':latex.append(r'\clearpage')
        if i==0:latex.append(r'{\LARGE\bfseries '+tex(p['title'])+r'}\par\addvspace{4pt}')
        else:latex.append(r'\section*{'+tex(p['title'])+'}')
        md.append(('# ' if i==0 else '## ')+p['title'])
        for block in p['blocks']:
            kind=block['type']
            if kind=='paragraph':latex.append(tex(block['text'])+'\n\n');md.append(block['text'])
            elif kind=='heading':latex.append(r'\par\addvspace{4pt}\textbf{'+tex(block['text'])+r'}\par\nobreak');md.append('### '+block['text'])
            elif kind=='table':
                latex.append(latex_table(block));md.append('### '+block['title'])
                rows=[block['headers'],['---']*len(block['headers']),*block['rows']]
                md.append('\n'.join('| '+' | '.join(map(str,row))+' |' for row in rows))
            elif kind=='equation':latex.append(r'\[\small '+block['tex']+r'\]');md.append('$$\n'+block['tex']+'\n$$')
            elif kind=='source':latex.append(r'\noindent\href{'+block['path']+'}{'+tex(block['label'])+r'}\par');md.append(f"[{block['label']}]({block['path']})")
            elif kind=='url':latex.append(r'\noindent\href{'+block['url']+'}{'+tex(block['label'])+r'}\par');md.append(f"[{block['label']}]({block['url']})")
            elif kind=='code':latex.append(r'\noindent\texttt{'+tex(block['text'])+r'}\par');md.append('`'+block['text']+'`')
    latex.append(r'\end{document}')
    STEM.with_suffix('.tex').write_text('\n'.join(latex)+'\n')
    STEM.with_suffix('.md').write_text('\n\n'.join(md)+'\n')
    inputs=[MASTER,registry_path,*workbooks,taxonomy_path,ROOT/'analysis/research_inputs.py',Path(__file__),BRIEF/'run_manifest.json',*[BRIEF/name for name in provenance['outputs']]]
    outputs=[STEM.with_suffix('.tex'),STEM.with_suffix('.md'),*sorted(OUT.glob('*.csv')),OUT/'language_check.json']
    manifest={'scope':'2026 Q1-Q3; 113 observed IPOs; English report and mentor draft','simulation_used':False,
        'inputs':{str(p.relative_to(ROOT)):sha(p) for p in inputs},'outputs':{str(p.relative_to(ROOT)):sha(p) for p in outputs}}
    (OUT/'run_manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
    print(f'Created {len(pages)} report sections with automatic pagination. Language check: maximum {max(max(x["word_counts"]) for x in audit)} words per sentence.')


if __name__=='__main__':
    main()
