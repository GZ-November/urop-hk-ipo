"""GEM source layout, actual template mapping, evidence gates and board isolation."""
from __future__ import annotations

import datetime as dt
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

PIPE = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(PIPE), str(PIPE / "src"), str(PIPE / "tools")]

from boards import report_source_dir  # noqa: E402
from cohort import build_cohort_workbook, load_cfg, write_cohort_config  # noqa: E402
from extraction_targets import target_paths  # noqa: E402
from gem_template import GEM_EXCLUDED_HEADERS, update_gem_view  # noqa: E402
from listing_reports import report_as_of, report_links, reports_for_interval  # noqa: E402
from sample_builder import load_nlr_candidates  # noqa: E402
from state import require, save_record  # noqa: E402
from validate import validate_all  # noqa: E402
from write_back import write_all  # noqa: E402


def gem_report(path, year=2025):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet['A5'] = f'Year : {year} (up to {year}/10/10)'
    sheet['B9'] = 'Stock'
    sheet['B10'] = 'code'
    for column, value in {1: dt.datetime(year, 10, 10), 2: 8549, 3: 'Fixture GEM Limited',
                          6: 0.5, 8: 11464.72, 10: 50000000, 11: 400000000,
                          13: 200000000, 15: 'Construction', 17: 'Cayman Islands',
                          18: 'Offer for subscription', 20: 'Sponsor', 22: 'Auditor'}.items():
        sheet.cell(12, column).value = value
    sheet['R13'] = 'Offer for placing'  # continuation must not create another issuer
    workbook.save(path)
    workbook.close()


class GemPipelineTests(unittest.TestCase):
    def test_links_and_source_caches_are_board_specific(self):
        html = '<a href="/-/media/HKEXnews/New-Listing-Report/GEM/e_newlistings2025.xlsx">GEM</a><a href="/-/media/HKEXnews/New-Listing-Report/Main/NLR2025_Eng.xlsx">MB</a>'
        self.assertIn('e_newlistings', report_links(html, 'gem')[2025])
        self.assertIn('NLR2025', report_links(html)[2025])
        self.assertNotEqual(report_source_dir(Path('/tmp/x'), 'gem'), report_source_dir(Path('/tmp/x'), 'main'))

    def test_sparse_historical_gem_report_layout_and_coverage(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / 'e_newlistings2025.xlsx'
            gem_report(path)
            self.assertEqual(report_as_of(path), dt.date(2025, 10, 10))
            rows = load_nlr_candidates(path, 2025, board='gem')
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]['stock_code'], '8549.HK')
            self.assertEqual(rows[0]['gem_funds_raised_hkd'], 50000000)
            self.assertIsNone(rows[0]['funds_raised_public_hkd'])
            self.assertIsNone(rows[0]['prospectus_date'])
            self.assertEqual(reports_for_interval(dt.date(2025, 1, 1), dt.date(2025, 12, 31),
                                                 Path(temporary), today=dt.date(2026, 10, 12), board='gem'), [path])

    def test_empty_official_gem_interval_exports_zero_rows(self):
        from codebook import export_all
        datasets = PIPE / "datasets"
        datasets.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=datasets) as temporary:
            root = Path(temporary)
            source = root / "sources"
            source.mkdir()
            gem_report(source / "e_newlistings2025.xlsx")
            book = build_cohort_workbook(dt.date(2025, 1, 1), dt.date(2025, 3, 31),
                                        source_dir=source, workbook_path=root / "HKIPO-GEM.xlsx",
                                        today=dt.date(2026, 10, 12), board="gem")
            config = write_cohort_config(dt.date(2025, 1, 1), dt.date(2025, 3, 31), book, source_dir=source, board="gem")
            exports = export_all(load_cfg(config_path=config), out_dir=root / "delivery")
            self.assertEqual(exports["sample_size"], 0)
            self.assertEqual(len(Path(exports["csv_path"]).read_text().splitlines()), 1)

    def test_current_year_stale_gem_report_does_not_claim_full_coverage(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / 'e_newlistings2025.xlsx'
            gem_report(path)
            with patch('listing_reports.requests.Session') as session:
                session.return_value.get.side_effect = OSError('offline')
                with self.assertRaisesRegex(ValueError, 'coverage missing'):
                    reports_for_interval(dt.date(2025, 1, 1), dt.date(2025, 12, 31),
                                         Path(temporary), today=dt.date(2025, 12, 31), board='gem')

    def test_clean_template_preserves_46_student_headers_and_deletes_main_only_fields(self):
        original = openpyxl.load_workbook(PIPE.parent / 'templates/HKIPO-GEM-template-students.xlsx', read_only=True)
        template = openpyxl.load_workbook(PIPE.parent / 'templates/HKIPO-GEM-template-final.xlsx', read_only=True)
        try:
            student = next(original.active.iter_rows(max_row=1, max_col=46, values_only=True))
            actual = next(template['GEM'].iter_rows(max_row=1, max_col=46, values_only=True))
            self.assertEqual(student, actual)
            for sheet in template:
                headers = next(sheet.iter_rows(max_row=1, values_only=True))
                self.assertFalse(GEM_EXCLUDED_HEADERS.intersection(headers))
                self.assertEqual(sheet.max_row, 1)
                self.assertLess(sheet.max_column, 220)
        finally:
            original.close()
            template.close()

    def test_gem_extension_requires_independent_current_review_and_preserves_layout(self):
        datasets = PIPE / 'datasets'
        datasets.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=datasets) as temporary:
            directory = Path(temporary)
            source = directory / 'sources'
            source.mkdir()
            gem_report(source / 'e_newlistings2025.xlsx')
            book = build_cohort_workbook(dt.date(2025, 10, 1), dt.date(2025, 10, 31),
                                        source_dir=source, workbook_path=directory / 'HKIPO-GEM.xlsx',
                                        today=dt.date(2026, 10, 12), board='gem')
            config = write_cohort_config(dt.date(2025, 10, 1), dt.date(2025, 10, 31), book, source_dir=source, board='gem')
            cfg = load_cfg(config_path=config)
            schema_path, output, packets = target_paths(cfg, 'gem')
            schema = json.loads(schema_path.read_text())
            payload = {'code': '8549.HK', 'fields': {f['key']: {'value': 'NA' if f['kind'] in ('date', 'text') else 'NaN', 'page': None, 'quote': '', 'confidence': 'low'} for f in schema['fields']}}
            payload['fields']['gem_eligibility_rule'] = {'value': 'GEM Rule 11.12A', 'page': 1, 'quote': 'We meet the requirements of GEM Rule 11.12A.', 'confidence': 'high'}
            (output / 'extracted').mkdir(parents=True)
            (output / 'extracted/HKIPO-MB8549.json').write_text(json.dumps(payload))
            (packets / 'HKIPO-MB8549.md').write_text('<<<PAGE 1>>>\nWe meet the requirements of GEM Rule 11.12A.\n')
            result = validate_all(cfg, target='gem', log=lambda *_: None)
            self.assertTrue(result['gate_pass'], result['errors'])
            record = dict(result['records'][0], target='gem', gate_pass=True)
            save_record(cfg, 'gem', '8549.HK', 'extracted', record)
            with self.assertRaises(SystemExit):
                write_all(cfg, target='gem', log=lambda *_: None)
            # Test-only review fixture; production approvals are generated by actual reviewers.
            save_record(cfg, 'gem', '8549.HK', 'reviewed', dict(record, verdict='pass', reviewer='independent-test-fixture'))
            self.assertEqual(write_all(cfg, target='gem', log=lambda *_: None), 1)
            workbook = openpyxl.load_workbook(book)
            try:
                self.assertEqual(workbook['GEM_details']['F2'].value, 'GEM Rule 11.12A')
                update_gem_view(workbook)
                self.assertEqual(workbook['GEM']['F2'].value, 50000000)
                self.assertEqual(workbook['GEM']['G2'].value, 400000000)
                self.assertEqual(workbook['GEM']['H2'].value, 200000000)
                self.assertIsNone(workbook['NLR']['I2'].value)  # no invented tranche split
            finally:
                workbook.close()
            with self.assertRaises(ValueError):
                require(cfg, 'gem', '8549.HK', 'changed-payload', 'reviewed')


if __name__ == '__main__':
    unittest.main()
