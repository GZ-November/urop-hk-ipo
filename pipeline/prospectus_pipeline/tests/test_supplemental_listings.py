"""Reviewed supplement entries extend an annual listing report's coverage."""
from __future__ import annotations

import datetime as dt
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "src"))

from cohort import build_cohort_workbook  # noqa: E402
from listing_reports import load_supplement, reports_and_supplement  # noqa: E402
import listing_reports  # noqa: E402
import openpyxl  # noqa: E402
from unittest.mock import patch  # noqa: E402

TEMPLATE = ROOT.parent / "templates" / "HKIPO-MB-template-final.xlsx"
TODAY = dt.date(2026, 10, 1)

SUPPLEMENT = """\
year: 2026
report_as_of: "2026-06-30"
entries:
  - stock_code: "6802.HK"
    company_name: "Supplement IPO Ltd - H Shares"
    prospectus_date: "2026-09-22"
    listing_date: "2026-09-30"
    offer_price_hkd: 58.85
    sources:
      - url: "https://www1.hkexnews.hk/listedco/listconews/sehk/2026/0929/2026092901927.pdf"
        retrieved: "2026-10-01"
"""


def make_report(path: Path, year: int, as_of: str, rows: list[list]) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([f"Main Board New Listings up to {as_of}"])
    sheet.append(["File no", "Stock Code", "Company Name", "Date of Prospectus",
                  "Date of Listing", "Sponsor(s)", "Accountant", "Valuer",
                  "Funds Raised HK (a)", "Offer Price", "Fund Type"])
    for row in rows:
        sheet.append(row)
    workbook.save(path / f"NLR{year}_Eng.xlsx")
    workbook.close()


class SupplementTests(unittest.TestCase):
    def setUp(self):
        self.temporary_source = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary_source.cleanup)
        self.source = Path(self.temporary_source.name)
        make_report(self.source, 2026, "31 July 2026", [
            [1, 6082, "First IPO Ltd", dt.datetime(2026, 7, 20),
             dt.datetime(2026, 7, 27), "Sponsor", "Accountant", None, 100, 10],
        ])
        (self.source / "NLR2026_Supplement_20260930.yaml").write_text(SUPPLEMENT, encoding="utf-8")

    def test_supplement_only_covers_days_after_the_report(self):
        entries = load_supplement(
            self.source, 2026, dt.date(2026, 1, 1), dt.date(2026, 9, 30),
            covered_through=dt.date(2026, 6, 30),
        )
        self.assertEqual([entry["stock_code"] for entry in entries], ["6802.HK"])
        self.assertEqual(entries[0]["listing_date"], dt.date(2026, 9, 30))
        self.assertEqual(entries[0]["offer_price_hkd"], 58.85)
        self.assertTrue(entries[0]["supplement_sources"])

    def test_supplement_entry_inside_report_coverage_fails_closed(self):
        with self.assertRaisesRegex(ValueError, "outside the gap"):
            load_supplement(
                self.source, 2026, dt.date(2026, 1, 1), dt.date(2026, 9, 30),
                covered_through=dt.date(2026, 9, 30),
            )

    def test_supplement_entry_without_source_is_rejected(self):
        text = SUPPLEMENT.replace('url: "https://www1.hkexnews.hk/listedco/listconews/sehk/2026/0929/2026092901927.pdf"', "url: null")
        (self.source / "NLR2026_Supplement_20260930.yaml").write_text(text, encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "no official source URL"):
            load_supplement(
                self.source, 2026, dt.date(2026, 1, 1), dt.date(2026, 9, 30),
                covered_through=dt.date(2026, 6, 30),
            )

    def test_reports_and_supplement_extends_stale_annual_report(self):
        stale = ValueError("report covers through 2026-07-31, requested through 2026-09-30")
        with patch.object(listing_reports, "reports_for_interval", side_effect=stale):
            reports, supplement = reports_and_supplement(
                dt.date(2026, 7, 1), dt.date(2026, 9, 30), self.source, today=TODAY,
            )
        self.assertEqual([report.name for report in reports], ["NLR2026_Eng.xlsx"])
        self.assertEqual([entry["stock_code"] for entry in supplement], ["6802.HK"])

    def test_reports_and_supplement_without_file_stays_fail_closed(self):
        (self.source / "NLR2026_Supplement_20260930.yaml").unlink()
        stale = ValueError("report covers through 2026-07-31, requested through 2026-09-30")
        with patch.object(listing_reports, "reports_for_interval", side_effect=stale):
            with self.assertRaisesRegex(ValueError, "coverage missing"):
                reports_and_supplement(
                    dt.date(2026, 7, 1), dt.date(2026, 9, 30), self.source, today=TODAY,
                )

    def test_build_cohort_workbook_reconciles_supplemented_membership(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "cohort.xlsx"
            stale = ValueError("report covers through 2026-07-31, requested through 2026-09-30")
            with patch.object(listing_reports, "reports_for_interval", side_effect=stale):
                build_cohort_workbook(
                    dt.date(2026, 7, 1), dt.date(2026, 9, 30),
                    source_dir=self.source, template_path=TEMPLATE,
                    workbook_path=path, today=TODAY,
                )
                workbook = openpyxl.load_workbook(path, read_only=True)
                sheet = workbook["NLR"]
                codes = [value for (value,) in sheet.iter_rows(
                    min_row=2, min_col=2, max_col=2, values_only=True
                ) if value]
                workbook.close()
                self.assertEqual(codes, ["6082.HK", "6802.HK"])
                # Re-run: the membership now matches and the run stays read-only.
                result = build_cohort_workbook(
                    dt.date(2026, 7, 1), dt.date(2026, 9, 30),
                    source_dir=self.source, template_path=TEMPLATE,
                    workbook_path=path, today=TODAY,
                )
            self.assertEqual(result, path)


if __name__ == "__main__":
    unittest.main()
