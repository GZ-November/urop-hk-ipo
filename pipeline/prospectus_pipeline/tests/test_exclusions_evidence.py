import datetime as dt
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / "src")]
WS = ROOT.parent

from evidence import build_evidence_manifest  # noqa: E402
from exclusions import build_exclusion_log  # noqa: E402
from master_panel import _coverage_pct, build_registry  # noqa: E402


class DeclaredDtypeTests(unittest.TestCase):
    def test_coverage_pct_parsing(self):
        self.assertEqual(_coverage_pct("100% 完备"), 100.0)
        self.assertEqual(_coverage_pct("Adequate (14/23 (60.87%))"), 60.87)
        self.assertEqual(_coverage_pct("Sparse (2/23 (8.7%))"), 8.7)
        self.assertEqual(_coverage_pct("Reserved / Unmatured"), 0.0)

    def test_declared_dtype_uses_book_with_data(self):
        def codebook_md(tag: str, dtype: str, coverage: str) -> str:
            return (
                f"# codebook {tag}\n\n"
                "| 列 | 变量 | 释义 | 层级 | 类型 | 时点 | 样本状态 | 统计 |\n"
                "|---|---|---|---|---|---|---|---|\n"
                f"| **ER** | `6-month post-IPO close price (HK$)` | 测试 | 深蓝 | "
                f"`{dtype}` | Post-IPO | {coverage} | stats |\n"
            )

        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            (tmp / "codebooks").mkdir(exist_ok=True)
            (tmp / "codebooks" / "HKIPO_2026Q1_Codebook.md").write_text(
                codebook_md("Q1", "numeric", "84.4% 完备"), encoding="utf-8")
            (tmp / "codebooks" / "HKIPO_2026Q2_Codebook.md").write_text(
                codebook_md("Q2", "string", "Reserved / Unmatured (0/23 (0.0%))"), encoding="utf-8")
            registry = build_registry(tmp, out_path=tmp / "reg.yaml")
            er = [v for v in registry["variables"] if v["letter"] == "ER"][0]
            self.assertEqual(er["declared_dtype"], "numeric")
            self.assertEqual(er["declared_from"], "HKIPO_2026Q1_Codebook.md")


class EvidenceManifestTests(unittest.TestCase):
    def test_manifest_covers_matching_files_only(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            (tmp / "cohorts").mkdir(exist_ok=True)
            (tmp / "registry").mkdir(exist_ok=True)
            (tmp / "cohorts" / "HKIPO-MB2025Q1.xlsx").write_bytes(b"workbook")
            (tmp / "registry" / "HKIPO_Variable_Registry.yaml").write_text("meta: {}\n", encoding="utf-8")
            (tmp / "unrelated.txt").write_text("ignore me", encoding="utf-8")
            out_path, entries = build_evidence_manifest(ws=tmp)
            names = [name for name, _ in entries]
            self.assertEqual(names, ["cohorts/HKIPO-MB2025Q1.xlsx", "registry/HKIPO_Variable_Registry.yaml"])
            text = out_path.read_text(encoding="utf-8")
            self.assertIn("cohorts/HKIPO-MB2025Q1.xlsx", text)
            self.assertNotIn("unrelated.txt", text)


class ExclusionLogTests(unittest.TestCase):
    def make_cfg(self, workbook: Path) -> dict:
        return {
            "sheet": "NLR",
            "_workbook_path": workbook,
            "dataset": {
                "cohort": "2025 Q1",
                "period_start": "2025-01-01",
                "period_end": "2025-03-31",
                "report_source_dir": "sources",
            },
        }

    def test_log_reports_exclusions_and_gaps(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            workbook = tmp / "HKIPO-MB2025Q1.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "NLR"
            ws["B1"] = "Stock Code"
            ws["B2"] = "0001.HK"
            wb.save(workbook)

            candidates = [
                {"stock_code": "0001.HK", "company_name": "In Book Ltd",
                 "listing_date": dt.datetime(2025, 2, 10), "inclusion_status": "INCLUDED"},
                {"stock_code": "0002.HK", "company_name": "Missing Ltd",
                 "listing_date": dt.datetime(2025, 3, 5), "inclusion_status": "INCLUDED"},
                {"stock_code": "0003.HK", "company_name": "SPAC Ltd",
                 "listing_date": dt.datetime(2025, 1, 20), "inclusion_status": "EXCLUDED",
                 "exclusion_reason": "SPAC offering under Chapter 18B",
                 "listing_route": "SPAC"},
                {"stock_code": "0004.HK", "company_name": "Outside Period Ltd",
                 "listing_date": dt.datetime(2025, 5, 1), "inclusion_status": "INCLUDED"},
            ]
            out_path = tmp / "HKIPO_2025Q1_Exclusions.md"
            _, summary = build_exclusion_log(
                self.make_cfg(workbook), candidates=candidates, out_path=out_path)
            self.assertEqual(summary["included"], 2)
            self.assertEqual(summary["excluded"], 1)
            self.assertEqual(summary["missing_from_workbook"], ["0002.HK"])
            text = out_path.read_text(encoding="utf-8")
            self.assertIn("SPAC offering under Chapter 18B", text)
            self.assertIn("0002.HK", text)
            self.assertNotIn("0004.HK", text)  # 区间外不进入日志


if __name__ == "__main__":
    unittest.main()
