"""state.py 门禁链的单元测试：凭证存取、哈希绑定、三段授权。"""
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / "src")]

from state import (
    CONTRACT_VERSION,
    authorized,
    cells_digest,
    digest,
    read_record,
    record_path,
    require,
    save_record,
    state_dir,
)


def make_cfg(tmp: Path) -> dict:
    return {
        "_root": tmp,
        "_ws": tmp,
        "_config_path": str(tmp / "config.yaml"),
        "dataset": {"id": "TEST"},
    }


class StateDigestTests(unittest.TestCase):
    def test_digest_stable_and_sensitive(self):
        a = digest({"k": 1}, {"s": 2}, {"e": 3})
        self.assertEqual(a, digest({"k": 1}, {"s": 2}, {"e": 3}))
        self.assertNotEqual(a, digest({"k": 2}, {"s": 2}, {"e": 3}))

    def test_cells_digest_tracks_value_changes(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as tmp:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws["A1"] = 1
            before = cells_digest(ws, ["A1"])
            ws["A1"] = 2
            self.assertNotEqual(before, cells_digest(ws, ["A1"]))
            self.assertEqual(len(before), 64)


class StateRecordTests(unittest.TestCase):
    def test_record_path_normalizes_code_and_rejects_digitless(self):
        with tempfile.TemporaryDirectory() as tmp:
            cfg = make_cfg(Path(tmp))
            # 数字代码被归一为 4 位 + .HK（"12" -> 0012.HK）
            path = record_path(cfg, "prospectus", "12", "extracted")
            self.assertIn("0012.HK.json", str(path))
            # 无数字代码无法归一，必须被拒
            with self.assertRaises(ValueError):
                record_path(cfg, "prospectus", "AB.CD", "extracted")

    def test_save_read_round_trip_enriches_record(self):
        with tempfile.TemporaryDirectory() as tmp:
            cfg = make_cfg(Path(tmp))
            save_record(cfg, "prospectus", "0001.HK", "reviewed", {"code": "0001.HK", "gate_pass": True})
            rec = read_record(cfg, "prospectus", "0001.HK", "reviewed")
            self.assertEqual(rec["code"], "0001.HK")
            self.assertEqual(rec["contract_version"], CONTRACT_VERSION)
            self.assertEqual(rec["stage"], "reviewed")
            self.assertIn("recorded_at", rec)

    def test_require_rejects_hash_mismatch(self):
        with tempfile.TemporaryDirectory() as tmp:
            cfg = make_cfg(Path(tmp))
            save_record(cfg, "prospectus", "0001.HK", "validated",
                        {"code": "0001.HK", "target": "prospectus", "hash": "aaa", "gate_pass": True})
            with self.assertRaises(ValueError):
                require(cfg, "prospectus", "0001.HK", "bbb", "validated")
            rec = require(cfg, "prospectus", "0001.HK", "aaa", "validated")
            self.assertEqual(rec["code"], "0001.HK")

    def test_require_rejects_gate_not_passed(self):
        with tempfile.TemporaryDirectory() as tmp:
            cfg = make_cfg(Path(tmp))
            save_record(cfg, "prospectus", "0001.HK", "validated",
                        {"code": "0001.HK", "target": "prospectus", "hash": "aaa", "gate_pass": False})
            with self.assertRaises(ValueError):
                require(cfg, "prospectus", "0001.HK", "aaa", "validated")

    def test_authorized_requires_all_three_stages(self):
        with tempfile.TemporaryDirectory() as tmp:
            cfg = make_cfg(Path(tmp))
            current = {"status": "pass", "hash": "h1"}
            # 缺少任何一段凭证都要被拒
            with self.assertRaises(ValueError):
                authorized(cfg, "0001.HK", "prospectus", current)
            for stage in ("extracted", "validated", "reviewed"):
                save_record(cfg, "prospectus", "0001.HK", stage,
                            {"code": "0001.HK", "target": "prospectus", "hash": "h1", "gate_pass": True})
            authorized(cfg, "0001.HK", "prospectus", current)  # 不抛即通过

    def test_authorized_rejects_failed_snapshot(self):
        with tempfile.TemporaryDirectory() as tmp:
            cfg = make_cfg(Path(tmp))
            with self.assertRaises(ValueError):
                authorized(cfg, "0001.HK", "prospectus", {"status": "fail", "hash": "h1", "errors": ["x"]})

    def test_state_dir_isolates_non_default_configs(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            (tmp / "config.yaml").write_text("workbook: x\n", encoding="utf-8")
            (tmp / "other.yaml").write_text("workbook: y\n", encoding="utf-8")
            base = state_dir(make_cfg(tmp))
            isolated = state_dir({**make_cfg(tmp), "_config_path": str(tmp / "other.yaml"),
                                  "dataset": {"id": "OTHER"}})
            self.assertEqual(base, tmp / ".pipeline_state")
            self.assertEqual(isolated, tmp / ".pipeline_state" / "OTHER")


if __name__ == "__main__":
    unittest.main()
