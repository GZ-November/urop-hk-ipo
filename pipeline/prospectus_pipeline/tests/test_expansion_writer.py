import csv
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "src"), str(ROOT)]
from write_back_expansion import (
    ExpansionOverwriteError,
    OverwriteConflict,
    WorkbookExpansionWriter,
    find_overwrite_conflicts,
)


def _build_cohort(root, stab_rows=(), curated=None):
    """单行样本工作簿 (row 2 = 1234.HK) + 最小 master 输出；curated={col: value} 预填 row 2。"""
    workbook_path = root / "cohort.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "NLR"
    sheet.append(["File", "Code", "Name", "Prospectus", "Listing", None, None, None, None, None, "Offer", *CORNERSTONE_HEADERS])
    sheet.append([1, "1234.HK", "Selected", datetime(2026, 1, 2), datetime(2026, 1, 10), None, None, None, None, None, 10.0, 0.4, datetime(2026, 7, 10)])
    for col, value in (curated or {}).items():
        sheet.cell(2, col).value = value
    workbook.save(workbook_path)
    workbook.close()

    out_master = root / "out" / "master"
    out_master.mkdir(parents=True)
    (root / "out" / "allot" / "extracted").mkdir(parents=True, exist_ok=True)
    with (out_master / "stabilization_events.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=["stock_code", "stabilizing_manager", "stabilization_period_end"])
        writer.writeheader()
        writer.writerows(stab_rows)
    with (out_master / "daily_market_panel.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        csv.DictWriter(stream, fieldnames=["stock_code", "amihud_illiq", "zero_volume_flag", "daily_return", "max_drawdown"]).writeheader()
    with (out_master / "horizon_summary.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        csv.DictWriter(stream, fieldnames=["stock_code", "horizon", "bhr_from_day1", "wr_hsi"]).writeheader()
    with (out_master / "lockup_events.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        csv.DictWriter(stream, fieldnames=["stock_code", "lockup_category"]).writeheader()

    cfg = {
        "workbook": str(workbook_path),
        "workbook_path": workbook_path,
        "_workbook_path": workbook_path,
        "sheet": "NLR",
        "data_start_row": 2,
        "id_columns": {"file_no": "A", "stock_code": "B", "name": "C",
                       "prospectus_date": "D", "listing_date": "E", "offer_price": "K"},
        "dataset": {"period_start": "2026-01-01", "period_end": "2026-01-31"},
        "filter_period": True,
        "paths": {"out": root / "out", "allot_out": root / "out" / "allot"},
    }
    return workbook_path, cfg


def _row2(workbook_path, *cols):
    book = openpyxl.load_workbook(workbook_path)
    try:
        return [book["NLR"].cell(2, col).value for col in cols]
    finally:
        book.close()


CORNERSTONE_HEADERS = ["Final cornerstone allocation (% of base offer)", "Earliest cornerstone unlock date (dd/mm/yy)"]


def _write_master(out_master, lockup_rows=()):
    out_master.mkdir(parents=True)
    with (out_master / "stabilization_events.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=["stock_code", "stabilizing_manager", "stabilization_period_end"],
        )
        writer.writeheader()
        writer.writerow({"stock_code": "1234.HK", "stabilizing_manager": "Example Sponsor", "stabilization_period_end": "2026-02-09"})

    with (out_master / "daily_market_panel.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        csv.DictWriter(stream, fieldnames=["stock_code", "amihud_illiq", "zero_volume_flag", "daily_return", "max_drawdown"]).writeheader()
    with (out_master / "horizon_summary.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        csv.DictWriter(stream, fieldnames=["stock_code", "horizon", "bhr_from_day1", "wr_hsi"]).writeheader()
    with (out_master / "lockup_events.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=["stock_code", "lockup_category", "car_m5_p5", "car_m20_p20", "volume_shock_ratio"],
        )
        writer.writeheader()
        writer.writerows(lockup_rows)


def _cfg(root, workbook_path, period_end="2026-01-31"):
    return {
        "workbook": str(workbook_path),
        "workbook_path": workbook_path,
        "_workbook_path": workbook_path,
        "sheet": "NLR",
        "data_start_row": 2,
        "id_columns": {
            "file_no": "A",
            "stock_code": "B",
            "name": "C",
            "prospectus_date": "D",
            "listing_date": "E",
            "offer_price": "K",
        },
        "dataset": {"period_start": "2026-01-01", "period_end": period_end},
        "filter_period": True,
        "paths": {"out": root / "out", "allot_out": root / "allot_out"},
    }


class ExpansionWriterTests(unittest.TestCase):
    def test_writer_uses_mapper_and_only_writes_selected_cohort_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook_path = root / "cohort.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "NLR"
            sheet.append(["File", "Code", "Name", "Prospectus", "Listing", None, None, None, None, None, "Offer", *CORNERSTONE_HEADERS])
            sheet.append([1, "1234.HK", "Selected", datetime(2026, 1, 2), datetime(2026, 1, 10), None, None, None, None, None, 10.0, 0.4, datetime(2026, 7, 10)])
            sheet.append([2, "5678.HK", "Outside", datetime(2026, 2, 2), datetime(2026, 2, 10), None, None, None, None, None, 20.0, 0.4, datetime(2026, 8, 10)])
            workbook.save(workbook_path)
            workbook.close()

            _write_master(root / "out" / "master")
            result = WorkbookExpansionWriter(cfg=_cfg(root, workbook_path)).write_expansion()

            check = openpyxl.load_workbook(workbook_path, data_only=True)
            output = check["NLR"]
            self.assertEqual(result, 0)
            self.assertEqual(output.cell(1, 162).value, "Stabilizing manager")
            self.assertEqual(output.cell(2, 162).value, "Example Sponsor")
            self.assertEqual(output.cell(2, 163).value, "2026-02-09")
            self.assertIsNone(output.cell(3, 162).value)
            check.close()

    def test_cornerstone_unlock_metrics_blank_for_issuers_without_cornerstones(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook_path = root / "cohort.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "NLR"
            sheet.append(["File", "Code", "Name", "Prospectus", "Listing", None, None, None, None, None, "Offer", *CORNERSTONE_HEADERS])
            # 有基石 / 解禁日 NA / 最终配售为 0 但解禁日仍存在（2026Q2 形态）
            sheet.append([1, "1111.HK", "With CS", datetime(2026, 1, 2), datetime(2026, 1, 5), None, None, None, None, None, 10.0, 0.35, datetime(2026, 7, 5)])
            sheet.append([2, "2222.HK", "No CS NA", datetime(2026, 1, 2), datetime(2026, 1, 6), None, None, None, None, None, 10.0, 0, "NA"])
            sheet.append([3, "3333.HK", "No CS zero", datetime(2026, 1, 2), datetime(2026, 1, 7), None, None, None, None, None, 10.0, 0, datetime(2026, 7, 7)])
            # 预先存在的陈旧值必须被清空
            for row in (3, 4):
                for col in (187, 188, 189):
                    sheet.cell(row, col, 9.99)
            workbook.save(workbook_path)
            workbook.close()

            lockups = [
                {"stock_code": code, "lockup_category": "Cornerstone_Lockup",
                 "car_m5_p5": "-0.05", "car_m20_p20": "0.12", "volume_shock_ratio": "1.7"}
                for code in ("1111.HK", "2222.HK", "3333.HK")
            ]
            _write_master(root / "out" / "master", lockups)
            # 陈旧值属于受保护的非空单元格：清除需要显式 force_overwrite
            with self.assertRaises(ExpansionOverwriteError):
                WorkbookExpansionWriter(cfg=_cfg(root, workbook_path)).write_expansion()
            WorkbookExpansionWriter(cfg=_cfg(root, workbook_path)).write_expansion(force_overwrite=True)

            check = openpyxl.load_workbook(workbook_path, data_only=True)
            output = check["NLR"]
            self.assertEqual([output.cell(2, c).value for c in (187, 188, 189)], [-0.05, 0.12, 1.7])
            for row in (3, 4):
                self.assertEqual([output.cell(row, c).value for c in (187, 188, 189)], [None, None, None])
            check.close()

    def test_writer_fails_closed_without_cornerstone_columns(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook_path = root / "cohort.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "NLR"
            sheet.append(["File", "Code", "Name", "Prospectus", "Listing", None, None, None, None, None, "Offer"])
            sheet.append([1, "1234.HK", "Selected", datetime(2026, 1, 2), datetime(2026, 1, 10), None, None, None, None, None, 10.0])
            workbook.save(workbook_path)
            workbook.close()

            _write_master(root / "out" / "master")
            with self.assertRaisesRegex(ValueError, "cornerstone column"):
                WorkbookExpansionWriter(cfg=_cfg(root, workbook_path)).write_expansion()

    def test_column_subset_rewrites_regimes_and_leaves_other_columns(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook_path = root / "cohort.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "NLR"
            sheet.append(["File", "Code", "Name", "Prospectus", "Listing"])
            sheet.append([1, "1234.HK", "Selected", datetime(2024, 12, 30), datetime(2025, 1, 8)])
            sheet.cell(2, 162, "Curated Stabilizing Manager")
            sheet.cell(2, 202, "POST_2025_REFORM")
            workbook.save(workbook_path)
            workbook.close()

            out_master = root / "out" / "master"
            out_master.mkdir(parents=True)
            for name, fields in (
                ("stabilization_events.csv", ["stock_code"]),
                ("horizon_summary.csv", ["stock_code", "horizon"]),
                ("daily_market_panel.csv", ["stock_code", "amihud_illiq", "zero_volume_flag", "daily_return", "max_drawdown"]),
                ("lockup_events.csv", ["stock_code", "lockup_category"]),
            ):
                with (out_master / name).open("w", newline="", encoding="utf-8-sig") as stream:
                    csv.DictWriter(stream, fieldnames=fields).writeheader()

            cfg = {
                "workbook": str(workbook_path),
                "workbook_path": workbook_path,
                "_workbook_path": workbook_path,
                "sheet": "NLR",
                "data_start_row": 2,
                "id_columns": {"file_no": "A", "stock_code": "B", "name": "C", "prospectus_date": "D", "listing_date": "E"},
                "dataset": {"period_start": "2025-01-01", "period_end": "2025-03-31"},
                "filter_period": True,
                "paths": {"out": root / "out", "allot_out": root / "allot_out"},
            }
            writer = WorkbookExpansionWriter(cfg=cfg)
            with self.assertRaises(ValueError):
                writer.write_expansion({999})
            self.assertEqual(writer.write_expansion({201, 202}), 0)

            check = openpyxl.load_workbook(workbook_path, data_only=True)
            output = check["NLR"]
            self.assertEqual(output.cell(2, 162).value, "Curated Stabilizing Manager")
            self.assertEqual(output.cell(2, 201).value, "POST_FINI")
            self.assertEqual(output.cell(2, 202).value, "PRE_2025_REFORM")
            check.close()



class OverwriteGuardTests(unittest.TestCase):
    CURATED = {163: "2025-02-08", 171: 0.042, 185: "2025-07-10", 190: "Real Sponsor Ltd"}

    def test_refuses_to_clear_curated_cells_and_leaves_workbook_untouched(self):
        with tempfile.TemporaryDirectory() as tmp:
            workbook_path, cfg = _build_cohort(Path(tmp), curated=self.CURATED)
            before = workbook_path.read_bytes()
            with self.assertRaises(ExpansionOverwriteError) as ctx:
                WorkbookExpansionWriter(cfg=cfg).write_expansion()
            self.assertEqual(workbook_path.read_bytes(), before)
            self.assertEqual(sorted(c.col for c in ctx.exception.conflicts), sorted(self.CURATED))
            self.assertIn("--force-overwrite", str(ctx.exception))

    def test_force_overwrite_clears_curated_cells(self):
        with tempfile.TemporaryDirectory() as tmp:
            workbook_path, cfg = _build_cohort(Path(tmp), curated=self.CURATED)
            self.assertEqual(WorkbookExpansionWriter(cfg=cfg).write_expansion(force_overwrite=True), 0)
            self.assertEqual(_row2(workbook_path, *self.CURATED), [None] * len(self.CURATED))

    def test_source_backed_values_may_replace_curated_cells(self):
        with tempfile.TemporaryDirectory() as tmp:
            workbook_path, cfg = _build_cohort(
                Path(tmp),
                stab_rows=[{"stock_code": "1234.HK", "stabilizing_manager": "New Manager",
                            "stabilization_period_end": "2026-02-09"}],
                curated={162: "Old Manager", 163: "2026-02-01"},
            )
            WorkbookExpansionWriter(cfg=cfg).write_expansion()
            self.assertEqual(_row2(workbook_path, 162, 163), ["New Manager", "2026-02-09"])

    def test_placeholder_source_value_is_refused_even_into_empty_cell(self):
        with tempfile.TemporaryDirectory() as tmp:
            workbook_path, cfg = _build_cohort(Path(tmp), stab_rows=[{
                "stock_code": "1234.HK", "stabilizing_manager": "China International Capital Corporation / Sponsor-OC",
                "stabilization_period_end": "2026-02-09"}])
            before = workbook_path.read_bytes()
            with self.assertRaises(ExpansionOverwriteError) as ctx:
                WorkbookExpansionWriter(cfg=cfg).write_expansion()
            self.assertEqual([c.col for c in ctx.exception.conflicts], [162])
            self.assertEqual(workbook_path.read_bytes(), before)

    def test_legacy_placeholder_cell_may_be_cleared(self):
        with tempfile.TemporaryDirectory() as tmp:
            workbook_path, cfg = _build_cohort(Path(tmp), curated={162: "CICC / Sponsor-OC"})
            WorkbookExpansionWriter(cfg=cfg).write_expansion()
            self.assertEqual(_row2(workbook_path, 162), [None])

    def test_dry_run_never_saves(self):
        with tempfile.TemporaryDirectory() as tmp:
            workbook_path, cfg = _build_cohort(Path(tmp), stab_rows=[{
                "stock_code": "1234.HK", "stabilizing_manager": "Example Sponsor",
                "stabilization_period_end": "2026-02-09"}])
            before = workbook_path.read_bytes()
            self.assertEqual(WorkbookExpansionWriter(cfg=cfg).write_expansion(dry_run=True), 0)
            self.assertEqual(workbook_path.read_bytes(), before)

            (Path(tmp) / "curated").mkdir()
            curated_path, curated_cfg = _build_cohort(Path(tmp) / "curated", curated=self.CURATED)
            curated_before = curated_path.read_bytes()
            with self.assertRaises(ExpansionOverwriteError):
                WorkbookExpansionWriter(cfg=curated_cfg).write_expansion(dry_run=True)
            self.assertEqual(curated_path.read_bytes(), curated_before)

    def test_find_overwrite_conflicts_rules(self):
        sheet = openpyxl.Workbook().active
        sheet.cell(2, 162).value = "Curated"
        sheet.cell(2, 163).value = "   "
        sheet.cell(2, 164).value = 0
        sheet.cell(2, 165).value = "CICC / Sponsor-OC"
        planned = [
            (2, "1234.HK", 162, None),        # 清空人工值 → 冲突
            (2, "1234.HK", 163, None),        # 现有为空白 → 允许
            (2, "1234.HK", 164, ""),          # 0 是真实值，空串覆盖 → 冲突
            (2, "1234.HK", 165, None),        # 清除旧占位值 → 允许
            (2, "1234.HK", 166, "X / Sponsor-OC"),  # 占位值写入空格 → 冲突
            (2, "1234.HK", 165, "CICC / Sponsor-OC"),  # 与现值相同 → 无变更
        ]
        conflicts = find_overwrite_conflicts(sheet, planned)
        self.assertEqual([c.col for c in conflicts], [162, 164, 166])
        self.assertEqual(conflicts[0], OverwriteConflict(2, "1234.HK", 162, "Curated", None))

if __name__ == "__main__":
    unittest.main()
