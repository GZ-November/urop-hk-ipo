import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / "src")]
WS = ROOT.parent
WORKBOOK_PATH = WS / "cohorts" / "HKIPO-MB2026Q1.xlsx"
FIXTURE_WORKBOOK = ROOT / "tests" / "fixtures" / "mock_workbook.xlsx"

from report import generate_report
from run import load_cfg


class ReportTests(unittest.TestCase):
    def config_for(self, workbook=FIXTURE_WORKBOOK, dataset_id="TEST_FIXTURE", cohort="TEST FIXTURE"):
        cfg = load_cfg(config_path=ROOT / "config.yaml")
        cfg["workbook_path"] = workbook
        cfg["dataset"] = {**cfg.get("dataset", {}), "id": dataset_id, "cohort": cohort}
        return cfg

    def test_generate_report_with_mock_fixture(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            cfg = self.config_for()
            out_md = Path(tmp_dir) / "mock_report.md"
            stats = generate_report(cfg, out_path=out_md)
            self.assertEqual(stats["n_companies"], 2)
            self.assertTrue(out_md.exists())
            content = out_md.read_text(encoding="utf-8")
            self.assertIn("香港主板 TEST FIXTURE IPO 全景学术与市场分析报告", content)
            self.assertIn("2 家", content)

    def test_missing_returns_and_actual_offer_price(self):
        import math
        import openpyxl
        with tempfile.TemporaryDirectory() as temporary:
            book = Path(temporary) / "report.xlsx"
            workbook = openpyxl.load_workbook(FIXTURE_WORKBOOK)
            sheet = workbook["NLR"]
            headers = {str(sheet.cell(1, c).value or "").strip(): c for c in range(1, sheet.max_column + 1)}
            offer = next(c for h, c in headers.items() if "IPO Subscription Price" in h)
            close = headers["First trading day closing price (HK$)"]
            for row in range(2, sheet.max_row + 1):
                if sheet.cell(row, 2).value:
                    sheet.cell(row, offer).value = 10
                    sheet.cell(row, close).value = 20
            workbook.save(book)
            stats = generate_report(self.config_for(book), out_path=Path(temporary) / "report.md")
            self.assertAlmostEqual(stats["avg_first_day_return"], 1.0)
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, close).value = None
            workbook.save(book)
            workbook.close()
            stats = generate_report(self.config_for(book), out_path=Path(temporary) / "missing.md")
            self.assertTrue(math.isnan(stats["avg_first_day_return"]))

    @unittest.skipUnless(WORKBOOK_PATH.exists(), "Requires local production dataset HKIPO-MB2026Q1.xlsx")
    def test_generate_report_metrics_production(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            out_md = Path(tmp_dir) / "prod_report.md"
            cfg = self.config_for(WORKBOOK_PATH, dataset_id="HKIPO-MB2026Q1", cohort="2026 Q1")
            stats = generate_report(cfg, out_path=out_md)
            self.assertEqual(stats["n_companies"], 38)
            self.assertGreater(stats["total_net_proceeds_hkd"], 90e9)  # > 90 billion HKD
            self.assertGreater(stats["total_market_cap_hkd"], 1e12)     # > 1 trillion HKD
            self.assertGreater(stats["cornerstone_coverage_pct"], 0.8)  # > 80% coverage
            self.assertEqual(stats["pathways"]["ch18c_specialist_tech"], 6)
            self.assertEqual(stats["pathways"]["ch18a_biotech"], 3)
            self.assertEqual(stats["pathways"]["ah_dual_listing"], 15)  # 含招股书兜底的 2701、9611、9980

            # Verify markdown report exists and has substantial content
            self.assertTrue(out_md.exists())
            content = out_md.read_text(encoding="utf-8")
            self.assertIn("香港主板 2026 Q1 IPO 全景学术与市场分析报告", content)
            self.assertIn("38 家", content)
            self.assertIn("資訊科技業", content)


if __name__ == "__main__":
    unittest.main()
