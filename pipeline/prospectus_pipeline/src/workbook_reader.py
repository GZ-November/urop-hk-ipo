#!/usr/bin/env python3
"""NLR 工作簿读取的唯一 seam：表头规范化、表头映射、发行人行与代码列。

此前"打开 NLR → 规范化表头 → 迭代发行人行"的 implementation 在
cohort / codebook / exclusions / market_panel / enrich_master_dataset /
write_back / validate 七处各有一份变体，且 audit（只读）反向依赖
write_back（写回）的工具函数。本 module 收敛为一个 interface；
写回与审计共享同一实现，口径变更只改这里。
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import Any, Iterator

import openpyxl

DEFAULT_SHEET = "NLR"
CODE_COLUMN = 2          # Stock Code
LISTING_DATE_COLUMN = 5  # Date of Listing (dd/mm/yy)


def norm_header(v: Any) -> str:
    """表头单元格 -> 规范化键（压缩空白 + 小写）。"""
    return " ".join(str(v or "").replace("\n", " ").split()).strip().lower()


def open_sheet(path: Path | str, sheet: str = DEFAULT_SHEET):
    """只读模式打开工作簿并返回指定 sheet（调用方负责 close）。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return wb, wb[sheet]
    except Exception:
        wb.close()
        raise


def header_map(ws) -> dict[str, list[int]]:
    """表头行 -> {规范化表头: [1 基列号, ...]}。"""
    mapping: dict[str, list[int]] = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(1, col).value
        if val in (None, ""):
            continue
        mapping.setdefault(norm_header(val), []).append(col)
    return mapping


def iter_issuer_rows(
    path: Path | str,
    sheet: str = DEFAULT_SHEET,
    code_col: int = CODE_COLUMN,
    start_row: int = 2,
    max_col: int | None = None,
) -> Iterator[tuple[int, int, tuple[Any, ...]]]:
    """逐个发行人迭代 (行号, code 列号, 该行指定列区间的值)。

    跳过 code 列为空的行；max_col=None 时读到最后一列。
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        upper = max_col or ws.max_column
        for row_no, row in enumerate(
            ws.iter_rows(min_row=start_row, min_col=1, max_col=upper, values_only=True),
            start=start_row,
        ):
            code = row[code_col - 1] if code_col - 1 < len(row) else None
            if code in (None, ""):
                continue
            yield row_no, code_col, row
    finally:
        wb.close()


def read_codes(path: Path | str, sheet: str = DEFAULT_SHEET, code_col: int = CODE_COLUMN) -> list[str]:
    """数据行的股票代码列表（保序、去空白）。"""
    codes = []
    for _, _, row in iter_issuer_rows(path, sheet, code_col=code_col, max_col=code_col):
        codes.append(str(row[code_col - 1]).strip())
    return codes


def codes_and_listing_dates(
    path: Path | str, sheet: str = DEFAULT_SHEET
) -> list[tuple[str, dt.date | None]]:
    """(股票代码, 上市日期) 列表；上市日缺失或非法时为 None。"""
    out: list[tuple[str, dt.date | None]] = []
    for _, _, row in iter_issuer_rows(path, sheet, max_col=LISTING_DATE_COLUMN):
        listing = row[LISTING_DATE_COLUMN - 1]
        if isinstance(listing, dt.datetime):
            listing = listing.date()
        if not isinstance(listing, dt.date):
            listing = None
        out.append((str(row[CODE_COLUMN - 1]).strip(), listing))
    return out


def resolve_columns(ws, schema: dict) -> tuple[dict, list[str]]:
    """把 schema 的 key 通过规范化表头解析成真实列字母；返回 (映射, 问题列表)。"""
    header_to_cols: dict[str, list[str]] = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(1, col).value
        if val in (None, ""):
            continue
        if ws.cell(1, col).fill.fill_type != "solid":
            continue
        header_to_cols.setdefault(norm_header(val), []).append(
            openpyxl.utils.get_column_letter(col))
    mapping, issues = {}, []
    for field in schema["fields"]:
        cols = header_to_cols.get(norm_header(field["header"]), [])
        if len(cols) == 1:
            mapping[field["key"]] = cols[0]
        elif not cols:
            issues.append(f"{field['key']}: header not found: {field['header']!r}")
        else:
            issues.append(f"{field['key']}: header ambiguous: {field['header']!r} -> {cols}")
    return mapping, issues

