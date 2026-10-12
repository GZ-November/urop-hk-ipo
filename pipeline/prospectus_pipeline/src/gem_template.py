"""GEM student-layout view over canonical data and GEM-specific evidence fields."""
from __future__ import annotations

from copy import copy
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
WS = ROOT.parent
# Source headers, never physical schema keys: the 46 student columns are preserved.
ORIGINAL_TO_NLR = ["E", "B", "C", "K", "DA", None, "L", None, "AU", "CE", "AV", "F", "G",
                   "L", "M", "N", "O", "P", "Q", "R", "S", "T", "U", "Y", "Z", "AA", "AB",
                   "AC", "AD", "AE", "AF", "AG", "AH", "AI", "AJ", "AK", "AL", "AM", "AN",
                   "AO", "AP", "AQ", "BC", "AR", "AT", "EH"]

# Retain storage compatibility while excluding Main Board-only concepts from the GEM research view.
GEM_EXCLUDED_HEADERS = {"WVR flag", "Chapter 18A flag", "Chapter 18C flag",
                        "Technology commercialization stage", "2025 pricing reform regime",
                        "Controlling shareholder 6-month disposal lockup expiry date",
                        "Controlling shareholder 12-month cessation of control expiry date"}


def original_headers() -> list[str | None]:
    workbook = openpyxl.load_workbook(WS / "templates/HKIPO-MB-template-final.xlsx", read_only=True)
    try:
        headers = next(workbook["NLR"].iter_rows(max_row=1, values_only=True))
        return [headers[openpyxl.utils.column_index_from_string(letter) - 1] if letter else None
                for letter in ORIGINAL_TO_NLR]
    finally:
        workbook.close()


def common_columns(sheet) -> list[int]:
    mapped = set(original_headers())
    return [c for c in range(1, sheet.max_column + 1) if (sheet.cell(1, c).value not in mapped or str(sheet.cell(1, c).value or "").startswith("Underwriting Commission"))
            and str(sheet.cell(1, c).value or "").strip() not in GEM_EXCLUDED_HEADERS
            and "[Reserved]" not in str(sheet.cell(1, c).value or "")]


def detail_schema() -> dict:
    return json.loads((ROOT / "schema/gem_fields.json").read_text(encoding="utf-8"))


def update_gem_view(workbook) -> None:
    """Refresh values only; canonical NLR and independent GEM_details remain authoritative."""
    if "GEM" not in workbook.sheetnames:
        return
    view, source = workbook["GEM"], workbook["NLR"]
    extra_columns = common_columns(source)
    header_columns = {source.cell(1, c).value: c for c in range(1, source.max_column + 1)}
    headers = original_headers()
    for row in range(2, source.max_row + 1):
        if source.cell(row, 2).value in (None, ""):
            continue
        for column, header in enumerate(headers, 1):
            value = source.cell(row, header_columns[header]).value if header in header_columns else None
            if column == 44:
                rates = [source[f"{letter}{row}"].value for letter in ("AR", "AS")]
                disclosed = [float(rate) for rate in rates if isinstance(rate, (int, float))]
                positive = [rate for rate in disclosed if rate != 0]
                if positive and max(positive) - min(positive) > 1e-12:
                    value = "NA"  # unequal tranche rates do not define one overall fee
                else:
                    value = positive[0] if positive else (0 if disclosed else None)
            if column == 6:  # combined official gross funds: never assign total to HK or Int. tranche
                funds = [source[f"{letter}{row}"].value for letter in ("I", "J")]
                value = sum(x for x in funds if isinstance(x, (int, float))) if any(isinstance(x, (int, float)) for x in funds) else None
                if "GEM_details" in workbook.sheetnames:
                    official = workbook["GEM_details"].cell(row, 21).value
                    if isinstance(official, (int, float)):
                        value = official
            if column in (7, 8) and "GEM_details" in workbook.sheetnames:
                official = workbook["GEM_details"].cell(row, 22 if column == 7 else 23).value
                if isinstance(official, (int, float)):
                    value = official
            if column == 8 and value is None:
                shares, price = source[f"L{row}"].value, source[f"K{row}"].value
                if isinstance(shares, (int, float)) and isinstance(price, (int, float)):
                    value = shares * price
            view.cell(row, column).value = value
        for offset, column in enumerate(extra_columns, 47):
            view.cell(row, offset).value = source.cell(row, column).value
        details = workbook["GEM_details"]
        for column in range(1, 6):
            details.cell(row, column).value = source.cell(row, column).value
        for offset, field in enumerate(detail_schema()["fields"], 47 + len(extra_columns)):
            detail_column = 6 + detail_schema()["fields"].index(field)
            view.cell(row, offset).value = details.cell(row, detail_column).value
    workbook.active = workbook.sheetnames.index("GEM")


def build_template(destination: Path | None = None) -> Path:
    """Create a compact blank template; original historical GEM workbook is untouched."""
    destination = destination or WS / "templates/HKIPO-GEM-template-final.xlsx"
    original = openpyxl.load_workbook(WS / "templates/HKIPO-GEM-template-students.xlsx", read_only=True)
    canonical = openpyxl.load_workbook(WS / "templates/HKIPO-MB-template-final.xlsx", read_only=True)
    workbook = openpyxl.Workbook()
    view = workbook.active
    view.title = "GEM"
    nlr = workbook.create_sheet("NLR")
    details = workbook.create_sheet("GEM_details")
    for col in range(1, 203):
        source = canonical["NLR"].cell(1, col)
        nlr.cell(1, col).value = source.value
        for attribute in ("font", "fill", "border", "alignment", "protection", "number_format"):
            setattr(nlr.cell(1, col), attribute, copy(getattr(source, attribute)))
    for column in range(1, nlr.max_column + 1):
        if nlr.cell(1, column).value == "HK ordinary IPO count in 90 calendar days before prospectus":
            nlr.cell(1, column).value = "GEM ordinary IPO count in 90 calendar days before prospectus"
    for column in range(nlr.max_column, 0, -1):
        if str(nlr.cell(1, column).value or "").strip() in GEM_EXCLUDED_HEADERS:
            nlr.delete_cols(column)
    for col in range(1, 47):
        source = original.active.cell(1, col)
        view.cell(1, col).value = source.value
        for attribute in ("font", "fill", "border", "alignment", "protection", "number_format"):
            setattr(view.cell(1, col), attribute, copy(getattr(source, attribute)))
    extra_columns = common_columns(nlr)
    for offset, col in enumerate(extra_columns, 47):
        view.cell(1, offset).value = nlr.cell(1, col).value
    for col in range(1, 6):
        details.cell(1, col).value = nlr.cell(1, col).value
    for column, field in enumerate(detail_schema()["fields"], 6):
        details.cell(1, column).value = field["header"]
        view.cell(1, 47 + len(extra_columns) + column - 6).value = field["header"]
    for col, header in enumerate(("Official GEM funds raised (HK$)", "Official GEM shares at listing", "Official GEM market capitalisation (HK$)", "Official GEM source report", "GEM prospectus date source"), 21):
        details.cell(1, col).value = header
    for sheet in workbook:
        sheet.freeze_panes = "D2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            if cell.fill.fill_type != "solid":
                cell.fill = openpyxl.styles.PatternFill("solid", fgColor="D9EAF7")
            cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
            sheet.column_dimensions[cell.column_letter].width = 22
        sheet.row_dimensions[1].height = 75
    original.close()
    canonical.close()
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    workbook.close()
    return destination


def sync_gem_view(cfg: dict) -> None:
    from boards import normalize_board
    if normalize_board(cfg.get("dataset", {}).get("board")) != "gem":
        return
    from workbook_transaction import workbook_transaction
    book = Path(cfg["_workbook_path"])
    with workbook_transaction(book, operation="gem-view") as workbook:
        update_gem_view(workbook)
