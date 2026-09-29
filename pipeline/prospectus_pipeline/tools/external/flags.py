#!/usr/bin/env python3
"""从已有产物确定性推导 7 个空列（不花 token）。

推导来源：
  BH Listing board            ← col_AS（上市途径）是否含 Main Board / GEM
  BJ A+H issuer flag          ← col_AS 是否肯定表明已有 A 股上市（A+H / other listed shares / A Shares …，见 is_a_plus_h）
  BK WVR flag                 ← col_AS 是否肯定援引 "Chapter 8A"（忽略否定式表述）
  BL Chapter 18A flag         ← col_AS 是否肯定援引 "Chapter 18A"（忽略否定式表述）
  BM Chapter 18C flag         ← col_AS 是否肯定援引 "Chapter 18C"（忽略否定式表述）
  BQ Place of incorporation   ← 招股书前几页 "incorporated in ..."
  CL Earliest cornerstone unlock date ← 上市日 + 6 个月（港交所规则；
                                无基石者填 NA，依据配发抽取 col_CK == 0
                                或 cornerstone_absence.json verdict absent/none）

只赋值，不改 font/fill/number_format。写前备份。
证据落到 out/derived_flags.json 供审计。
"""
from __future__ import annotations

import datetime as dt
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2] if Path(__file__).resolve().parent.name == "external" else Path(__file__).resolve().parent
WS = ROOT.parent
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))

from contracts import normalize_code  # noqa: E402
from cornerstone import confirmed_absent  # noqa: E402
from run import load_cfg  # noqa: E402

HEADERS = {
    "BH": "Listing board",
    "BJ": "A+H issuer flag",
    "BK": "WVR flag",
    "BL": "Chapter 18A flag",
    "BM": "Chapter 18C flag",
    "BQ": "Place of incorporation",
    "CL": "Earliest cornerstone unlock date (dd/mm/yy)",
}
LOCKUP_MONTHS = 6


def norm(s) -> str:
    return " ".join(str(s or "").replace("\n", " ").split()).strip().lower()


def add_months(d: dt.date, months: int) -> dt.date:
    """月末安全加月：2026-08-31 + 6 → 2027-02-28。"""
    y, m = divmod(d.month - 1 + months, 12)
    y, m = d.year + y, m + 1
    import calendar
    return dt.date(y, m, min(d.day, calendar.monthrange(y, m)[1]))


def digits(code: str) -> str:
    return "".join(c for c in str(code) if c.isdigit())


NEGATION_BEFORE = re.compile(
    r"\b(?:not|no|neither|nor|without|except|other\s+than|rather\s+than)\b\s*"
    r"(?:[^.,;()]{0,40}\s+)?$",
    re.I,
)
NEGATION_AFTER = re.compile(
    r"^\s*(?:[^.,;()]{0,20}\s+)?(?:does\s+not|do\s+not|is\s+not|are\s+not|not)\b",
    re.I,
)


def cites_chapter(text: str, label: str) -> bool:
    """判断上市途径文本是否「肯定地」依据某一章上市。

    招股书抽取文本里大量出现否定式澄清，例如
      "Main Board standard listing (not Chapter 18C / 18A)"
      "H share listing; no Chapter 18C basis disclosed"
    单纯子串匹配会把这类表述误判为肯定，造成 18A/18C 标记假阳性
    （2025Q2 cohort 的 2605.HK、9678.HK 即因此被误标为 18C）。
    这里同时检查命中位置前后的否定词，命中否定式则跳过该处提及。
    """
    return _affirmed(text, re.compile(rf"Chapter\s+{label}", re.I))


def _affirmed(text: str, pattern: re.Pattern) -> bool:
    """pattern 在 text 中至少有一处命中不处于否定语境（前后 NEGATION_* 均不触发）。"""
    for m in pattern.finditer(text):
        before = text[max(0, m.start() - 60): m.start()]
        after = text[m.end(): m.end() + 40]
        if NEGATION_BEFORE.search(before) or NEGATION_AFTER.match(after):
            continue
        if re.search(r"[非无未不]\s*$", before):
            continue
        return True
    return False


# 上市途径文本里 A+H 发行人的几种写法（2025–2026 cohort 实测）：
#   "A+H dual listing ..." / "Chapter 19A PRC issuer (A+H)"
#   "Chapter 19A of the Listing Rules (PRC issuer with other listed shares)"
#   "PRC issuer with A Shares listed on the Shenzhen Stock Exchange"
#   "PRC joint stock company already listed on the SZSE main board"
# 仅写 "Chapter 19A PRC issuer" / "H Share issuer" 的不算——19A 同样适用于
# 只发 H 股、无 A 股的中国发行人。
A_PLUS_H_PATTERNS = [
    re.compile(r"\bA\s*\+\s*H\b"),
    re.compile(r"\bother\s+listed\s+shares\b", re.I),
    re.compile(r"(?<![\w.])A[-\s]?[Ss]hares?\b"),       # "A Shares"/"A-share"；大写 A 避开冠词 "a share"
    re.compile(r"A\s*股"),
    re.compile(r"\blisted\s+on\s+(?:the\s+)?(?:ChiNext|STAR\s+Market|SSE|SZSE|BSE|"
               r"(?:Shanghai|Shenzhen|Beijing)\s+Stock\s+Exchange)\b", re.I),
]


def is_a_plus_h(text: str) -> bool:
    """上市途径文本是否肯定地表明发行人已有 A 股上市（A+H）。

    旧规则只认字面量 "A+H"，漏掉了 2026Q2/Q3 抽取统一写成的
    "PRC issuer with other listed shares"（2476.HK、3308.HK 等 12 家）
    以及 "with A Shares listed on the Shenzhen Stock Exchange"（1081.HK）。
    """
    return any(_affirmed(text, p) for p in A_PLUS_H_PATTERNS)


def place_of_incorporation(code: str, text_dir: Path | None = None) -> tuple[str, str]:
    """扫前若干页判断注册地；返回 (值, 证据)。"""
    f = (text_dir or ROOT / "data" / "text") / f"HKIPO-MB{digits(code)}.jsonl"
    if not f.exists():
        return "NA", "缺页级文本"
    head = ""
    for i, line in enumerate(f.open(encoding="utf-8")):
        if i >= 6:
            break
        head += json.loads(line)["text"] + "\n"
    flat = " ".join(head.split())
    # 封面措辞不统一：incorporated / established / organised / formed / registered 都见过
    VERB = r"(?:incorporated|established|organi[sz]ed|formed|registered|constituted)"
    pats = [
        (rf"{VERB}\s+in\s+the\s+People.{{0,3}}s\s+Republic\s+of\s+China", "PRC"),
        (rf"{VERB}\s+in\s+(?:the\s+)?Cayman\s+Islands", "Cayman Islands"),
        (rf"{VERB}\s+in\s+(?:the\s+)?Bermuda", "Bermuda"),
        (rf"{VERB}\s+in\s+the\s+British\s+Virgin\s+Islands", "British Virgin Islands"),
        (rf"{VERB}\s+in\s+Hong\s+Kong", "Hong Kong"),
        # 兜底：直接看 "People's Republic of China with limited liability"
        (r"People.{0,3}s\s+Republic\s+of\s+China\s+with\s+limited\s+liability", "PRC"),
    ]
    for pat, val in pats:
        m = re.search(pat, flat, re.I)
        if m:
            return val, flat[max(0, m.start() - 60):m.end() + 60]
    return "NA", "前 6 页未见注册地表述"


import argparse
from workbook_transaction import workbook_transaction


def main() -> int:
    ap = argparse.ArgumentParser(description="Derive statutory flags")
    ap.add_argument("--config", help="Cohort config (otherwise PIPELINE_CONFIG)")
    ap.add_argument("--dry-run", action="store_true", help="Do not mutate workbook")
    ap.add_argument("--book", help="Path to workbook")
    ap.add_argument("--only", nargs="*", default=None, help="Filter by stock code(s)")
    args = ap.parse_args()
    cfg = load_cfg(config_path=args.config)

    book = Path(args.book) if args.book else cfg["_workbook_path"]
    dry = args.dry_run

    # 基石「确认无」的公司（配发抽取 col_CK == 0 是权威依据，见 cornerstone.confirmed_absent）
    no_cornerstone = confirmed_absent(cfg)

    wb_read = openpyxl.load_workbook(book, data_only=True)
    ws_read = wb_read[cfg["sheet"]]
    col_of = {}
    for c in range(1, ws_read.max_column + 1):
        v = ws_read.cell(1, c).value
        if v in (None, ""):
            continue
        for key, h in HEADERS.items():
            if norm(v) == norm(h):
                col_of[key] = c
    missing = [k for k in HEADERS if k not in col_of]
    if missing:
        wb_read.close()
        raise SystemExit(f"工作簿找不到列：{missing}")

    # 行号 + 上市日
    row_of, listing = {}, {}
    ci_code = openpyxl.utils.column_index_from_string(cfg["id_columns"]["stock_code"])
    ci_list = openpyxl.utils.column_index_from_string(cfg["id_columns"]["listing_date"])
    selected = {normalize_code(x) for x in args.only} if args.only else None
    for r in range(cfg["data_start_row"], ws_read.max_row + 1):
        code = ws_read.cell(r, ci_code).value
        if code in (None, ""):
            continue
        norm_c = normalize_code(str(code).strip())
        if selected is not None and norm_c not in selected:
            continue
        row_of[norm_c] = r
        ld = ws_read.cell(r, ci_list).value
        if isinstance(ld, dt.datetime):
            ld = ld.date()
        listing[norm_c] = ld if isinstance(ld, dt.date) else None
    wb_read.close()
    if not row_of:
        raise SystemExit("没有匹配的公司；检查 --config、--book 和 --only 参数")
    if selected is not None and selected != set(row_of):
        raise SystemExit(f"--only 中有代码未匹配工作簿：{sorted(selected - set(row_of))}")

    audit, n = {}, 0
    print(f"{'code':9s} {'board':10s} {'A+H':4s} {'WVR':4s} {'18A':4s} {'18C':4s} "
          f"{'注册地':14s} 基石解禁")
    for code, r in sorted(row_of.items()):
        jf = cfg["paths"]["out"] / "extracted" / f"HKIPO-MB{digits(code)}.json"
        if not jf.exists():
            raise SystemExit(f"{code}: 缺抽取 JSON：{jf}")
        asv = str(json.loads(jf.read_text(encoding="utf-8"))["fields"]["col_AS"]["value"])
        board = "GEM" if re.search(r"\bGEM\b", asv, re.I) else "Main Board"
        a_plus_h = 1 if is_a_plus_h(asv) else 0
        wvr = 1 if cites_chapter(asv, "8A") else 0
        c18a = 1 if cites_chapter(asv, "18A") else 0
        c18c = 1 if cites_chapter(asv, "18C") else 0
        inc, inc_ev = place_of_incorporation(code, cfg["paths"]["text"])
        ld = listing.get(code)
        if code in no_cornerstone:
            unlock, unlock_note = "NA", "确认无基石，无解禁日"
        elif ld:
            unlock = add_months(ld, LOCKUP_MONTHS)
            unlock_note = f"上市日 {ld} + {LOCKUP_MONTHS} 个月"
        else:
            unlock, unlock_note = "NA", "缺上市日"
        vals = {"BH": board, "BJ": a_plus_h, "BK": wvr, "BL": c18a, "BM": c18c,
                "BQ": inc, "CL": unlock}
        audit[code] = {"values": {k: str(v) for k, v in vals.items()},
                       "evidence": {"col_AS": asv, "incorporation": inc_ev,
                                    "unlock": unlock_note}}
        n += 1
        print(f"{code:9s} {board:10s} {a_plus_h:4d} {wvr:4d} {c18a:4d} {c18c:4d} "
              f"{inc:14s} {str(unlock)}")

    outj = cfg["paths"]["out"] / "derived_flags.json"
    outj.write_text(json.dumps(audit, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
                    encoding="utf-8")
    print(f"\n推导 {n} 家；证据 -> {outj}")

    if dry:
        print("--dry-run：未写回")
        return 0

    with workbook_transaction(book, operation="flags") as wb:
        ws = wb[cfg["sheet"]]
        for code, rec in audit.items():
            r = row_of[code]
            for k in ["BH", "BJ", "BK", "BL", "BM", "BQ"]:
                v = rec["values"][k]
                ws.cell(r, col_of[k]).value = int(v) if v.lstrip("-").isdigit() else v
            u = rec["values"]["CL"]
            if u == "NA":
                ws.cell(r, col_of["CL"]).value = "NA"
            else:
                y, m, d = (int(x) for x in u.split("-"))
                ws.cell(r, col_of["CL"]).value = dt.date(y, m, d)

    print(f"已写回 7 列（{len(audit)} 家） -> {book.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
